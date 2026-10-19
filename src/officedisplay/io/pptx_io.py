from __future__ import annotations

import tempfile
from pathlib import Path

from officedisplay.model.presentation import (
    Background,
    Presentation,
    Slide,
    SlideObject,
    flatten,
    new_id,
)


EMU_PER_PIXEL = 9525  # 96 dpi


def _emu(value: float) -> int:
    return int(value * EMU_PER_PIXEL)


def export_pptx(presentation: Presentation, path: str | Path) -> None:
    from pptx import Presentation as PptxPresentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu

    package = PptxPresentation()
    package.slide_width = Emu(_emu(presentation.width))
    package.slide_height = Emu(_emu(presentation.height))
    blank = package.slide_layouts[6]
    shapes = {
        "rect": MSO_SHAPE.RECTANGLE,
        "roundrect": MSO_SHAPE.ROUNDED_RECTANGLE,
        "ellipse": MSO_SHAPE.OVAL,
        "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
        "diamond": MSO_SHAPE.DIAMOND,
        "arrow": MSO_SHAPE.RIGHT_ARROW,
        "star": MSO_SHAPE.STAR_5_POINT,
        "line": MSO_SHAPE.LINE_INVERSE,
    }
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    for slide in presentation.slides:
        page = package.slides.add_slide(blank)
        _paint_background(page, slide, presentation)
        for obj, x, y, w, h in flatten(slide):
            left, top, width, height = Emu(_emu(x * presentation.width)), Emu(_emu(y * presentation.height)), Emu(_emu(max(w, 0.01) * presentation.width)), Emu(_emu(max(h, 0.01) * presentation.height))
            if obj.kind == "text":
                box = page.shapes.add_textbox(left, top, width, height)
                frame = box.text_frame
                frame.word_wrap = True
                box.text_frame.paragraphs[0].text = str(obj.props.get("text") or "")
                paragraph = frame.paragraphs[0]
                paragraph.alignment = aligns.get(obj.props.get("align") or "left", PP_ALIGN.LEFT)
                font = paragraph.font
                font.size = _pt(obj.props.get("font_size") or 18)
                font.bold = bool(obj.props.get("bold"))
                font.italic = bool(obj.props.get("italic"))
                font.underline = bool(obj.props.get("underline"))
                font.name = obj.props.get("font_family") or "Calibri"
                font.color.rgb = _rgb(obj.props.get("color") or "#1B1F2A")
                if obj.name:
                    box.name = obj.name
            elif obj.kind == "image" and obj.props.get("path") and Path(obj.props["path"]).is_file():
                picture = page.shapes.add_picture(obj.props["path"], left, top, width, height)
                if obj.name:
                    picture.name = obj.name
            elif obj.kind == "shape":
                kind = shapes.get(obj.props.get("shape") or "rect", MSO_SHAPE.RECTANGLE)
                shape = page.shapes.add_shape(kind, left, top, width, height)
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(obj.props.get("fill") or "#0B63E5")
                if obj.name:
                    shape.name = obj.name
            elif obj.kind == "chart":
                _chart(page, obj, left, top, width, height)
        if slide.notes:
            notes = page.notes_slide
            notes.notes_text_frame.text = slide.notes
    package.save(str(path))


def _pt(size: int):
    from pptx.util import Pt

    return Pt(max(8, min(int(size), 96)))


def _rgb(color: str):
    from pptx.dml.color import RGBColor

    hex_color = (color or "#000000").replace("#", "")
    if len(hex_color) == 8:
        hex_color = hex_color[2:]
    if len(hex_color) != 6:
        hex_color = "000000"
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def _paint_background(page, slide: Slide, presentation: Presentation) -> None:
    fill = page.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(slide.background.color or "#FFFFFF")


def _chart(page, obj, left, top, width, height) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    data = CategoryChartData()
    data.categories = obj.props.get("categories") or ["A", "B"]
    for serie in obj.props.get("series") or [{"name": "Werte", "values": [1, 2]}]:
        data.add_series(serie.get("name") or "Werte", serie.get("values") or [0])
    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
    if obj.props.get("chart_type") == "pie":
        chart_type = XL_CHART_TYPE.PIE
    elif obj.props.get("chart_type") == "line":
        chart_type = XL_CHART_TYPE.LINE
    page.shapes.add_chart(chart_type, left, top, width, height, data)


class PptxReadError(ValueError):
    """Die Datei ist keine lesbare PowerPoint-Präsentation."""


def import_pptx(path: str | Path, media_dir: str | Path | None = None) -> Presentation:
    from pptx import Presentation as PptxPresentation

    source = Path(path)
    try:
        package = PptxPresentation(str(source))
    except Exception as exc:
        raise PptxReadError(f"Die PowerPoint-Datei lässt sich nicht laden: {source.name}") from exc
    slide_w = float(package.slide_width or 12192000)
    slide_h = float(package.slide_height or 6858000)
    width = max(1, int(slide_w / EMU_PER_PIXEL))
    height = max(1, int(slide_h / EMU_PER_PIXEL))
    presentation = Presentation(title=source.stem, width=width, height=height)
    presentation.slides = []
    media = {"folder": None, "hint": media_dir, "source": source}
    for index, page in enumerate(package.slides, start=1):
        slide = Slide(id=new_id(), name=f"Folie {index}", background=Background(color="#FFFFFF"))
        try:
            notes = page.notes_slide.notes_text_frame.text
            slide.notes = notes or ""
        except Exception:
            slide.notes = ""
        try:
            shapes = list(page.shapes)
        except Exception:
            shapes = []
        shapes.extend(_alternate_shapes(getattr(page, "shapes", None)))
        for shape in shapes:
            _take_shape(shape, slide, slide_w, slide_h, media)
        presentation.slides.append(slide)
    if not presentation.slides:
        presentation.slides.append(Slide(id=new_id(), name="Folie 1"))
    return presentation


def _take_shape(shape, slide: Slide, slide_w: float, slide_h: float, media: dict, transforms=()) -> None:
    """Eine Form übernehmen. Unbekannte oder defekte Formen werden übersprungen."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        kind = shape.shape_type
    except Exception:
        return
    if kind == MSO_SHAPE_TYPE.GROUP:
        try:
            children = list(shape.shapes)
        except Exception:
            children = []
        children.extend(_alternate_shapes(getattr(shape, "shapes", None)))
        frame = _group_frame(shape)
        nested = transforms + (frame,) if frame is not None else transforms
        for child in children:
            _take_shape(child, slide, slide_w, slide_h, media, nested)
        return
    box = _box(shape, slide_w, slide_h, transforms)
    if kind == MSO_SHAPE_TYPE.PICTURE:
        if box is not None:
            _take_picture(shape, slide, box, media)
        return
    if getattr(shape, "has_table", False):
        if box is not None:
            _take_table(shape, slide, box)
        return
    text = _shape_text(shape)
    if kind == MSO_SHAPE_TYPE.TEXT_BOX:
        if text and box is not None:
            slide.objects.append(_text_object(shape, box, text))
        return
    if kind == MSO_SHAPE_TYPE.PLACEHOLDER or getattr(shape, "is_placeholder", False):
        # PowerPoint lässt die Maße oft weg. python-pptx liefert dann None,
        # der Text steht aber im Platzhalter. Ohne Kasten bleibt die Folie leer.
        if box is None and text:
            box = _placeholder_fallback(shape)
        if text and box is not None:
            slide.objects.append(_text_object(shape, box, text))
        elif box is not None:
            _take_picture(shape, slide, box, media)
        return
    if kind in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE}:
        if box is None:
            return
        preset = "line" if kind == MSO_SHAPE_TYPE.LINE else _auto_kind(shape)
        slide.objects.append(_shape_object(shape, box, preset))
        if text:
            slide.objects.append(_text_object(shape, box, text))
        return
    if text and box is not None:
        slide.objects.append(_text_object(shape, box, text))


def _box(shape, slide_w: float, slide_h: float, transforms=()):
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if left is None or top is None or width is None or height is None:
        return None
    if slide_w <= 0 or slide_h <= 0:
        return None
    x, y, w, h = float(left), float(top), float(width), float(height)
    for off_x, off_y, ch_x, ch_y, scale_x, scale_y in reversed(tuple(transforms)):
        x = off_x + (x - ch_x) * scale_x
        y = off_y + (y - ch_y) * scale_y
        w = abs(w * scale_x)
        h = abs(h * scale_y)
    return (
        x / slide_w,
        y / slide_h,
        max(0.01, w / slide_w),
        max(0.01, h / slide_h),
    )


def _group_frame(shape):
    """Gruppenkinder stehen im lokalen chOff/chExt-Raum, nicht in Folien-EMUs."""
    element = getattr(shape, "_element", None)
    xfrm = getattr(element, "xfrm", None) if element is not None else None
    if xfrm is None:
        props = getattr(element, "grpSpPr", None)
        xfrm = getattr(props, "xfrm", None) if props is not None else None
    if xfrm is None:
        return None

    def num(value) -> float:
        if value is None:
            return 0.0
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    off_x, off_y = num(getattr(xfrm, "x", None)), num(getattr(xfrm, "y", None))
    ext_cx, ext_cy = num(getattr(xfrm, "cx", None)), num(getattr(xfrm, "cy", None))
    ch_off, ch_ext = getattr(xfrm, "chOff", None), getattr(xfrm, "chExt", None)
    ch_x = num(getattr(ch_off, "x", None)) if ch_off is not None else 0.0
    ch_y = num(getattr(ch_off, "y", None)) if ch_off is not None else 0.0
    ch_cx = num(getattr(ch_ext, "cx", None)) if ch_ext is not None else 0.0
    ch_cy = num(getattr(ch_ext, "cy", None)) if ch_ext is not None else 0.0
    if ext_cx <= 0 or ext_cy <= 0 or ch_cx <= 0 or ch_cy <= 0:
        return None
    return (off_x, off_y, ch_x, ch_y, ext_cx / ch_cx, ext_cy / ch_cy)


def _placeholder_fallback(shape) -> tuple[float, float, float, float]:
    ph_type = None
    index = 0
    try:
        fmt = shape.placeholder_format
        ph_type = fmt.type
        index = int(getattr(fmt, "idx", 0) or 0)
    except Exception:
        ph_type = None
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER

        if ph_type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}:
            return (0.06, 0.12, 0.88, 0.22)
        if ph_type == PP_PLACEHOLDER.SUBTITLE:
            return (0.08, 0.38, 0.84, 0.16)
    except Exception:
        pass
    return (0.08, min(0.72, 0.36 + max(0, index - 1) * 0.06), 0.84, 0.46)


def _alternate_shapes(collection) -> list:
    """python-pptx überspringt AlternateContent. Fallback-Formen trotzdem lesen."""
    if collection is None:
        return []
    try:
        from pptx.oxml.ns import qn

        tree = collection._spTree
        factory = collection._shape_factory
        alternate = qn("ve:AlternateContent")
        fallback = qn("ve:Fallback")
        choice = qn("ve:Choice")
        tags = {qn("p:sp"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp"), qn("p:pic")}
        found = []
        for child in list(tree):
            if child.tag != alternate:
                continue
            holder = None
            for name in (fallback, choice):
                candidate = child.find(name)
                if candidate is not None and any(element.tag in tags for element in candidate):
                    holder = candidate
                    break
            if holder is None:
                continue
            for element in holder:
                if element.tag not in tags:
                    continue
                try:
                    found.append(factory(element))
                except Exception:
                    continue
        return found
    except Exception:
        return []


def _shape_text(shape) -> str:
    try:
        if not shape.has_text_frame:
            return ""
        return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
    except Exception:
        return ""


def _font_size(shape) -> int:
    try:
        size = shape.text_frame.paragraphs[0].font.size
        if size is not None:
            return max(8, min(int(size.pt), 96))
    except Exception:
        pass
    return 18


def _text_object(shape, box, text: str) -> SlideObject:
    x, y, w, h = box
    return SlideObject(
        id=new_id(),
        kind="text",
        x=x,
        y=y,
        w=w,
        h=h,
        name=getattr(shape, "name", "") or "",
        props={"text": text, "font_size": _font_size(shape), "color": "#1B1F2A", "font_family": "Calibri"},
    )


def _auto_kind(shape) -> str:
    try:
        from pptx.enum.shapes import MSO_SHAPE

        preset = shape.auto_shape_type
    except Exception:
        return "rect"
    pairs = (
        ("RECTANGLE", "rect"),
        ("ROUNDED_RECTANGLE", "roundrect"),
        ("OVAL", "ellipse"),
        ("ISOSCELES_TRIANGLE", "triangle"),
        ("DIAMOND", "diamond"),
        ("RIGHT_ARROW", "arrow"),
        ("STAR_5_POINT", "star"),
        ("LINE_INVERSE", "line"),
        ("STRAIGHT_CONNECTOR_1", "line"),
    )
    for enum_name, kind in pairs:
        member = getattr(MSO_SHAPE, enum_name, None)
        if member is not None and preset == member:
            return kind
    return "rect"


def _fill_hex(shape) -> str:
    try:
        from pptx.enum.dml import MSO_FILL

        fill = shape.fill
        if fill.type != MSO_FILL.SOLID:
            return "#0B63E5"
        rgb = fill.fore_color.rgb
        return "#{:02X}{:02X}{:02X}".format(int(rgb[0]), int(rgb[1]), int(rgb[2]))
    except Exception:
        return "#0B63E5"


def _shape_object(shape, box, preset: str) -> SlideObject:
    x, y, w, h = box
    return SlideObject(
        id=new_id(),
        kind="shape",
        x=x,
        y=y,
        w=w,
        h=h,
        name=getattr(shape, "name", "") or preset,
        props={"shape": preset, "fill": _fill_hex(shape), "stroke": "#00000000", "stroke_width": 0},
    )


def _take_picture(shape, slide: Slide, box, media: dict) -> None:
    try:
        image = shape.image
        blob = image.blob
        ext = image.ext or "png"
    except Exception:
        _picture_stand_in(shape, slide, box)
        return
    folder = _media_folder(media)
    file_path = folder / f"{new_id()}.{ext}"
    try:
        file_path.write_bytes(blob)
    except OSError:
        _picture_stand_in(shape, slide, box)
        return
    x, y, w, h = box
    slide.objects.append(
        SlideObject(
            id=new_id(),
            kind="image",
            x=x,
            y=y,
            w=w,
            h=h,
            name=getattr(shape, "name", "") or "",
            props={"path": str(file_path)},
        )
    )


def _picture_stand_in(shape, slide: Slide, box) -> None:
    """Bild ohne Blob bleibt als Fläche sichtbar, die Folie wird nicht leer."""
    x, y, w, h = box
    slide.objects.append(
        SlideObject(
            id=new_id(),
            kind="shape",
            x=x,
            y=y,
            w=w,
            h=h,
            name=getattr(shape, "name", "") or "bild",
            props={"shape": "rect", "fill": "#E8EEF8", "stroke": "#0B63E5", "stroke_width": 2},
        )
    )


def _take_table(shape, slide: Slide, box) -> None:
    lines = []
    try:
        for row in shape.table.rows:
            cells = []
            for cell in row.cells:
                cells.append((cell.text or "").strip())
            if any(cells):
                lines.append(" | ".join(cells))
    except Exception:
        return
    text = "\n".join(lines).strip()
    if text:
        slide.objects.append(_text_object(shape, box, text))


def _media_folder(media: dict) -> Path:
    if media.get("folder") is not None:
        return media["folder"]
    hint = media.get("hint")
    source = Path(media["source"])
    candidates = []
    if hint:
        candidates.append(Path(hint))
    else:
        candidates.append(Path(str(source.with_suffix("")) + "_media"))
    for folder in candidates:
        try:
            folder.mkdir(parents=True, exist_ok=True)
            probe = folder / ".od-write"
            probe.write_bytes(b"ok")
            probe.unlink()
            media["folder"] = folder
            return folder
        except OSError:
            continue
    folder = Path(tempfile.mkdtemp(prefix="officedisplay-pptx-"))
    media["folder"] = folder
    return folder
