from __future__ import annotations

import csv
from pathlib import Path


def read_table(path: str, sheet: str | None = None, limit: int = 500) -> tuple[list[str], list[dict]]:
    suffix = Path(path).suffix.lower()
    if suffix in {".csv", ".txt"}:
        return _csv(path, limit)
    return _xlsx(path, sheet, limit)


def _csv(path: str, limit: int) -> tuple[list[str], list[dict]]:
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        rows = list(csv.reader(handle, delimiter=delimiter))
    return _rows(rows, limit)


def _xlsx(path: str, sheet: str | None, limit: int) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        if index > limit:
            break
        rows.append([("" if cell is None else cell) for cell in row])
    workbook.close()
    return _rows(rows, limit)


def _rows(rows: list[list], limit: int) -> tuple[list[str], list[dict]]:
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        return [], []
    header = [str(cell).strip() or f"Spalte {i + 1}" for i, cell in enumerate(rows[0])]
    series = [{"name": name, "values": []} for name in header[1:]] or [{"name": "Werte", "values": []}]
    categories = []
    for row in rows[1 : limit + 1]:
        categories.append(str(row[0]) if row else "")
        for index, serie in enumerate(series):
            cell = row[index + 1] if index + 1 < len(row) else 0
            try:
                serie["values"].append(float(str(cell).replace(",", ".")))
            except (TypeError, ValueError):
                serie["values"].append(0.0)
    return categories, series


def bind_chart(obj, path: str, sheet: str | None = None) -> None:
    categories, series = read_table(path, sheet)
    obj.kind = "chart"
    obj.props["chart_type"] = obj.props.get("chart_type") or "bar"
    obj.props["categories"] = categories
    obj.props["series"] = series
    obj.props["binding"] = {"path": path, "sheet": sheet}
    obj.props["example"] = False


def refresh_bindings(presentation) -> int:
    count = 0
    for slide in presentation.slides:
        for obj in slide.objects:
            binding = (obj.props or {}).get("binding")
            if obj.kind == "chart" and binding and binding.get("path"):
                bind_chart(obj, binding["path"], binding.get("sheet"))
                count += 1
    return count
