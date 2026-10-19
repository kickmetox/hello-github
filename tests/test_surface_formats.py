import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


class SurfaceFormatTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def _pptx_bytes(self, folder: Path) -> bytes:
        from pptx import Presentation

        package = Presentation()
        slide = package.slides.add_slide(package.slide_layouts[0])
        slide.shapes.title.text = "Titel aus Paket"
        target = folder / "quelle.pptx"
        package.save(str(target))
        return target.read_bytes()

    def test_potx_or_ppsx_keeps_title(self):
        from officedisplay.io.open_file import load_presentation

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            blob = self._pptx_bytes(root)
            potx = root / "vorlage.potx"
            potx.write_bytes(blob)
            loaded = load_presentation(potx)
            texts = " ".join(obj.props.get("text", "") for slide in loaded.slides for obj in slide.objects)
            self.assertIn("Titel aus Paket", texts)
            ppsx = root / "schau.ppsx"
            ppsx.write_bytes(blob)
            again = load_presentation(ppsx)
            texts = " ".join(obj.props.get("text", "") for slide in again.slides for obj in slide.objects)
            self.assertIn("Titel aus Paket", texts)

    def test_pdrx_zip_imports(self):
        from officedisplay.io.open_file import load_presentation

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            target = root / "beispiel.pdrx"
            target.write_bytes(self._pptx_bytes(root))
            self.assertTrue(zipfile.is_zipfile(target))
            loaded = load_presentation(target)
            texts = " ".join(obj.props.get("text", "") for slide in loaded.slides for obj in slide.objects)
            self.assertIn("Titel aus Paket", texts)

    def test_png_becomes_image_object(self):
        from officedisplay.io.incoming import place_on_slide
        from officedisplay.model.presentation import Slide, new_id

        with tempfile.TemporaryDirectory() as folder:
            image = Path(folder) / "klein.png"
            Image.new("RGB", (12, 8), (20, 40, 60)).save(image)
            slide = Slide(id=new_id(), name="Folie")
            kind = place_on_slide(slide, str(image), 0.1, 0.1)
            self.assertEqual(kind, "image")
            self.assertEqual(slide.objects[-1].kind, "image")
            self.assertTrue(Path(slide.objects[-1].props["path"]).is_file())

    def test_mp4_and_avi_are_video_objects(self):
        from officedisplay.io.incoming import classify, place_on_slide
        from officedisplay.model.presentation import Slide, new_id

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name in ("clip.mp4", "film.avi"):
                path = root / name
                path.write_bytes(b"not a real stream")
                self.assertEqual(classify(str(path)), "video")
                slide = Slide(id=new_id(), name="Folie")
                place_on_slide(slide, str(path), 0.2, 0.2)
                self.assertEqual(slide.objects[-1].kind, "video")
                self.assertTrue(slide.objects[-1].props["path"].endswith(path.suffix))

    def test_ppt_without_soffice_reads_text_and_png(self):
        import struct

        from PIL import Image

        from officedisplay.io import open_file
        from officedisplay.io.open_file import FormatReadError, load_presentation

        def record(version: int, kind: int, body: bytes) -> bytes:
            return struct.pack("<HHI", version & 0xF, kind, len(body)) + body

        def compound(stream: bytes) -> bytes:
            sector = 512
            count = max(1, (len(stream) + sector - 1) // sector)
            fat = [0xFFFFFFFF] * 128
            fat[0] = 0xFFFFFFFD
            fat[1] = 0xFFFFFFFE
            for index in range(count):
                sid = 2 + index
                fat[sid] = 0xFFFFFFFE if index == count - 1 else sid + 1
            header = bytearray(512)
            header[0:8] = bytes.fromhex("D0CF11E0A1B11AE1")
            struct.pack_into("<HHHHH", header, 0x18, 0x003E, 0x0003, 0xFFFE, 9, 6)
            struct.pack_into("<I", header, 0x2C, 1)
            struct.pack_into("<I", header, 0x30, 1)
            struct.pack_into("<I", header, 0x38, 0)
            struct.pack_into("<I", header, 0x3C, 0xFFFFFFFE)
            struct.pack_into("<I", header, 0x44, 0xFFFFFFFE)
            for index in range(109):
                struct.pack_into("<I", header, 0x4C + index * 4, 0xFFFFFFFF)
            struct.pack_into("<I", header, 0x4C, 0)

            def entry(name: str, kind: int, start: int, size: int, child: int = 0xFFFFFFFF) -> bytes:
                raw = bytearray(128)
                encoded = name.encode("utf-16le")
                raw[0 : len(encoded)] = encoded
                struct.pack_into("<H", raw, 64, len(encoded) + 2)
                raw[66] = kind
                raw[67] = 1
                struct.pack_into("<III", raw, 68, 0xFFFFFFFF, 0xFFFFFFFF, child)
                struct.pack_into("<I", raw, 116, start)
                struct.pack_into("<Q", raw, 120, size)
                return bytes(raw)

            directory = bytearray(512)
            directory[0:128] = entry("Root Entry", 5, 0xFFFFFFFE, 0, child=1)
            directory[128:256] = entry("PowerPoint Document", 2, 2, len(stream))
            data = stream.ljust(count * sector, b"\x00")
            return bytes(header) + struct.pack("<128I", *fat) + bytes(directory) + data

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "punkt.png"
            Image.new("RGB", (8, 8), (10, 20, 30)).save(image)
            text = record(0, 4000, "Bericht".encode("utf-16le"))
            slide = record(0xF, 1006, text + image.read_bytes())
            path = root / "alt.ppt"
            path.write_bytes(compound(slide))
            broken = root / "kaputt.ppt"
            broken.write_bytes(b"\xd0\xcf\x11\xe0kein powerpoint")
            original = open_file.shutil.which
            open_file.shutil.which = lambda _name: None
            try:
                loaded = load_presentation(path)
                with self.assertRaises(FormatReadError) as caught:
                    load_presentation(broken)
            finally:
                open_file.shutil.which = original
            texts = [obj.props.get("text", "") for slide in loaded.slides for obj in slide.objects if obj.kind == "text"]
            self.assertTrue(any("Bericht" in text for text in texts))
            images = [obj for slide in loaded.slides for obj in slide.objects if obj.kind == "image"]
            self.assertTrue(images)
            self.assertTrue(Path(images[0].props["path"]).read_bytes().startswith(b"\x89PNG"))
            message = str(caught.exception)
            self.assertNotIn("soffice", message)
            self.assertNotIn("LibreOffice", message)
            self.assertNotIn("Traceback", message)
            self.assertTrue(message.strip())

    def test_menu_layout_and_notes(self):
        from PySide6.QtCore import QSettings

        from PySide6.QtWidgets import QWidget

        from officedisplay.model.presentation import Slide, new_id
        from officedisplay.ui.main_window import MainWindow

        settings = QSettings("Sellerbach", "OfficeDisplay")
        settings.setValue("menubar/visible", True)
        window = MainWindow()
        try:
            self.assertTrue(window.menuBar().isVisible())
            self.assertTrue(window.menu_button.isVisible())
            window.hide_menu_bar()
            self.app.processEvents()
            self.assertFalse(window.menuBar().isVisible())
            window.menu_button.click()
            self.app.processEvents()
            self.assertTrue(window.menuBar().isVisible())
            window.presentation.slides.append(Slide(id=new_id(), name="Zweite"))
            window.presentation.slides[0].objects.clear()
            window.presentation.slides[0].objects  # keep the list
            from officedisplay.model.presentation import SlideObject

            window.presentation.slides[1].objects = [
                SlideObject(id=new_id(), kind="text", x=0.1, y=0.1, w=0.4, h=0.2, props={"text": "Bleibt stehen"})
            ]
            window.canvas.set_index(0)
            window.layout_list.setCurrentRow(4)
            window.apply_current_layout()
            self.assertEqual(window.presentation.slides[0].objects, [])
            self.assertIn("Bleibt stehen", window.presentation.slides[1].objects[0].props["text"])
            self.assertEqual(len(window.presentation.slides), 2)
            window.notes_edit.setPlainText("Notiz vom Rand")
            window.commit_notes()
            self.assertEqual(window.current_slide().notes, "Notiz vom Rand")
            window.undo_change()
            self.assertEqual(window.current_slide().notes, "")
            self.assertIn("Folie", window.slide_pos.text())
            self.assertIn("%", window.zoom_label.text())
            titles = [widget.property("groupTitle") for widget in window.findChildren(QWidget)]
            for name in ("Folie", "Schrift", "Absatz", "Zeichnen", "Objekte"):
                self.assertIn(name, titles)
        finally:
            window.show_menu_bar()
            window.show_all_toolbars()
            window.close()


if __name__ == "__main__":
    unittest.main()
