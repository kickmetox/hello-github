import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image


os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _texts(presentation) -> str:
    return " ".join(
        str(obj.props.get("text") or "")
        for slide in presentation.slides
        for obj in slide.objects
        if obj.kind == "text"
    )


def _rich(folder: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    image = folder / "bild.png"
    Image.new("RGB", (32, 24), (20, 90, 200)).save(image)
    package = Presentation()
    title = package.slides.add_slide(package.slide_layouts[0])
    title.shapes.title.text = "CyberSec Titel"
    title.placeholders[1].text = "Untertitel im Saal"
    body = package.slides.add_slide(package.slide_layouts[1])
    body.shapes.title.text = "Punkte"
    frame = body.placeholders[1].text_frame
    frame.text = "Erster Punkt"
    frame.add_paragraph().text = "Zweiter Punkt"
    field = body.shapes.add_textbox(Inches(0.5), Inches(4.6), Inches(4), Inches(0.4))
    field.text_frame.paragraphs[0].text = "Freies Textfeld"
    body.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.2), Inches(1.2), Inches(2), Inches(1))
    table = body.shapes.add_table(2, 2, Inches(0.5), Inches(2.2), Inches(4.2), Inches(1.2)).table
    table.cell(0, 0).text = "SpalteA"
    table.cell(0, 1).text = "SpalteB"
    chart_slide = package.slides.add_slide(package.slide_layouts[5])
    chart_slide.shapes.title.text = "Diagramm Titel"
    data = CategoryChartData()
    data.categories = ["Nord", "Süd"]
    data.add_series("Werte", (4, 7))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(6), Inches(4), data
    )
    chart_slide.notes_slide.notes_text_frame.text = "Notiz Saal"
    picture = package.slides.add_slide(package.slide_layouts[6])
    picture.shapes.add_picture(str(image), Inches(0.6), Inches(0.6), Inches(3), Inches(2))
    package.slides.add_slide(package.slide_layouts[6])
    target = folder / "CyberSec ATCEUC.pptx"
    package.save(str(target))
    return target


def _replace_part(path: Path, predicate, payload: bytes) -> None:
    temporary = path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            blob = payload if predicate(info.filename) else source.read(info.filename)
            target.writestr(info, blob)
    temporary.replace(path)


class PptxOpenTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_filename_with_spaces_and_unicode(self):
        from officedisplay.io.pptx_io import import_pptx
        from officedisplay.ui import main_window as ui
        from officedisplay.ui.main_window import MainWindow

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            sample = _rich(root)
            loaded = import_pptx(sample)
            blob = _texts(loaded)
            self.assertIn("CyberSec Titel", blob)
            self.assertIn("Erster Punkt", blob)
            self.assertIn("Freies Textfeld", blob)
            self.assertIn("SpalteA", blob)
            self.assertIn("Diagramm Titel", blob)
            self.assertEqual(loaded.slides[2].notes, "Notiz Saal")
            self.assertGreaterEqual(len(loaded.slides), 5)
            unicode_path = root / "Fölie Übersicht.pptx"
            unicode_path.write_bytes(sample.read_bytes())
            again = import_pptx(unicode_path)
            self.assertIn("CyberSec Titel", _texts(again))
            messages = []
            original = ui.show_text
            ui.show_text = lambda _parent, _title, text: messages.append(text)
            guarded = ui.load_odisp

            def refuse_odisp(*_args, **_kwargs):
                raise AssertionError("pptx wurde als odisp gelesen")

            ui.load_odisp = refuse_odisp
            window = MainWindow()
            try:
                window.load_path(unicode_path)
                self.assertIn("CyberSec Titel", _texts(window.presentation))
                self.assertFalse(messages)
            finally:
                ui.show_text = original
                ui.load_odisp = guarded
                window.close()

    def test_table_or_chart_keeps_title(self):
        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            sample = _rich(root)
            broken = root / "chart kaputt.pptx"
            broken.write_bytes(sample.read_bytes())
            _replace_part(broken, lambda name: name.startswith("ppt/charts/") and name.endswith(".xml"), b"<not xml")
            loaded = import_pptx(broken)
            blob = _texts(loaded)
            self.assertIn("Diagramm Titel", blob)
            self.assertIn("SpalteA", blob)
            self.assertIn("CyberSec Titel", blob)

    def test_broken_shape_on_first_slide_keeps_second(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls
        from pptx.util import Inches

        from officedisplay.io import pptx_io
        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            package = Presentation()
            first = package.slides.add_slide(package.slide_layouts[5])
            first.shapes.title.text = "Titel Folie Eins"
            broken = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(2), Inches(1))
            broken.name = "Kaputt"
            second = package.slides.add_slide(package.slide_layouts[5])
            second.shapes.title.text = "Zweite Folie bleibt"
            target = root / "zwei folien.pptx"
            package.save(str(target))
            graphic = f'''
            <p:graphicFrame {nsdecls("a", "p", "r")}>
              <p:nvGraphicFramePr>
                <p:cNvPr id="80" name="Ohne Graphic"/>
                <p:cNvGraphicFramePr/>
                <p:nvPr/>
              </p:nvGraphicFramePr>
              <p:xfrm>
                <a:off x="914400" y="914400"/>
                <a:ext cx="2000000" cy="1000000"/>
              </p:xfrm>
            </p:graphicFrame>
            '''
            package = Presentation(str(target))
            package.slides[0].shapes._spTree.append(parse_xml(graphic))
            package.save(str(target))
            original = pptx_io._take_shape

            def boom(shape, slide, slide_w, slide_h, media, transforms=()):
                try:
                    name = shape.name
                except Exception:
                    name = ""
                if name == "Kaputt":
                    raise RuntimeError("kaputtes Shape")
                return original(shape, slide, slide_w, slide_h, media, transforms)

            pptx_io._take_shape = boom
            try:
                loaded = import_pptx(target)
            finally:
                pptx_io._take_shape = original
            blob = _texts(loaded)
            self.assertIn("Titel Folie Eins", blob)
            self.assertIn("Zweite Folie bleibt", blob)
            self.assertGreaterEqual(len(loaded.slides), 2)

            destroyed = root / "folie1 kaputt.pptx"
            destroyed.write_bytes(target.read_bytes())
            _replace_part(destroyed, lambda name: name.endswith("ppt/slides/slide1.xml"), b"<p:sld><kaputt")
            survived = import_pptx(destroyed)
            self.assertIn("Zweite Folie bleibt", _texts(survived))

    def test_zip_fallback_reads_title_when_presentation_fails(self):
        from officedisplay.io.pptx_io import PptxReadError, import_pptx
        from officedisplay.ui import main_window as ui
        from officedisplay.ui.main_window import MainWindow

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            target = root / "nur zip.pptx"
            slide = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:cSld><p:spTree>
                <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
                <p:grpSpPr/>
                <p:sp>
                  <p:nvSpPr>
                    <p:cNvPr id="2" name="Titel"/>
                    <p:cNvSpPr/>
                    <p:nvPr><p:ph type="title"/></p:nvPr>
                  </p:nvSpPr>
                  <p:txBody>
                    <a:bodyPr/>
                    <a:lstStyle/>
                    <a:p><a:r><a:t>Titel aus slide1.xml</a:t></a:r></a:p>
                  </p:txBody>
                </p:sp>
              </p:spTree></p:cSld>
            </p:sld>
            """
            presentation = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                            xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst>
              <p:sldSz cx="12192000" cy="6858000"/>
            </p:presentation>
            """
            rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
              <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
            </Relationships>
            """
            with zipfile.ZipFile(target, "w") as archive:
                archive.writestr("[Content_Types].xml", b"<Types>")
                archive.writestr("ppt/presentation.xml", presentation.encode("utf-8"))
                archive.writestr("ppt/_rels/presentation.xml.rels", rels.encode("utf-8"))
                archive.writestr("ppt/slides/slide1.xml", slide.encode("utf-8"))
            loaded = import_pptx(target)
            self.assertIn("Titel aus slide1.xml", _texts(loaded))

            missing = root / "beziehung fehlt.pptx"
            with zipfile.ZipFile(missing, "w") as archive:
                archive.writestr(
                    "[Content_Types].xml",
                    b'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"></Types>',
                )
                archive.writestr("ppt/presentation.xml", presentation.encode("utf-8"))
                archive.writestr(
                    "ppt/_rels/presentation.xml.rels",
                    b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>',
                )
            with self.assertRaises(PptxReadError) as caught:
                import_pptx(missing)
            self.assertIn("lässt sich nicht laden", str(caught.exception))
            self.assertIn("leer", str(caught.exception))

            secret = root / "geheim.pptx"
            secret.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1EncryptedPackage EncryptionInfo")
            with self.assertRaises(PptxReadError) as encrypted:
                import_pptx(secret)
            self.assertEqual(str(encrypted.exception), "Die Datei ist verschlüsselt und lässt sich hier nicht öffnen.")

            messages = []
            original = ui.show_text
            ui.show_text = lambda _parent, _title, text: messages.append(text)
            window = MainWindow()
            try:
                window.load_path(secret)
                self.assertEqual(messages[-1], "Die Datei ist verschlüsselt und lässt sich hier nicht öffnen.")
                plain = root / "keine pptx.pptx"
                plain.write_text("keine pptx", encoding="utf-8")
                window.load_path(plain)
                self.assertIn("lässt sich nicht laden", messages[-1])
                self.assertIn("kein ZIP", messages[-1])
            finally:
                ui.show_text = original
                window.close()


    def test_cover_photo_logo_and_title_keep_their_place(self):
        from pptx import Presentation
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt

        from officedisplay.io.pptx_io import import_pptx
        from officedisplay.ui.main_window import MainWindow
        from officedisplay.ui.render import render_thumb

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            wide = root / "himmel.png"
            logo = root / "logo.png"
            Image.new("RGB", (64, 36), (220, 20, 20)).save(wide)
            Image.new("RGB", (16, 16), (20, 40, 180)).save(logo)
            package = Presentation()
            package.slide_width = Emu(12192000)
            package.slide_height = Emu(6858000)
            slide = package.slides.add_slide(package.slide_layouts[6])
            width, height = int(package.slide_width), int(package.slide_height)
            slide.shapes.add_picture(str(wide), Emu(0), Emu(0), package.slide_width, package.slide_height)
            slide.shapes.add_picture(
                str(logo),
                Emu(int(width * 0.05)),
                Emu(int(height * 0.06)),
                Emu(int(width * 0.16)),
                Emu(int(height * 0.22)),
            )
            box = slide.shapes.add_textbox(
                Emu(int(width * 0.18)),
                Emu(int(height * 0.36)),
                Emu(int(width * 0.64)),
                Emu(int(height * 0.28)),
            )
            frame = box.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = "Impact"
            run.font.size = Pt(40)
            frame.add_paragraph().text = "Zeile"
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            path = root / "impact.pptx"
            package.save(path)
            loaded = import_pptx(path)
            page = loaded.slides[0]
            images = [obj for obj in page.objects if obj.kind == "image"]
            cover = page.background.mode == "image" and bool(page.background.image)
            self.assertTrue(
                cover
                or any(obj.x < 0.05 and obj.y < 0.05 and obj.w > 0.9 and obj.h > 0.9 for obj in images)
            )
            small = [obj for obj in images if obj.x < 0.3 and (obj.w < 0.9 or obj.h < 0.9)]
            self.assertTrue(small)
            text = next(obj for obj in page.objects if obj.kind == "text" and "Impact" in str(obj.props.get("text") or ""))
            self.assertGreaterEqual(int(text.props.get("font_size") or 0), 28)
            self.assertIn("Zeile", str(text.props.get("text") or ""))
            self.assertEqual(text.props.get("align"), "center")
            self.assertEqual(text.props.get("valign"), "middle")
            self.assertGreater(text.y, 0.2)
            thumb = render_thumb(loaded, page, 160)
            red = 0
            for y in range(0, thumb.height(), 3):
                for x in range(0, thumb.width(), 3):
                    color = thumb.pixelColor(x, y)
                    if color.red() > 180 and color.green() < 80:
                        red += 1
            self.assertGreater(red, 20)
            window = MainWindow()
            try:
                window.load_path(path)
                self.assertEqual(window.canvas.slide_index, 0)
                self.assertGreaterEqual(window.strip.count(), 1)
                current = window.presentation.slides[0]
                self.assertTrue(any("Impact" in str(obj.props.get("text") or "") for obj in current.objects))
            finally:
                window.close()

    def test_layout_background_image_reaches_the_slide(self):
        from pptx import Presentation
        from pptx.oxml import parse_xml
        from pptx.util import Emu

        from officedisplay.io.pptx_io import import_pptx
        from officedisplay.ui.render import render_thumb

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "layout.png"
            Image.new("RGB", (80, 45), (0, 0, 220)).save(image)
            package = Presentation()
            package.slide_width = Emu(12192000)
            package.slide_height = Emu(6858000)
            layout = package.slide_layouts[6]
            package.slides.add_slide(layout)
            _part, rid = layout.part.get_or_add_image_part(str(image))
            xml = (
                '<p:bg xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                f'<p:bgPr><a:blipFill><a:blip r:embed="{rid}"><a:alphaModFix amt="50000"/></a:blip>'
                "<a:stretch><a:fillRect/></a:stretch></a:blipFill></p:bgPr></p:bg>"
            )
            common = layout._element.cSld
            if common.bg is not None:
                common._remove_bg()
            common.spTree.addprevious(parse_xml(xml))
            path = root / "layout-hintergrund.pptx"
            package.save(path)
            loaded = import_pptx(path)
            page = loaded.slides[0]
            images = [obj for obj in page.objects if obj.kind == "image"]
            full = any(obj.x < 0.05 and obj.y < 0.05 and obj.w > 0.9 and obj.h > 0.9 for obj in images)
            self.assertTrue(bool(page.background.image) or full)
            self.assertAlmostEqual(page.background.opacity, 0.5, places=2)
            self.assertNotIn("Click to edit", _texts(loaded))
            thumb = render_thumb(loaded, page, 120)
            ink = 0
            for y in range(0, thumb.height(), 4):
                color = thumb.pixelColor(thumb.width() // 2, y)
                if color.blue() > color.red() + 20:
                    ink += 1
            self.assertGreater(ink, 3)


    def test_title_layout_stays_one_slide_and_not_a_text_wall(self):
        from lxml import etree
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Pt

        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            banner = root / "banner.png"
            left = root / "links.png"
            right = root / "rechts.png"
            Image.new("RGB", (80, 20), (20, 90, 180)).save(banner)
            Image.new("RGB", (16, 16), (0, 80, 160)).save(left)
            Image.new("RGB", (16, 16), (180, 40, 40)).save(right)
            package = Presentation()
            package.slide_width = Emu(12192000)
            package.slide_height = Emu(6858000)
            self.assertGreater(len(package.slide_layouts), 1)
            slide = package.slides.add_slide(package.slide_layouts[6])
            width, height = int(package.slide_width), int(package.slide_height)
            slide.shapes.add_picture(str(banner), Emu(0), Emu(0), package.slide_width, Emu(int(height * 0.16)))
            title = slide.shapes.add_textbox(
                Emu(int(width * 0.15)),
                Emu(int(height * 0.32)),
                Emu(int(width * 0.70)),
                Emu(int(height * 0.18)),
            )
            paragraph = title.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            p_props = paragraph._p.get_or_add_pPr()
            default_run = p_props.makeelement(qn("a:defRPr"))
            default_run.set("sz", "1200")
            p_props.append(default_run)
            run = paragraph.add_run()
            run.text = "Anforderungen"
            run.font.size = Pt(36)
            subtitle = slide.shapes.add_textbox(
                Emu(int(width * 0.18)),
                Emu(int(height * 0.52)),
                Emu(int(width * 0.64)),
                Emu(int(height * 0.12)),
            )
            sub = subtitle.text_frame.paragraphs[0]
            lead = sub.add_run()
            lead.text = "Nebenzeile "
            marked = sub.add_run()
            marked.text = "gelb"
            marked_props = marked._r.get_or_add_rPr()
            highlight = marked_props.makeelement(qn("a:highlight"))
            color = highlight.makeelement(qn("a:srgbClr"))
            color.set("val", "FFFF00")
            highlight.append(color)
            marked_props.append(highlight)
            plain = sub.add_run()
            plain.text = " fremd"
            slide.shapes.add_picture(
                str(left),
                Emu(int(width * 0.04)),
                Emu(int(height * 0.78)),
                Emu(int(width * 0.12)),
                Emu(int(height * 0.14)),
            )
            slide.shapes.add_picture(
                str(right),
                Emu(int(width * 0.82)),
                Emu(int(height * 0.78)),
                Emu(int(width * 0.12)),
                Emu(int(height * 0.14)),
            )
            slide.notes_slide.notes_text_frame.text = "Nur Notiz"
            background = (
                '<p:bg xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                '<p:bgPr><a:solidFill><a:srgbClr val="003399"/></a:solidFill></p:bgPr></p:bg>'
            )
            common = slide._element.cSld
            if common.bg is not None:
                common._remove_bg()
            common.spTree.addprevious(parse_xml(background))
            tree = slide.shapes._spTree
            keep = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
            moving = [child for child in list(tree) if child.tag not in keep]
            alternate = etree.Element(qn("ve:AlternateContent"))
            choice = etree.SubElement(alternate, qn("ve:Choice"))
            choice.set("Requires", "a")
            fallback = etree.SubElement(alternate, qn("ve:Fallback"))
            wall = f'''
            <p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                  xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:nvSpPr>
                <p:cNvPr id="90" name="Textwand"/>
                <p:cNvSpPr txBox="1"/>
                <p:nvPr/>
              </p:nvSpPr>
              <p:spPr>
                <a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>
              </p:spPr>
              <p:txBody>
                <a:bodyPr/><a:lstStyle/>
                <a:p><a:r><a:rPr sz="1200"/><a:t>Textwand</a:t></a:r></a:p>
              </p:txBody>
            </p:sp>
            '''
            fallback.append(parse_xml(wall))
            for child in moving:
                tree.remove(child)
                choice.append(child)
            tree.append(alternate)
            path = root / "anforderungen.pptx"
            package.save(path)
            ghost = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:cSld><p:spTree>
                <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
                <p:grpSpPr/>
                <p:sp><p:nvSpPr><p:cNvPr id="2" name="Geist"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
                  <p:txBody><a:bodyPr/><a:p><a:r><a:t>LayoutGeist</a:t></a:r></a:p></p:txBody>
                </p:sp>
              </p:spTree></p:cSld>
            </p:sld>
            """
            _add_zip_entry(path, "ppt/slides/slide99.xml", ghost.encode("utf-8"))
            loaded = import_pptx(path)
            self.assertEqual(len(loaded.slides), 1)
            page = loaded.slides[0]
            blob = _texts(loaded)
            self.assertIn("Anforderungen", blob)
            self.assertNotIn("Textwand", blob)
            self.assertNotIn("LayoutGeist", blob)
            self.assertNotIn("Nur Notiz", blob)
            self.assertEqual(page.notes, "Nur Notiz")
            self.assertEqual(page.background.color.upper(), "#003399")
            images = [obj for obj in page.objects if obj.kind == "image"]
            self.assertTrue(any(obj.y < 0.2 and obj.h < 0.4 and obj.w > 0.6 for obj in images))
            self.assertTrue(any(obj.x < 0.25 and obj.y > 0.65 and obj.w < 0.4 for obj in images))
            self.assertTrue(any(obj.x > 0.6 and obj.y > 0.65 and obj.w < 0.4 for obj in images))
            heading = next(obj for obj in page.objects if obj.kind == "text" and "Anforderungen" in str(obj.props.get("text") or ""))
            self.assertNotIn("Nebenzeile", str(heading.props.get("text") or ""))
            self.assertGreaterEqual(int(heading.props.get("font_size") or 0), 28)
            self.assertLess(heading.h, 0.6)
            self.assertEqual(heading.props.get("align"), "center")
            other = next(obj for obj in page.objects if obj.kind == "text" and "Nebenzeile" in str(obj.props.get("text") or ""))
            self.assertNotEqual(other.id, heading.id)
            runs = other.props.get("runs") or []
            self.assertTrue(any(item.get("text") == "gelb" and item.get("highlight") for item in runs))
            self.assertTrue(any("fremd" in str(item.get("text") or "") and not item.get("highlight") for item in runs))
            self.assertFalse(any(obj.kind == "shape" and str(obj.props.get("fill") or "").upper() == "#FFFF00" for obj in page.objects))

    def test_notes_layouts_and_extra_slide_files_are_not_slides(self):
        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "liste.pptx"
            slide = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:cSld><p:spTree>
                <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
                <p:grpSpPr/>
                <p:sp>
                  <p:nvSpPr><p:cNvPr id="2" name="Titel"/><p:cNvSpPr txBox="1"/><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
                  <p:spPr><a:xfrm><a:off x="400000" y="400000"/><a:ext cx="8000000" cy="1200000"/></a:xfrm></p:spPr>
                  <p:txBody><a:bodyPr/><a:lstStyle/>
                    <a:p><a:r><a:rPr sz="3600"/><a:t>Anforderungen</a:t></a:r></a:p>
                  </p:txBody>
                </p:sp>
              </p:spTree></p:cSld>
            </p:sld>
            """
            notes = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                     xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:cSld><p:spTree>
                <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
                <p:grpSpPr/>
                <p:sp><p:nvSpPr><p:cNvPr id="3" name="Notizen"/><p:cNvSpPr/><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>
                  <p:txBody><a:bodyPr/><a:p><a:r><a:t>NurNotiz</a:t></a:r></a:p></p:txBody>
                </p:sp>
              </p:spTree></p:cSld>
            </p:notes>
            """
            layout = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:sldLayout xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                         xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:cSld><p:spTree>
                <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
                <p:grpSpPr/>
                <p:sp><p:txBody><a:bodyPr/><a:p><a:r><a:t>Mastertext</a:t></a:r></a:p></p:txBody></p:sp>
              </p:spTree></p:cSld>
            </p:sldLayout>
            """
            extra = slide.replace("Anforderungen", "LayoutGeist")
            presentation = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <p:presentation xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                            xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst>
              <p:sldSz cx="12192000" cy="6858000"/>
            </p:presentation>
            """
            rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
            <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
              <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
              <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide" Target="notesSlides/notesSlide1.xml"/>
              <Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="slideLayouts/slideLayout1.xml"/>
            </Relationships>
            """
            with zipfile.ZipFile(target, "w") as archive:
                archive.writestr("[Content_Types].xml", b"<Types>")
                archive.writestr("ppt/presentation.xml", presentation.encode("utf-8"))
                archive.writestr("ppt/_rels/presentation.xml.rels", rels.encode("utf-8"))
                archive.writestr("ppt/slides/slide1.xml", slide.encode("utf-8"))
                archive.writestr("ppt/slides/slide99.xml", extra.encode("utf-8"))
                archive.writestr("ppt/notesSlides/notesSlide1.xml", notes.encode("utf-8"))
                archive.writestr("ppt/slideLayouts/slideLayout1.xml", layout.encode("utf-8"))
            loaded = import_pptx(target)
            self.assertEqual(len(loaded.slides), 1)
            blob = _texts(loaded)
            self.assertIn("Anforderungen", blob)
            self.assertNotIn("LayoutGeist", blob)
            self.assertNotIn("Mastertext", blob)
            self.assertNotIn("NurNotiz", blob)
            heading = next(obj for obj in loaded.slides[0].objects if "Anforderungen" in str(obj.props.get("text") or ""))
            self.assertGreaterEqual(int(heading.props.get("font_size") or 0), 28)
            self.assertLess(heading.w, 0.95)


def _add_zip_entry(path: Path, name: str, payload: bytes) -> None:
    temporary = path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(name, payload)
    temporary.replace(path)


if __name__ == "__main__":
    unittest.main()
