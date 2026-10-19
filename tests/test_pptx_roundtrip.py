import tempfile
import unittest
from pathlib import Path

from PIL import Image

from officedisplay.io.native_io import load_odisp, save_odisp
from officedisplay.io.odp_io import export_odp, import_odp
from officedisplay.io.pptx_io import export_pptx, import_pptx
from officedisplay.model.presentation import Presentation, Slide, SlideObject, new_id


class RoundtripTest(unittest.TestCase):
    def _deck(self, image_path: Path) -> Presentation:
        presentation = Presentation(title="Runde", width=1280, height=720)
        slide = Slide(id=new_id(), name="Folie 1", notes="Notiz bleibt")
        slide.objects.append(
            SlideObject(
                id=new_id(),
                kind="text",
                x=0.1,
                y=0.1,
                w=0.6,
                h=0.2,
                name="titel",
                props={
                    "text": "Hallo Andreas",
                    "font_size": 28,
                    "bold": True,
                    "color": "#112233",
                    "font_family": "Calibri",
                    "align": "left",
                },
            )
        )
        slide.objects.append(
            SlideObject(
                id=new_id(),
                kind="image",
                x=0.1,
                y=0.4,
                w=0.3,
                h=0.3,
                name="foto",
                props={"path": str(image_path)},
            )
        )
        presentation.slides = [slide]
        return presentation

    def test_pptx_text_and_image(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "bild.png"
            Image.new("RGB", (40, 30), (10, 80, 200)).save(image)
            target = root / "deck.pptx"
            export_pptx(self._deck(image), target)
            loaded = import_pptx(target)
            blob = " ".join(obj.props.get("text", "") for obj in loaded.slides[0].objects)
            self.assertIn("Hallo Andreas", blob)
            self.assertEqual(loaded.slides[0].notes, "Notiz bleibt")
            images = [obj for obj in loaded.slides[0].objects if obj.kind == "image"]
            self.assertTrue(images)
            self.assertTrue(Path(images[0].props["path"]).is_file())

    def test_odp_text(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "bild.png"
            Image.new("RGB", (8, 8), (0, 0, 0)).save(image)
            target = root / "deck.odp"
            export_odp(self._deck(image), target)
            loaded = import_odp(target)
            blob = " ".join(str(obj.props.get("text") or "") for obj in loaded.slides[0].objects)
            self.assertIn("Hallo Andreas", blob)

    def test_odisp_keeps_objects(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "bild.png"
            Image.new("RGB", (8, 8), (1, 2, 3)).save(image)
            target = root / "deck.odisp"
            original = self._deck(image)
            save_odisp(original, target)
            loaded = load_odisp(target)
            self.assertEqual(loaded.slides[0].objects[0].props["text"], "Hallo Andreas")
            self.assertEqual(loaded.title, "Runde")


class OpenPptxTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import os

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def _sample(self, folder: Path) -> Path:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        image = folder / "bild.png"
        Image.new("RGB", (48, 24), (20, 90, 200)).save(image)
        package = Presentation()
        title = package.slides.add_slide(package.slide_layouts[0])
        title.shapes.title.text = "Quartal Nord"
        title.placeholders[1].text = "Untertitel im Saal"
        body = package.slides.add_slide(package.slide_layouts[1])
        body.shapes.title.text = "Punkte"
        frame = body.placeholders[1].text_frame
        frame.text = "Erster Satz."
        frame.add_paragraph().text = "Zweiter Satz."
        body.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(2), Inches(2), Inches(1))
        broken = body.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(4), Inches(1), Inches(0.6))
        transform = broken._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
        if transform is not None:
            broken._element.spPr.remove(transform)
        picture = package.slides.add_slide(package.slide_layouts[6])
        box = picture.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(5), Inches(0.7))
        box.text_frame.paragraphs[0].text = "Bleibt stehen"
        picture.shapes.add_picture(str(image), Inches(0.4), Inches(1.4), Inches(3), Inches(1.6))
        target = folder / "real.pptx"
        package.save(target)
        return target

    def test_import_keeps_slides_text_shape_and_image(self):
        from officedisplay.io.pptx_io import PptxReadError, import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            loaded = import_pptx(self._sample(root))
            self.assertGreaterEqual(len(loaded.slides), 3)
            blob = " ".join(obj.props.get("text", "") for slide in loaded.slides for obj in slide.objects if obj.kind == "text")
            self.assertIn("Quartal Nord", blob)
            self.assertIn("Erster Satz.", blob)
            self.assertIn("Bleibt stehen", blob)
            shapes = [obj for slide in loaded.slides for obj in slide.objects if obj.kind == "shape"]
            self.assertTrue(any(obj.props.get("shape") == "roundrect" for obj in shapes))
            images = [obj for slide in loaded.slides for obj in slide.objects if obj.kind == "image"]
            self.assertTrue(images)
            self.assertTrue(Path(images[0].props["path"]).is_file())
            bad = root / "kaputt.pptx"
            bad.write_text("keine pptx", encoding="utf-8")
            with self.assertRaises(PptxReadError) as caught:
                import_pptx(bad)
            self.assertIn("lässt sich nicht laden", str(caught.exception))

    def test_import_from_readonly_directory(self):
        import os

        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            sample = self._sample(root)
            locked = root / "nurlesen"
            locked.mkdir()
            target = locked / "deck.pptx"
            target.write_bytes(sample.read_bytes())
            os.chmod(locked, 0o555)
            try:
                loaded = import_pptx(target)
            finally:
                os.chmod(locked, 0o755)
            self.assertGreaterEqual(len(loaded.slides), 3)
            images = [obj for slide in loaded.slides for obj in slide.objects if obj.kind == "image"]
            self.assertTrue(images)
            self.assertTrue(Path(images[0].props["path"]).is_file())

    def test_open_path_drop_and_error_keeps_deck(self):
        from officedisplay.ui import main_window as ui
        from officedisplay.ui.main_window import OPEN_FILTER, PPTX_FILTER, MainWindow

        self.assertIn("*.pptx", OPEN_FILTER)
        self.assertIn("*.pptx", PPTX_FILTER)
        messages = []
        original = ui.show_text
        ui.show_text = lambda _parent, _title, text: messages.append(text)
        window = MainWindow()
        try:
            with tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                sample = self._sample(root)
                window.load_path(sample)
                blob = " ".join(obj.props.get("text", "") for slide in window.presentation.slides for obj in slide.objects)
                self.assertIn("Quartal Nord", blob)
                kept = len(window.presentation.slides)
                before_drop = kept
                window._drop_files([str(sample)], 0.1, 0.1)
                self.assertGreater(len(window.presentation.slides), before_drop)
                bad = root / "kaputt.pptx"
                bad.write_text("keine pptx", encoding="utf-8")
                window.load_path(bad)
                self.assertEqual(len(window.presentation.slides), kept + 3)
                self.assertTrue(messages)
                self.assertIn("lässt sich nicht laden", messages[-1])
        finally:
            ui.show_text = original
            window.close()

    def test_title_placeholder_reaches_editor(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        from officedisplay.io.pptx_io import import_pptx
        from officedisplay.ui.main_window import MainWindow
        from officedisplay.ui.render import render_thumb

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            package = Presentation()
            slide = package.slides.add_slide(package.slide_layouts[1])
            slide.shapes.title.text = "Titelplatzhalter Nord"
            frame = slide.placeholders[1].text_frame
            frame.text = "Erster Punkt"
            frame.add_paragraph().text = "Zweiter Punkt"
            slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.2), Inches(4.5), Inches(2.2), Inches(1.2))
            path = root / "titel.pptx"
            package.save(path)
            loaded = import_pptx(path)
            texts = [obj.props.get("text", "") for slide in loaded.slides for obj in slide.objects if obj.kind == "text"]
            self.assertTrue(any("Titelplatzhalter Nord" in text for text in texts))
            self.assertTrue(any(obj.kind == "shape" for slide in loaded.slides for obj in slide.objects))

            bare = Presentation()
            drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

            def strip(shapes):
                for shape in shapes:
                    props = getattr(shape._element, "spPr", None)
                    if props is None:
                        continue
                    transform = props.find(f"{drawing}xfrm")
                    if transform is not None:
                        props.remove(transform)

            strip(bare.slide_master.shapes)
            for layout in bare.slide_layouts:
                strip(layout.shapes)
            empty = bare.slides.add_slide(bare.slide_layouts[1])
            empty.shapes.title.text = "Titel ohne Mass"
            empty.placeholders[1].text = "Punkt ohne Mass"
            bare_path = root / "ohne_mass.pptx"
            bare.save(bare_path)
            inherited = import_pptx(bare_path)
            blob = " ".join(obj.props.get("text", "") for slide in inherited.slides for obj in slide.objects)
            self.assertIn("Titel ohne Mass", blob)
            self.assertIn("Punkt ohne Mass", blob)
            self.assertTrue(all(0 <= obj.x <= 1 and obj.w >= 0.1 for slide in inherited.slides for obj in slide.objects))

            window = MainWindow()
            try:
                window.load_path(path)
                self.assertEqual(window.canvas.slide_index, 0)
                self.assertGreaterEqual(len(window.presentation.slides), 1)
                self.assertGreaterEqual(window.strip.count(), 1)
                current = window.presentation.slides[window.canvas.slide_index]
                self.assertTrue(any("Titelplatzhalter Nord" in str(obj.props.get("text") or "") for obj in current.objects))
                thumb = render_thumb(window.presentation, current, 320)
                ink = 0
                for y in range(0, thumb.height(), 2):
                    for x in range(0, thumb.width(), 2):
                        color = thumb.pixelColor(x, y)
                        if color.red() < 240 or color.green() < 240 or color.blue() < 240:
                            ink += 1
                self.assertGreater(ink, 30)
            finally:
                window.close()

    def test_group_children_use_slide_space(self):
        from pptx import Presentation
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls

        from officedisplay.io.pptx_io import import_pptx

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            package = Presentation()
            package.slides.add_slide(package.slide_layouts[6])
            target = root / "gruppe.pptx"
            package.save(target)
            package = Presentation(str(target))
            xml = f'''
            <p:grpSp {nsdecls("a", "p", "r")}>
              <p:nvGrpSpPr>
                <p:cNvPr id="20" name="Gruppe"/>
                <p:cNvGrpSpPr/>
                <p:nvPr/>
              </p:nvGrpSpPr>
              <p:grpSpPr>
                <a:xfrm>
                  <a:off x="914400" y="685800"/>
                  <a:ext cx="6400800" cy="4114800"/>
                  <a:chOff x="0" y="0"/>
                  <a:chExt cx="100000" cy="100000"/>
                </a:xfrm>
              </p:grpSpPr>
              <p:sp>
                <p:nvSpPr>
                  <p:cNvPr id="21" name="Titel in Gruppe"/>
                  <p:cNvSpPr txBox="1"/>
                  <p:nvPr/>
                </p:nvSpPr>
                <p:spPr>
                  <a:xfrm>
                    <a:off x="5000" y="8000"/>
                    <a:ext cx="80000" cy="20000"/>
                  </a:xfrm>
                  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                </p:spPr>
                <p:txBody>
                  <a:bodyPr/>
                  <a:lstStyle/>
                  <a:p><a:r><a:t>Gruppentitel Sichtbar</a:t></a:r></a:p>
                </p:txBody>
              </p:sp>
            </p:grpSp>
            '''
            package.slides[0].shapes._spTree.append(parse_xml(xml))
            package.save(target)
            loaded = import_pptx(target)
            texts = [obj for slide in loaded.slides for obj in slide.objects if obj.kind == "text"]
            self.assertTrue(texts)
            title = next(obj for obj in texts if "Gruppentitel" in str(obj.props.get("text") or ""))
            self.assertGreater(title.w, 0.2)
            self.assertGreater(title.h, 0.05)
            self.assertLess(title.x, 1)
            self.assertGreater(title.x, 0.02)


class ToolbarStartTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import os

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_toolbars_visible_and_reset_once(self):
        from PySide6.QtCore import QSettings

        from officedisplay.ui.main_window import TOOLBAR_KEYS, MainWindow
        from officedisplay.ui.theme import stylesheet

        self.assertIn("min-height: 40px", stylesheet("farbe"))
        self.assertNotIn("height: 0", stylesheet("farbe"))
        settings = QSettings("Sellerbach", "OfficeDisplay")
        settings.setValue("toolbar/force_visible_v1", False)
        for key in TOOLBAR_KEYS:
            settings.setValue(f"toolbar/{key}", False)
        window = MainWindow()
        try:
            for key in TOOLBAR_KEYS:
                self.assertTrue(window._toolbars[key].isVisible(), key)
                self.assertGreaterEqual(window._toolbars[key].minimumHeight(), 40)
            window._toolbar_actions["grafik"].setChecked(False)
            self.assertFalse(window._toolbars["grafik"].isVisible())
            again = MainWindow()
            try:
                self.assertFalse(again._toolbars["grafik"].isVisible())
                self.assertTrue(again._toolbars["text"].isVisible())
                self.assertTrue(again._toolbars["auswahl"].isVisible())
                again.show_all_toolbars()
                for key in TOOLBAR_KEYS:
                    self.assertTrue(again._toolbars[key].isVisible(), key)
            finally:
                again.close()
        finally:
            window.show_all_toolbars()
            window.close()


if __name__ == "__main__":
    unittest.main()
