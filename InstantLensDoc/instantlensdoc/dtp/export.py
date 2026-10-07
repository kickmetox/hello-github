"""DTP-Export: PDF via QPdfWriter/QPainter (kein ild_pdf-Writer), plus DOCX/ODT."""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from PySide6.QtCore import QMarginsF, QRectF, QSizeF, Qt
from PySide6.QtGui import (
    QColor,
    QFont,
    QImage,
    QPageLayout,
    QPageSize,
    QPainter,
    QPainterPath,
    QPdfWriter,
    QPen,
    QTextDocument,
)


def _qfont_for_frame(doc, fr) -> QFont:
    style = doc.styles.get(fr.style_id) or doc.styles.get("body")
    fam = fr.font_family or (style.font_family if style else "serif") or "serif"
    size = fr.font_size or (style.font_size if style else 11.0) or 11.0
    font = QFont(fam)
    font.setPointSizeF(float(size))
    weight = fr.font_weight or (style.weight if style else 400)
    try:
        font.setWeight(QFont.Weight(int(weight)) if int(weight) in (100, 200, 300, 400, 500, 600, 700, 800, 900) else QFont.Weight.Normal)
    except Exception:
        if int(weight) >= 600:
            font.setBold(True)
    stretch = fr.font_stretch or (style.stretch if style else 100)
    if stretch and stretch != 100:
        font.setStretch(int(stretch))
    axes = fr.font_axes or {}
    if "wght" in axes:
        w = float(axes["wght"])
        font.setBold(w >= 600)
        try:
            font.setWeight(QFont.Weight(int(max(100, min(900, w)))))
        except Exception:
            pass
    if "wdth" in axes:
        font.setStretch(int(max(50, min(200, axes["wdth"]))))
    if axes.get("slnt") or (style and style.italic):
        font.setItalic(True)
    if style and style.italic:
        font.setItalic(True)
    return font


def _shape_path(fr) -> QPainterPath:
    from .extrude import path_for_shape

    pts = getattr(fr, "weld_path", None) or []
    if pts and isinstance(pts[0], (list, tuple)) and len(pts[0]) >= 2:
        path = QPainterPath()
        path.moveTo(float(pts[0][0]), float(pts[0][1]))
        for pt in pts[1:]:
            if isinstance(pt, (list, tuple)) and len(pt) >= 2:
                path.lineTo(float(pt[0]), float(pt[1]))
        path.closeSubpath()
        if path.elementCount() > 2:
            return path
    return path_for_shape(fr.shape if fr.kind == "shape" else "rectangle", fr.width, fr.height)


def _text_path(doc, fr) -> QPainterPath:
    font = _qfont_for_frame(doc, fr)
    path = QPainterPath()
    path.addText(2.0, font.pointSizeF() + 2.0, font, fr.text or "")
    return path


def paint_page(painter: QPainter, doc, page: int, *, overlays: bool = False) -> None:
    """Eine Seite in Seitenkoordinaten (pt, Y nach unten) zeichnen."""
    g = doc.geometry
    painter.save()
    painter.fillRect(QRectF(0, 0, g.width_pt, g.height_pt), QColor("#ffffff"))

    master = None
    if doc.masters:
        mid = ""
        if 0 <= page < len(doc.page_master):
            mid = doc.page_master[page]
        master = next((m for m in doc.masters if m.id == mid), doc.masters[0])
    if master:
        painter.save()
        font = QFont("serif")
        font.setPointSizeF(9)
        painter.setFont(font)
        painter.setPen(QColor("#444444"))
        n, total = page + 1, max(1, doc.page_count)
        hr = QRectF(g.margin_left_pt, 8.0, g.width_pt - g.margin_left_pt - g.margin_right_pt, g.margin_top_pt - 10)
        painter.drawText(
            hr,
            Qt.AlignLeft | Qt.AlignVCenter,
            master.header_for(n, total, title=doc.title, number_system=getattr(doc, "number_system", "latin")),
        )
        fr = QRectF(
            g.margin_left_pt,
            g.height_pt - g.margin_bottom_pt + 4,
            g.width_pt - g.margin_left_pt - g.margin_right_pt,
            g.margin_bottom_pt - 8,
        )
        painter.drawText(
            fr,
            Qt.AlignRight | Qt.AlignVCenter,
            master.footer_for(n, total, title=doc.title, number_system=getattr(doc, "number_system", "latin")),
        )
        painter.restore()

    if overlays and doc.grid_visible:
        painter.save()
        pen = QPen(QColor(74, 144, 217, 70))
        pen.setWidthF(0.4)
        painter.setPen(pen)
        step = doc.grid_pt()
        x = 0.0
        while x <= g.width_pt + 0.1:
            painter.drawLine(x, 0.0, x, g.height_pt)
            x += step
        y = 0.0
        while y <= g.height_pt + 0.1:
            painter.drawLine(0.0, y, g.width_pt, y)
            y += step
        painter.restore()

    painted: set[str] = set()
    for fr in doc.sorted_frames(page):
        _paint_frame(painter, doc, fr)
        painted.add(fr.id)
    for fr in doc.frames:
        if fr.master and fr.id not in painted:
            _paint_frame(painter, doc, fr)
    painter.restore()


def _paint_text_on_path(painter: QPainter, doc, fr) -> None:
    from .text_path import place_text_on_path, polyline_for_kind

    font = _qfont_for_frame(doc, fr)
    size = float(font.pointSizeF() or 11.0)
    pts = polyline_for_kind(fr.path_kind, fr.width, fr.height)
    glyphs = place_text_on_path(fr.text or "", pts, font_size=size)
    color = QColor("#111111")
    painter.setFont(font)
    painter.setPen(color)
    painter.setBrush(color)
    for g in glyphs:
        painter.save()
        painter.translate(g["x"], g["y"])
        painter.rotate(g["angle"])
        if fr.as_outlines:
            gp = QPainterPath()
            gp.addText(0.0, 0.0, font, g["char"])
            painter.fillPath(gp, color)
        else:
            painter.drawText(0, 0, g["char"])
        painter.restore()


def paint_frame_local(painter: QPainter, doc, fr) -> None:
    """Rahmeninhalt in lokalen Koordinaten (0,0 = Rahmenecke)."""
    from .effects import apply_clip, apply_opacity, fill_brush, paint_drop_shadow
    from .envelope import warp_painter_path
    from .extrude import extrude_path, path_for_shape

    painter.save()
    apply_opacity(painter, fr)
    apply_clip(painter, doc, fr)
    if fr.rotation:
        painter.rotate(fr.rotation)

    path = _shape_path(fr) if fr.kind in ("shape", "image") else QPainterPath()
    if fr.kind == "text":
        path = _text_path(doc, fr)
        painter.setClipRect(QRectF(0, 0, fr.width, fr.height), Qt.IntersectClip)

    if fr.envelope and fr.envelope.corners:
        path = warp_painter_path(path, width=fr.width, height=fr.height, corners=fr.envelope.corners)

    paint_drop_shadow(painter, path, fr)

    if fr.extrude and fr.kind in ("shape", "text"):
        fill = fr.fill or "#4A90D9"
        data = extrude_path(
            path if path.elementCount() else path_for_shape(fr.shape, fr.width, fr.height),
            depth=fr.extrude.depth,
            angle_deg=fr.extrude.angle_deg,
            shading=fr.extrude.shading,
            fill=fill,
        )
        for face in data["faces"]:
            painter.setBrush(face["color"])
            painter.setPen(Qt.NoPen)
            painter.drawPolygon(face["poly"])
        painter.setBrush(data["base"])
        painter.setPen(QPen(QColor(fr.stroke or "#222"), max(0.4, fr.stroke_width)))
        painter.drawPath(data["front"])
        painter.restore()
        return

    if fr.kind == "image" and fr.image_path:
        img = QImage(fr.image_path)
        if not img.isNull():
            painter.drawImage(QRectF(0, 0, fr.width, fr.height), img)
        else:
            painter.fillRect(QRectF(0, 0, fr.width, fr.height), QColor(fr.fill or "#e8e8e8"))
    elif fr.kind == "shape":
        painter.setBrush(fill_brush(fr, fr.width, fr.height))
        painter.setPen(QPen(QColor(fr.stroke or "#1A5276"), max(0.4, fr.stroke_width)))
        painter.drawPath(path if path.elementCount() else path_for_shape(fr.shape, fr.width, fr.height))
    elif fr.kind == "ink" and fr.ink_points:
        pen = QPen(QColor(fr.stroke or "#111"), max(0.6, fr.stroke_width))
        pen.setCapStyle(Qt.RoundCap)
        painter.setPen(pen)
        pts = fr.ink_points
        for i in range(1, len(pts)):
            w = max(0.5, fr.stroke_width * (0.4 + 1.4 * (pts[i][2] if len(pts[i]) > 2 else 0.5)))
            pen.setWidthF(w)
            painter.setPen(pen)
            painter.drawLine(
                pts[i - 1][0] - fr.x,
                pts[i - 1][1] - fr.y,
                pts[i][0] - fr.x,
                pts[i][1] - fr.y,
            )
    elif fr.kind == "text":
        if fr.path_kind:
            _paint_text_on_path(painter, doc, fr)
        elif fr.as_outlines or (fr.envelope and fr.envelope.corners):
            painter.setPen(Qt.NoPen)
            painter.setBrush(QColor("#111111"))
            painter.drawPath(path)
        else:
            html = str(getattr(fr, "rich_html", "") or "").strip()
            if html:
                td = QTextDocument()
                td.setHtml(html)
                td.setTextWidth(max(8.0, fr.width - 4))
                painter.save()
                painter.translate(2, 2)
                td.drawContents(painter, QRectF(0, 0, fr.width - 4, fr.height - 4))
                painter.restore()
            else:
                painter.setPen(QColor("#111111"))
                painter.setFont(_qfont_for_frame(doc, fr))
                painter.drawText(
                    QRectF(2, 2, fr.width - 4, fr.height - 4),
                    Qt.TextWordWrap | Qt.AlignLeft | Qt.AlignTop,
                    fr.text or "",
                )
    elif fr.kind == "stamp":
        text_only = bool(getattr(fr, "stamp_text_only", False))
        try:
            sw = float(fr.stroke_width or 0.0)
        except (TypeError, ValueError):
            sw = 0.0
        outline = bool(getattr(fr, "stamp_outline", False)) and not text_only
        if getattr(fr, "shadow", False):
            painter.fillRect(QRectF(3, 3, fr.width, fr.height), QColor(0, 0, 0, 90))
        if fr.fill and not text_only:
            painter.fillRect(QRectF(0, 0, fr.width, fr.height), QColor(fr.fill))
        if (sw >= 0.5 or outline) and not text_only:
            painter.setPen(QPen(QColor(fr.stroke or "#1E8449"), max(0.6, sw if sw >= 0.5 else 2.0)))
            painter.setBrush(Qt.NoBrush)
            painter.drawRect(QRectF(0.5, 0.5, max(1.0, fr.width - 1), max(1.0, fr.height - 1)))
        painter.setPen(QColor(fr.stroke or "#1E8449"))
        painter.setFont(_qfont_for_frame(doc, fr))
        painter.drawText(
            QRectF(4, 4, max(4.0, fr.width - 8), max(4.0, fr.height - 8)),
            Qt.TextWordWrap | Qt.AlignCenter,
            fr.text or "STEMPEL",
        )
    painter.restore()


def _blend_modes_qt() -> dict[str, object]:
    try:
        C = QPainter.CompositionMode
        return {
            "multiply": C.CompositionMode_Multiply,
            "screen": C.CompositionMode_Screen,
            "overlay": C.CompositionMode_Overlay,
            "darken": C.CompositionMode_Darken,
            "lighten": C.CompositionMode_Lighten,
            "color_dodge": C.CompositionMode_ColorDodge,
            "color_burn": C.CompositionMode_ColorBurn,
        }
    except Exception:
        return {}


_BLEND_QT = _blend_modes_qt()


def _paint_frame(painter: QPainter, doc, fr) -> None:
    painter.save()
    ly = next((l for l in getattr(doc, "layers", []) if l.id == fr.layer_id), None)
    if ly is not None:
        try:
            painter.setOpacity(painter.opacity() * max(0.05, min(1.0, float(ly.opacity))))
        except Exception:
            pass
        mode = _BLEND_QT.get(str(getattr(ly, "blend_mode", "") or "").lower())
        if mode is not None:
            painter.setCompositionMode(mode)
    painter.translate(fr.x, fr.y)
    paint_frame_local(painter, doc, fr)
    painter.restore()


def export_pdf(doc, path: str | Path, *, overlays: bool = False) -> Path:
    """Vektor-PDF über QPdfWriter — nicht über ild_pdf-Exportwriter."""
    from PySide6.QtWidgets import QApplication
    import os

    if QApplication.instance() is None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        QApplication(["ild-dtp-export"])
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    writer = QPdfWriter(str(dest))
    g = doc.geometry
    writer.setTitle(doc.title or "InstantLens Doc DTP")
    writer.setPageSize(QPageSize(QSizeF(g.width_pt, g.height_pt), QPageSize.Unit.Point, g.name or "dtp"))
    writer.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Point)
    writer.setResolution(72)
    painter = QPainter()
    if not painter.begin(writer):
        raise RuntimeError("QPdfWriter: QPainter.begin fehlgeschlagen")
    try:
        for i in range(max(1, doc.page_count)):
            if i:
                writer.newPage()
            paint_page(painter, doc, i, overlays=overlays)
    finally:
        painter.end()
    return dest


def export_docx(doc, path: str | Path) -> Path:
    """Best-effort DOCX: Seitenformat + Rahmen als Absätze (keine absolute Position)."""
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, Twips
    except Exception as exc:
        raise RuntimeError(f"python-docx fehlt: {exc}") from exc

    d = Document()
    section = d.sections[0]
    # EMU: 1 pt = 12700 EMU; python-docx uses Emu via Pt/Twips
    section.page_width = Twips(int(doc.geometry.width_pt * 20))
    section.page_height = Twips(int(doc.geometry.height_pt * 20))
    section.top_margin = Twips(int(doc.geometry.margin_top_pt * 20))
    section.bottom_margin = Twips(int(doc.geometry.margin_bottom_pt * 20))
    section.left_margin = Twips(int(doc.geometry.margin_left_pt * 20))
    section.right_margin = Twips(int(doc.geometry.margin_right_pt * 20))
    d.core_properties.title = doc.title
    for page in range(max(1, doc.page_count)):
        if page:
            d.add_page_break()
        master = doc.masters[0] if doc.masters else None
        if master:
            p = d.add_paragraph(master.header_for(page + 1, doc.page_count, title=doc.title))
            p.runs[0].italic = True if p.runs else None
        for fr in doc.sorted_frames(page):
            if fr.kind == "text" and (fr.text or "").strip():
                p = d.add_paragraph(fr.text.replace("\n", " "))
                style = doc.styles.get(fr.style_id)
                if p.runs:
                    run = p.runs[0]
                    size = fr.font_size or (style.font_size if style else 11)
                    run.font.size = Pt(float(size))
                    if style and style.weight >= 600:
                        run.bold = True
                    if fr.font_family:
                        run.font.name = fr.font_family
                align = (style.alignment if style else "left") or "left"
                mapping = {
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                }
                p.alignment = mapping.get(align, WD_ALIGN_PARAGRAPH.LEFT)
            elif fr.kind in ("shape", "image"):
                d.add_paragraph(f"[{fr.kind}/{fr.shape} {fr.width:.0f}×{fr.height:.0f} pt]")
        if master:
            d.add_paragraph(master.footer_for(page + 1, doc.page_count, title=doc.title))
    d.save(str(dest))
    return dest


def export_odt(doc, path: str | Path) -> Path:
    """Minimales ODT (ZIP+XML), unabhängig von odfpy."""
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    paras: list[str] = []
    paras.append(f"<text:h text:style-name=\"Heading_20_1\" text:outline-level=\"1\">{escape(doc.title or 'DTP')}</text:h>")
    for page in range(max(1, doc.page_count)):
        paras.append(f"<text:p text:style-name=\"Standard\">Seite {page + 1}</text:p>")
        for fr in doc.sorted_frames(page):
            if fr.kind == "text" and (fr.text or "").strip():
                paras.append(
                    f"<text:p text:style-name=\"Standard\">{escape(fr.text)}</text:p>"
                )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        "<office:body><office:text>"
        + "".join(paras)
        + "</office:text></office:body></office:document-content>"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'office:version="1.2"><office:styles/></office:document-styles>'
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
        '<manifest:file-entry manifest:media-type="application/vnd.oasis.opendocument.text" manifest:full-path="/"/>'
        '<manifest:file-entry manifest:media-type="text/xml" manifest:full-path="content.xml"/>'
        '<manifest:file-entry manifest:media-type="text/xml" manifest:full-path="styles.xml"/>'
        "</manifest:manifest>"
    )
    with zipfile.ZipFile(dest, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text", compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content)
        zf.writestr("styles.xml", styles)
        zf.writestr("META-INF/manifest.xml", manifest)
    return dest


def export_pdfx3(doc, path: str | Path, *, min_dpi: float = 150.0) -> Path:
    """PDF/X-3: DTP-PDF + OutputIntent / GTS_PDFXVersion (pikepdf)."""
    from instantlensdoc.dtp.preflight import run_dtp_preflight

    dest = Path(path)
    report = run_dtp_preflight(doc, min_dpi=min_dpi)
    if not report.ok:
        raise ValueError(report.to_text())
    export_pdf(doc, dest)
    import pikepdf
    from pikepdf import Dictionary, Name, String

    with pikepdf.open(dest, allow_overwriting_input=True) as pdf:
        pdf.docinfo["/GTS_PDFXVersion"] = String("PDF/X-3:2003")
        pdf.docinfo["/Title"] = String(getattr(doc, "title", "") or "InstantLens Doc")
        pdf.Root.OutputIntents = [
            Dictionary(
                {
                    "/Type": Name.OutputIntent,
                    "/S": Name.GTS_PDFX,
                    "/OutputConditionIdentifier": String("Custom"),
                    "/Info": String("InstantLens Doc DTP PDF/X-3"),
                    "/RegistryName": String("http://www.color.org"),
                }
            )
        ]
        pdf.save(dest, min_version="1.4")
    return dest


def export_pdf_versions(doc, path: str | Path, versions: tuple[str, ...] = ("1.4", "1.7")) -> dict[str, Path]:
    """Multi-Version-PDF: stem-14.pdf und stem-17.pdf."""
    import pikepdf

    base = Path(path)
    src = export_pdf(doc, base)
    raw = src.read_bytes()
    out: dict[str, Path] = {}
    for ver in versions:
        dest = base.with_name(f"{base.stem}-{ver.replace('.', '')}{base.suffix}")
        dest.write_bytes(raw)
        with pikepdf.open(dest, allow_overwriting_input=True) as pdf:
            pdf.save(dest, min_version=ver)
        out[ver] = dest
    return out


def export_interactive_pdf(doc, path: str | Path) -> Path:
    """PDF mit AcroForm-Widgets (Text, Combo, Präsentation)."""
    from instantlensdoc.dtp.interactive import attach_widgets_to_pdf

    dest = export_pdf(doc, path)
    widgets = list(getattr(doc, "widgets", []) or [])
    if widgets:
        attach_widgets_to_pdf(dest, widgets, list(doc.frames))
    return dest
