"""Textimport in DTP-Rahmen: TXT/MD/HTML/DOCX/OCR-Sidecar."""

from __future__ import annotations

import re
from pathlib import Path


def read_import_text(path: str | Path) -> str:
    """Dateiinhalt als Fließtext. DOCX über python-docx; OCR-Sidecar als UTF-8."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(str(p))
    name = p.name.lower()
    suf = p.suffix.lower()
    if suf == ".docx":
        return _read_docx(p)
    raw = p.read_text(encoding="utf-8", errors="replace")
    if suf in (".html", ".htm") or name.endswith(".hocr") or ".ildocr.hocr" in name:
        return _strip_markup(raw)
    if suf == ".tsv" or name.endswith(".ildocr.tsv"):
        return _tsv_text(raw)
    return raw.replace("\r\n", "\n").replace("\r", "\n")


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx fehlt: {exc}") from exc
    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _strip_markup(html: str) -> str:
    text = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html)
    text = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</p>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.replace("\xa0", " ").strip()


def _tsv_text(raw: str) -> str:
    lines: list[str] = []
    for row in raw.splitlines():
        cols = row.split("\t")
        if not cols:
            continue
        # Tesseract TSV: text is last column when present
        word = cols[-1].strip() if cols else ""
        if word and word != "text":
            lines.append(word)
    return " ".join(lines)
