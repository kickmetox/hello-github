"""Seriendruck: Empfänger aus CSV/Excel → Briefe — 2.6.27 / Polish 2.6.27.

Platzhalter ``{{Feld}}`` / ``«Feld»`` im Template; lokal, ohne Cloud.

2.6.27: fehlende Felder melden, Delimiter, Preview, HTML/DOCX-Ausgabe,
optional eine Datei mit Trennseiten.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any, Sequence

_PLACEHOLDER_RE = re.compile(
    r"\{\{\s*([A-Za-z_][\w]*)\s*\}\}|«\s*([A-Za-z_][\w]*)\s*»"
)


def load_recipients_csv(
    path: str | Path,
    *,
    delimiter: str | None = None,
) -> list[dict[str, str]]:
    """CSV mit Kopfzeile laden → Liste von Empfänger-Dicts."""
    p = Path(path)
    text = p.read_text(encoding="utf-8-sig")
    sample = text[:4096]
    delim = delimiter
    if not delim:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = "," if sample.count(",") >= sample.count(";") else ";"
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    rows: list[dict[str, str]] = []
    for row in reader:
        if not isinstance(row, dict):
            continue
        cleaned = {
            str(k).strip(): ("" if v is None else str(v))
            for k, v in row.items()
            if k is not None and str(k).strip()
        }
        if any(v.strip() for v in cleaned.values()):
            rows.append(cleaned)
    return rows


def load_recipients_xlsx(path: str | Path) -> list[dict[str, str]]:
    """Erste Sheet als Empfängerliste (openpyxl falls vorhanden)."""
    p = Path(path)
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "openpyxl fehlt — CSV nutzen oder openpyxl installieren"
        ) from e
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    keys = [
        str(h).strip() if h is not None else f"col{i}"
        for i, h in enumerate(header)
    ]
    out: list[dict[str, str]] = []
    for raw in rows_iter:
        row = {
            keys[i]: ("" if (raw[i] is None) else str(raw[i]))
            for i in range(min(len(keys), len(raw or ())))
            if keys[i]
        }
        if any(v.strip() for v in row.values()):
            out.append(row)
    return out


def load_recipients(
    path: str | Path,
    *,
    delimiter: str | None = None,
) -> list[dict[str, str]]:
    p = Path(path)
    suf = p.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        return load_recipients_xlsx(p)
    return load_recipients_csv(p, delimiter=delimiter)


def find_placeholders(template: str) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for m in _PLACEHOLDER_RE.finditer(template or ""):
        name = m.group(1) or m.group(2) or ""
        if name and name not in seen:
            seen.add(name)
            names.append(name)
    return names


def missing_fields(
    template: str,
    recipients: Sequence[dict[str, Any]],
) -> dict[str, Any]:
    """Platzhalter prüfen: welche fehlen in Empfänger-Spalten / je Zeile."""
    ph = find_placeholders(template)
    if not recipients:
        return {
            "placeholders": ph,
            "columns": [],
            "missing_columns": list(ph),
            "rows_with_empty": [],
            "ok": not ph,
        }
    columns = sorted({str(k) for r in recipients for k in r.keys()})
    missing_cols = [p for p in ph if p not in columns]
    rows_empty: list[dict[str, Any]] = []
    for i, r in enumerate(recipients):
        empty = [p for p in ph if p in columns and not str(r.get(p, "")).strip()]
        if empty:
            rows_empty.append({"index": i, "empty": empty})
    return {
        "placeholders": ph,
        "columns": columns,
        "missing_columns": missing_cols,
        "rows_with_empty": rows_empty,
        "ok": not missing_cols,
    }


def merge_one(template: str, fields: dict[str, Any]) -> str:
    mapping = {str(k): ("" if v is None else str(v)) for k, v in (fields or {}).items()}

    def _repl(m: re.Match[str]) -> str:
        name = m.group(1) or m.group(2) or ""
        return mapping.get(name, m.group(0))

    return _PLACEHOLDER_RE.sub(_repl, template or "")


def mail_merge(
    template: str,
    recipients: Sequence[dict[str, Any]],
) -> list[str]:
    """Jeden Empfänger zu einem Brieftext mergen."""
    return [merge_one(template, r) for r in recipients]


def preview_merge(
    template: str,
    recipients: Sequence[dict[str, Any]],
    *,
    limit: int = 3,
) -> dict[str, Any]:
    """Vorschau der ersten N Briefe + Feldanalyse — 2.6.27."""
    analysis = missing_fields(template, recipients)
    letters = mail_merge(template, list(recipients)[: max(0, int(limit))])
    return {
        **analysis,
        "preview_count": len(letters),
        "total_recipients": len(recipients),
        "preview": letters,
    }


def _write_letter(path: Path, text: str, fmt: str) -> Path:
    fmt_l = (fmt or "txt").lower()
    if fmt_l == "html":
        body = (
            "<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<title>{path.stem}</title></head><body><pre>"
            + _html_escape(text)
            + "</pre></body></html>"
        )
        path = path.with_suffix(".html")
        path.write_text(body, encoding="utf-8")
        return path
    if fmt_l == "docx":
        path = path.with_suffix(".docx")
        try:
            from instantlensdoc.core.export import export_docx

            export_docx(text, path, title=path.stem)
        except Exception:
            # Fallback: plain text with .docx name avoided — write txt
            path = path.with_suffix(".txt")
            path.write_text(text, encoding="utf-8")
        return path
    path = path.with_suffix(".txt")
    path.write_text(text, encoding="utf-8")
    return path


def _html_escape(s: str) -> str:
    return (
        (s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def mail_merge_to_dir(
    template: str,
    recipients: Sequence[dict[str, Any]],
    out_dir: str | Path,
    *,
    stem: str = "letter",
    fmt: str = "txt",
    combined: bool = False,
    separator: str = "\n\n---\n\n",
) -> list[Path]:
    """Briefe als ``letter_001.txt`` … schreiben; optional eine kombinierte Datei."""
    dest = Path(out_dir)
    dest.mkdir(parents=True, exist_ok=True)
    letters = mail_merge(template, recipients)
    written: list[Path] = []
    if combined:
        joined = separator.join(letters)
        path = dest / f"{stem}_all"
        written.append(_write_letter(path, joined, fmt))
        return written
    for i, letter in enumerate(letters, start=1):
        path = dest / f"{stem}_{i:03d}"
        written.append(_write_letter(path, letter, fmt))
    return written


def mail_merge_from_files(
    template_path: str | Path,
    recipients_path: str | Path,
    out_dir: str | Path,
    *,
    stem: str = "letter",
    fmt: str = "txt",
    combined: bool = False,
    delimiter: str | None = None,
    strict: bool = False,
) -> dict[str, Any]:
    tpl_path = Path(template_path)
    template = tpl_path.read_text(encoding="utf-8")
    recipients = load_recipients(recipients_path, delimiter=delimiter)
    analysis = missing_fields(template, recipients)
    if strict and analysis["missing_columns"]:
        raise ValueError(
            "Fehlende Spalten für Platzhalter: "
            + ", ".join(analysis["missing_columns"])
        )
    paths = mail_merge_to_dir(
        template,
        recipients,
        out_dir,
        stem=stem,
        fmt=fmt,
        combined=combined,
    )
    return {
        "template": str(tpl_path),
        "recipients": str(Path(recipients_path)),
        "count": len(paths) if combined else len(recipients),
        "files": len(paths),
        "placeholders": find_placeholders(template),
        "missing_columns": analysis["missing_columns"],
        "rows_with_empty": len(analysis["rows_with_empty"]),
        "fmt": fmt,
        "combined": combined,
        "output": [str(p) for p in paths],
        "out_dir": str(Path(out_dir)),
    }
