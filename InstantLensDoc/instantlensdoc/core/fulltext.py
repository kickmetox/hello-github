"""Volltextsuche über mehrere geöffnete Dokumente / PDFs — 2.0.2: CSV Doc/Seite/Snippet/Match·BOM."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, List, Sequence

_TEXT_SUFFIXES = {".txt", ".md", ".html", ".htm", ".json", ".ildocr.txt"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}


class SearchPatternError(ValueError):
    """Ungültiger Suchausdruck (z. B. fehlerhafte Regex) — Multi-Doc-Suche."""


@dataclass
class SearchHit:
    path: str
    page: int | None  # 0-based PDF page; None for text files
    line: int | None
    snippet: str
    kind: str  # text | pdf | sidecar


def _is_word_char(ch: str) -> bool:
    if not ch:
        return False
    return ch.isalnum() or ch == "_"


def _compile_pattern(
    query: str,
    *,
    case_sensitive: bool = False,
    regex: bool = False,
) -> re.Pattern[str] | None:
    """Regex-Pattern oder None (Literal-Suche). Wirft SearchPatternError."""
    q = (query or "").strip()
    if not q or not regex:
        return None
    flags = 0 if case_sensitive else re.IGNORECASE
    try:
        return re.compile(q, flags)
    except re.error as e:
        raise SearchPatternError(str(e) or "ungültiger regulärer Ausdruck") from e


def _line_has_match(
    line: str,
    query: str,
    *,
    case_sensitive: bool = False,
    whole_word: bool = False,
    pattern: re.Pattern[str] | None = None,
) -> bool:
    """True wenn Zeile den Suchbegriff (Literal/Regex, Case, Whole-word) trifft."""
    if not line:
        return False
    q = (query or "").strip()
    if not q and pattern is None:
        return False
    if pattern is not None:
        for m in pattern.finditer(line):
            if not whole_word:
                return True
            pos, end = m.start(), m.end()
            left_ok = pos == 0 or not _is_word_char(line[pos - 1])
            right_ok = end >= len(line) or not _is_word_char(line[end])
            if left_ok and right_ok:
                return True
        return False
    hay = line if case_sensitive else line.casefold()
    needle = q if case_sensitive else q.casefold()
    if not needle:
        return False
    start = 0
    nlen = len(needle)
    while True:
        pos = hay.find(needle, start)
        if pos < 0:
            return False
        end = pos + nlen
        if whole_word:
            left_ok = pos == 0 or not _is_word_char(line[pos - 1])
            right_ok = end >= len(line) or not _is_word_char(line[end])
            if left_ok and right_ok:
                return True
            start = pos + 1
            continue
        return True


def _read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except Exception:
            continue
    return ""


def _resolve_ellipsis_style(style: str | None = None) -> str:
    """guillemets («») oder ellipsis (…)."""
    if style is None:
        try:
            from instantlensdoc.core.app_settings import get_search_snippet_ellipsis_style

            style = get_search_snippet_ellipsis_style()
        except Exception:
            style = "guillemets"
    s = str(style or "guillemets").strip().casefold()
    if s in ("ellipsis", "…", "...", "dots", "dot"):
        return "ellipsis"
    return "guillemets"


def truncate_display_text(
    text: str,
    max_len: int = 40,
    *,
    ellipsis_style: str | None = None,
) -> str:
    """
    Text für Listen kürzen — Ellipsis-Style aus Settings:
    guillemets → endet mit «…»; ellipsis → endet mit ….
    """
    raw = str(text or "").replace("\n", " ").replace("\r", " ").strip()
    if not raw:
        return ""
    limit = max(4, int(max_len))
    if len(raw) <= limit:
        return raw
    style = _resolve_ellipsis_style(ellipsis_style)
    if style == "ellipsis":
        return raw[: max(1, limit - 1)] + "…"
    # Guillemets: sichtbarer Kürzungs-Marker «…»
    marker = "«…»"
    keep = max(1, limit - len(marker))
    return raw[:keep] + marker


def _snippet_around(
    line: str,
    query: str,
    *,
    width: int = 120,
    context_chars: int = 40,
    mark_match: bool = True,
    ellipsis_style: str | None = None,
) -> str:
    """
    Kompaktes Kontext-Snippet um den Treffer (casefold-Match).
    context_chars: Zeichen links/rechts vom Match (wenn width groß genug).
    mark_match: Match hervorheben — Style «…» (guillemets) oder nur … (ellipsis).
    """
    style = _resolve_ellipsis_style(ellipsis_style)
    snippet = (line or "").replace("\t", " ").strip()
    if not snippet:
        return ""
    q = (query or "").strip()
    ql = q.casefold()
    ll = snippet.casefold()
    pos = ll.find(ql) if ql else -1
    if pos < 0:
        out = snippet[:width]
        return out + ("…" if len(snippet) > width else "")
    qlen = len(q)
    # Bevorzugt festen Kontext um Match; sonst width-Fenster
    ctx = max(8, int(context_chars))
    start = max(0, pos - ctx)
    end = min(len(snippet), pos + qlen + ctx)
    if end - start < min(width, len(snippet)):
        # Auf width auffüllen falls Zeile länger
        need = min(width, len(snippet)) - (end - start)
        extra_left = need // 2
        extra_right = need - extra_left
        start = max(0, start - extra_left)
        end = min(len(snippet), end + extra_right)
        if end - start < need:
            start = max(0, end - need)
    match_raw = snippet[pos : pos + qlen]
    before = snippet[start:pos]
    after = snippet[pos + qlen : end]
    if mark_match and q:
        if style == "ellipsis":
            mid = match_raw
        else:
            mid = f"«{match_raw}»"
    else:
        mid = match_raw
    out = f"{before}{mid}{after}".strip()
    if start > 0:
        out = "…" + out
    if end < len(snippet):
        out = out + "…"
    if len(out) > width + 8:
        # harte Obergrenze (Markierung zählt mit)
        out = out[: width + 5] + "…"
    return out


def format_hit_line(
    path_name: str,
    *,
    page: int | None,
    line: int | None,
    snippet: str,
    kind: str = "text",
    query: str = "",
) -> str:
    """Eine Trefferzeile für die Sidebar-Liste inkl. Kontext-Snippet."""
    loc = f"S.{page + 1}" if page is not None else (f"Z.{line}" if line else "?")
    kind_mark = "·OCR " if kind == "sidecar" else ""
    snip = (snippet or "").strip()
    if not snip and query:
        snip = query
    if len(snip) > 96:
        snip = snip[:93] + "…"
    return f"{path_name} {loc}: {kind_mark}{snip}"


def _pdf_pages_text(pdf_path: Path) -> List[tuple[int, str]]:
    """Alle PDF-Seiten in einem Document-Open extrahieren (Performance)."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        parts: List[tuple[int, str]] = []
        for i in range(len(doc)):
            page = doc[i]
            try:
                tp = page.get_textpage()
                try:
                    parts.append((i, tp.get_text_bounded() or ""))
                finally:
                    tp.close()
            finally:
                page.close()
        return parts
    finally:
        doc.close()


def extract_document_text(path: str | Path) -> List[tuple[int | None, str]]:
    """(page_index, text) pro Abschnitt — page_index None bei Fließtext."""
    path = Path(path)
    if not path.is_file():
        return []
    suf = path.suffix.lower()
    if suf in _TEXT_SUFFIXES or suf == ".docx":
        if suf == ".docx":
            try:
                from docx import Document as DocxDocument

                doc = DocxDocument(str(path))
                return [(None, "\n".join(p.text for p in doc.paragraphs))]
            except Exception:
                return []
        return [(None, _read_text_file(path))]
    if suf == ".pdf":
        parts: List[tuple[int | None, str]] = []
        sidecar = path.with_suffix(path.suffix + ".ildocr.txt")
        if sidecar.is_file():
            parts.append((None, _read_text_file(sidecar)))
        try:
            for i, blob in _pdf_pages_text(path):
                parts.append((i, blob))
        except Exception:
            pass
        return parts
    sidecar = path.with_suffix(path.suffix + ".ildocr.txt")
    if sidecar.is_file():
        return [(None, _read_text_file(sidecar))]
    return []


def filter_pdf_paths(paths: Sequence[str]) -> List[str]:
    """Nur existierende PDF-Pfade (Reihenfolge erhalten, dedupe)."""
    out: List[str] = []
    seen: set[str] = set()
    for raw in paths:
        p = Path(raw)
        if not p.is_file() or p.suffix.lower() != ".pdf":
            continue
        key = str(p.resolve()) if p.exists() else str(p)
        if key in seen:
            continue
        seen.add(key)
        out.append(str(p))
    return out


def search_paths(
    paths: Sequence[str],
    query: str,
    *,
    max_hits: int = 200,
    pdf_only: bool = False,
    case_sensitive: bool = False,
    whole_word: bool = False,
    regex: bool = False,
    on_progress: Callable[[int, int, str], None] | None = None,
) -> List[SearchHit]:
    """
    Volltextsuche über Pfade.
    pdf_only=True: nur PDFs (+ optional OCR-Sidecar am PDF).
    case_sensitive / whole_word / regex — 2.0.1.
    on_progress(i, total, path) bei vielen Docs — 2.0.1.
    """
    q = (query or "").strip()
    if not q:
        return []
    pattern = _compile_pattern(q, case_sensitive=case_sensitive, regex=regex)
    try:
        from instantlensdoc.core.app_settings import get_search_snippet_context_chars

        ctx = get_search_snippet_context_chars()
    except Exception:
        ctx = 40
    hits: List[SearchHit] = []
    iter_paths = filter_pdf_paths(paths) if pdf_only else list(paths)
    total = len(iter_paths)
    for idx, raw in enumerate(iter_paths):
        p = Path(raw)
        if on_progress is not None:
            try:
                on_progress(idx + 1, total, str(p))
            except Exception:
                pass
        if not p.is_file():
            continue
        if pdf_only and p.suffix.lower() != ".pdf":
            continue
        kind = "pdf" if p.suffix.lower() == ".pdf" else "text"
        try:
            sections = extract_document_text(p)
        except Exception:
            continue
        for page_idx, blob in sections:
            if not blob:
                continue
            # Schnellfilter nur bei Literal ohne Case/Wort (Regex braucht Zeilen)
            if (
                pattern is None
                and not case_sensitive
                and not whole_word
                and q.casefold() not in blob.casefold()
            ):
                continue
            for line_no, line in enumerate(blob.splitlines(), start=1):
                if not _line_has_match(
                    line,
                    q,
                    case_sensitive=case_sensitive,
                    whole_word=whole_word,
                    pattern=pattern,
                ):
                    continue
                snippet = _snippet_around(line, q, context_chars=ctx)
                hit_kind = kind
                if page_idx is None and p.suffix.lower() == ".pdf":
                    hit_kind = "sidecar"
                hits.append(
                    SearchHit(
                        path=str(p),
                        page=page_idx,
                        line=line_no if page_idx is None else line_no,
                        snippet=snippet,
                        kind=hit_kind,
                    )
                )
                if len(hits) >= max_hits:
                    return hits
    return hits


def search_open_pdfs(
    paths: Sequence[str],
    query: str,
    *,
    max_hits: int = 200,
    case_sensitive: bool = False,
    whole_word: bool = False,
    regex: bool = False,
    on_progress: Callable[[int, int, str], None] | None = None,
) -> List[SearchHit]:
    """Schnellsuche nur über geöffnete / gelistete PDFs — Optionen 2.0.1."""
    return search_paths(
        paths,
        query,
        max_hits=max_hits,
        pdf_only=True,
        case_sensitive=case_sensitive,
        whole_word=whole_word,
        regex=regex,
        on_progress=on_progress,
    )


def sidebar_document_paths(files_widget) -> List[str]:
    """Pfade aus Sidebar-Dokumentenliste."""
    out: List[str] = []
    for i in range(files_widget.count()):
        item = files_widget.item(i)
        if item is None:
            continue
        data = item.data(256)
        if data:
            out.append(str(data))
    return out


SEARCH_HIT_CSV_FIELDS = (
    "index",
    "label",
    "path",
    "page",
    "offset",
    "line",
    "kind",
    "query",
    "snippet",
)


def normalize_search_hit_record(raw: dict | None, *, index: int = 1, query: str = "") -> dict:
    """Eine Trefferzeile für CSV/JSON-Export normalisieren (inkl. Offset 0.9.4)."""
    src = dict(raw or {})
    page = src.get("page", "")
    line = src.get("line", "")
    offset = src.get("offset", "")
    if page is None:
        page = ""
    if line is None:
        line = ""
    if offset is None:
        offset = ""
    if page != "" and not isinstance(page, str):
        try:
            page = int(page)
        except (TypeError, ValueError):
            page = str(page)
    if line != "" and not isinstance(line, str):
        try:
            line = int(line)
        except (TypeError, ValueError):
            line = str(line)
    if offset != "" and not isinstance(offset, str):
        try:
            offset = int(offset)
        except (TypeError, ValueError):
            offset = str(offset)
    return {
        "index": int(src.get("index") or index),
        "label": str(src.get("label") or ""),
        "path": str(src.get("path") or ""),
        "page": page,
        "offset": offset,
        "line": line,
        "kind": str(src.get("kind") or "mark"),
        "query": str(src.get("query") if src.get("query") is not None else query),
        "snippet": str(src.get("snippet") or src.get("label") or ""),
    }


def export_search_hits_csv(
    path: str | Path,
    hits: Sequence[dict],
    *,
    query: str = "",
) -> Path:
    """Suchergebnisse als CSV exportieren (UTF-8)."""
    import csv

    dest = Path(path)
    if dest.suffix.lower() != ".csv":
        dest = dest.with_suffix(".csv")
    dest.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        normalize_search_hit_record(h, index=i + 1, query=query)
        for i, h in enumerate(hits or [])
    ]
    with dest.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=list(SEARCH_HIT_CSV_FIELDS), extrasaction="ignore"
        )
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in SEARCH_HIT_CSV_FIELDS})
    return dest


# Multi-Dokument-Suche CSV: Doc,Seite,Snippet,Match — optional UTF-8 BOM — 2.0.2
MULTI_DOC_SEARCH_CSV_FIELDS = (
    "Doc",
    "Seite",
    "Snippet",
    "Match",
)


def export_multi_doc_hits_csv(
    path: str | Path,
    hits: Sequence[SearchHit | dict],
    *,
    query: str = "",
    utf8_bom: bool = True,
) -> Path:
    """
    Multi-Doc-Treffer als CSV: Spalten Doc,Seite,Snippet,Match.
    SearchHit.page ist 0-basiert → CSV Seite 1-basiert.
    Encoding UTF-8 mit optionalem BOM — 2.0.2.
    """
    import csv

    dest = Path(path)
    if dest.suffix.lower() != ".csv":
        dest = dest.with_suffix(".csv")
    dest.parent.mkdir(parents=True, exist_ok=True)
    enc = "utf-8-sig" if utf8_bom else "utf-8"
    q = str(query or "")
    with dest.open("w", encoding=enc, newline="") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=list(MULTI_DOC_SEARCH_CSV_FIELDS), extrasaction="ignore"
        )
        writer.writeheader()
        for h in hits or []:
            if isinstance(h, SearchHit):
                doc = Path(h.path).name if h.path else ""
                page: int | str = (int(h.page) + 1) if h.page is not None else ""
                snippet = str(h.snippet or "")
                match = q
            else:
                src = dict(h or {})
                p = str(src.get("path") or "")
                doc = Path(p).name if p else str(src.get("Doc") or src.get("doc") or "")
                if "Seite" in src and src.get("Seite") not in (None, ""):
                    page = src["Seite"]
                elif src.get("page") is None or src.get("page") == "":
                    page = ""
                else:
                    try:
                        # Dict mit 0-basiertem page (wie SearchHit)
                        page = int(src["page"]) + 1
                    except (TypeError, ValueError):
                        page = str(src.get("page"))
                snippet = str(src.get("snippet") or src.get("Snippet") or "")
                match = str(
                    src.get("Match") or src.get("query") or src.get("match") or q
                )
            writer.writerow(
                {
                    "Doc": doc,
                    "Seite": page,
                    "Snippet": snippet,
                    "Match": match,
                }
            )
    return dest


# Felder je Treffer in ildsearch-v1 (kompatibel erweitert um offset — 0.9.4/0.9.5)
SEARCH_HIT_JSON_FIELDS = SEARCH_HIT_CSV_FIELDS


def export_search_hits_json(
    path: str | Path,
    hits: Sequence[dict],
    *,
    query: str = "",
) -> Path:
    """Suchergebnisse als JSON exportieren (Schema ildsearch-v1, inkl. Offset)."""
    import json
    from datetime import datetime, timezone

    dest = Path(path)
    if dest.suffix.lower() != ".json":
        dest = dest.with_suffix(".json")
    dest.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        normalize_search_hit_record(h, index=i + 1, query=query)
        for i, h in enumerate(hits or [])
    ]
    # Explizite Feldreihenfolge / Vollständigkeit (offset inkl.)
    ordered = [
        {k: row.get(k, "") for k in SEARCH_HIT_JSON_FIELDS}
        for row in rows
    ]
    payload = {
        "version": 1,
        "schema": "ildsearch-v1",
        "exported_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "query": str(query or ""),
        "count": len(ordered),
        "fields": list(SEARCH_HIT_JSON_FIELDS),
        "hits": ordered,
    }
    dest.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return dest
