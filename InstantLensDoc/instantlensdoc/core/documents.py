"""Dokument-I/O: TXT, MD, HTML, DOCX, CSV, Excel, PDF, Bilder."""

from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Optional

# Editor-Textcodierungen (Öffnen/Speichern); "auto" = BOM/chardet-Erkennung
TEXT_ENCODINGS = ("auto", "utf-8", "latin-1")
ENCODING_LABELS = {
    "auto": "Automatisch (BOM / chardet)",
    "utf-8": "UTF-8",
    "latin-1": "Latin-1 (ISO-8859-1)",
}


def normalize_text_encoding(encoding: str | None) -> str:
    """utf-8 / latin-1 / auto; Default utf-8."""
    enc = (encoding or "utf-8").strip().lower().replace("_", "-")
    if enc in ("auto", "detect", "automatic"):
        return "auto"
    if enc in ("utf8", "utf-8", "utf"):
        return "utf-8"
    if enc in ("latin-1", "latin1", "iso-8859-1", "iso8859-1", "cp1252", "windows-1252"):
        # cp1252-ähnlich bewusst auf latin-1 mappen (Aufgabenumfang)
        return "latin-1"
    return "utf-8"


def detect_file_encoding(path: str | Path, sample_size: int = 65536) -> str:
    """
    Datei-Encoding erkennen: BOM → UTF-8-Probe → optional chardet → latin-1.
    Rückgabe immer normalisiert (utf-8 oder latin-1), nie auto.
    """
    path = Path(path)
    try:
        raw = path.read_bytes()[: max(1024, int(sample_size))]
    except Exception:
        return "utf-8"
    if not raw:
        return "utf-8"
    # BOM
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8"
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        # UTF-16 nicht nativ unterstützt → als latin-1-Fallback vermeiden; utf-8 lesen mit replace
        return "utf-8"
    # Strikte UTF-8-Probe
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass
    # Optional chardet
    try:
        import chardet  # type: ignore

        guess = chardet.detect(raw) or {}
        enc = normalize_text_encoding(guess.get("encoding"))
        if enc == "auto":
            enc = "utf-8"
        # chardet liefert oft windows-1252 → latin-1
        if enc in ("utf-8", "latin-1"):
            return enc
    except Exception:
        pass
    return "latin-1"


def resolve_text_encoding(path: str | Path | None, encoding: str | None) -> str:
    """Encoding auflösen: auto → detect_file_encoding, sonst normalize."""
    enc = normalize_text_encoding(encoding)
    if enc == "auto":
        if path is None:
            return "utf-8"
        return detect_file_encoding(path)
    return enc


# Vorlagen für „Neues leeres Dokument“ (id → Titel, Text)
DOC_TEMPLATES: dict[str, tuple[str, str]] = {
    "empty": ("Unbenannt", ""),
    "brief": (
        "Brief",
        "{name}\n{street}\n{city}\n\n"
        "{date}\n\n"
        "Sehr geehrte Damen und Herren,\n\n"
        "[Text]\n\n"
        "Mit freundlichen Grüßen\n\n"
        "[Unterschrift]\n"
    ),
    "notiz": (
        "Notiz",
        "# Notiz\n\n"
        "Datum: {date}\n"
        "Thema: \n\n"
        "- \n"
        "- \n\n"
        "———\n",
    ),
}


def render_doc_template(template_id: str = "empty") -> tuple[str, str]:
    """
    Dokument-Vorlage rendern.
    Rückgabe: (Titel, Text). Unbekannte ID → leeres Dokument.
    Unterstützt eingebaute IDs und Nutzer-Vorlagen (user:… / UUID).
    """
    from datetime import date

    key = (template_id or "empty").strip()
    key_l = key.lower()
    title: str
    body: str
    if key_l in DOC_TEMPLATES:
        title, body = DOC_TEMPLATES[key_l]
    else:
        # Nutzer-Vorlage aus Einstellungen
        try:
            from instantlensdoc.core.app_settings import get_user_doc_template

            ut = get_user_doc_template(key)
        except Exception:
            ut = None
        if ut:
            title = str(ut.get("title") or "Vorlage")
            body = str(ut.get("body") or "")
        else:
            title, body = DOC_TEMPLATES["empty"]
    today = date.today().strftime("%d.%m.%Y")
    try:
        text = body.format(
            name="[Name]",
            street="[Straße Nr.]",
            city="[PLZ Ort]",
            date=today,
        )
    except (KeyError, ValueError, IndexError):
        # Nutzer-Vorlagen dürfen geschweifte Klammern enthalten
        text = body.replace("{date}", today)
    return title, text


def save_current_as_template(title: str, body: str) -> dict:
    """Aktuellen Editor-Text als Nutzer-Vorlage speichern."""
    from instantlensdoc.core.app_settings import save_user_doc_template

    return save_user_doc_template(title=title, body=body)


class DocKind(str, Enum):
    TEXT = "text"
    MARKDOWN = "markdown"
    HTML = "html"
    DOCX = "docx"
    RTF = "rtf"
    CSV = "csv"
    XLSX = "xlsx"
    PDF = "pdf"
    IMAGE = "image"
    UNKNOWN = "unknown"


# .xls ohne xlwt: Inhalt ist OOXML/.xlsx unter .xls-Endung
XLS_WRITE_VIA_XLSX_DE = (
    "xlwt fehlt — die Datei wurde im Excel-2007-Format (.xlsx) unter der "
    "Endung .xls gespeichert. Excel öffnet sie meist trotzdem; für echtes "
    "Excel-97-.xls bitte xlwt installieren."
)

RICH_EDITOR_KINDS = (
    DocKind.DOCX,
    DocKind.HTML,
    DocKind.MARKDOWN,
    DocKind.CSV,
    DocKind.XLSX,
)


@dataclass
class Document:
    path: Optional[Path] = None
    kind: DocKind = DocKind.TEXT
    text: str = ""
    dirty: bool = False
    title: str = "Unbenannt"
    meta: dict = field(default_factory=dict)

    @property
    def display_name(self) -> str:
        if self.path:
            return self.path.name
        return self.title

    @property
    def encoding(self) -> str:
        enc = normalize_text_encoding(self.meta.get("encoding"))
        return "utf-8" if enc == "auto" else enc


def detect_kind(path: Path) -> DocKind:
    ext = path.suffix.lower()
    if ext in {".txt", ".log", ".ild"}:
        # .ild = natives InstantLens-Doc (UTF-8-Text) — 2.6.43
        return DocKind.TEXT
    if ext == ".csv":
        return DocKind.CSV
    if ext in {".md", ".markdown"}:
        return DocKind.MARKDOWN
    if ext in {".html", ".htm"}:
        return DocKind.HTML
    if ext == ".docx":
        return DocKind.DOCX
    if ext == ".rtf":
        return DocKind.RTF
    if ext in {".xlsx", ".xlsm", ".xls"}:
        return DocKind.XLSX
    if ext == ".pdf":
        return DocKind.PDF
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}:
        return DocKind.IMAGE
    return DocKind.UNKNOWN


_SNIFF_KIND_TO_DOCKIND = {
    "zip-docx": DocKind.DOCX,
    "html": DocKind.HTML,
    "rtf": DocKind.RTF,
    "text": DocKind.TEXT,
    "png": DocKind.IMAGE,
    "jpeg": DocKind.IMAGE,
}


def detect_kind_mismatch(path: Path) -> tuple[DocKind, str] | None:
    """Datei mit ``.pdf``-Endung, deren **Inhalt** kein PDF ist (DOCX/HTML/RTF/Text/Bild).

    Rückgabe ``(echter DocKind, deutscher Hinweis)`` oder ``None``. Feldfall 2.6.53:
    ``Dunning_Kruger_Effekt_1.pdf`` war ein byte-identisches Word-DOCX — statt
    „Data format error“ wird die Datei mit ihrem echten Typ geöffnet — 2.6.54.
    """
    if path.suffix.lower() != ".pdf":
        return None
    try:
        from ild_pdf.pdf_sniff import describe_non_pdf_de, sniff_file

        sn = sniff_file(path)
    except Exception:
        return None
    if sn.is_pdf_like or sn.kind in ("empty", "unreadable"):
        return None
    real = _SNIFF_KIND_TO_DOCKIND.get(sn.kind)
    if real is None:
        return None
    return real, describe_non_pdf_de(sn)


def open_document(path: str | Path, *, encoding: str | None = None) -> Document:
    from instantlensdoc.core.fs_path import as_native_path

    path = as_native_path(path)
    if not path.exists():
        raise FileNotFoundError(f"Datei nicht gefunden: {path}")
    if not path.is_file():
        raise IsADirectoryError(f"Kein Dateipfad: {path}")
    kind = detect_kind(path)
    mismatch = detect_kind_mismatch(path)
    doc = Document(path=path, kind=kind, title=path.name)
    if mismatch is not None:
        real_kind, hint = mismatch
        import logging

        logging.getLogger("instantlensdoc.documents").warning(
            "Endung .pdf, Inhalt %s — Datei wird als %s geöffnet: %s",
            real_kind.name,
            real_kind.name,
            path,
        )
        kind = real_kind
        doc.kind = kind
        doc.meta["kind_mismatch"] = hint
        doc.meta["kind_mismatch_kind"] = real_kind.name
        # Speichern (Strg+S) würde den falsch benannten .pdf-Namen weiter als DOCX/HTML
        # beschreiben → Nutzer soll „Speichern unter“ mit richtiger Endung wählen
        doc.meta["readonly"] = True

    if encoding is None:
        try:
            from instantlensdoc.core.app_settings import get_editor_text_encoding

            encoding = get_editor_text_encoding()
        except Exception:
            encoding = "utf-8"
    enc = resolve_text_encoding(path, encoding)
    requested = normalize_text_encoding(encoding)
    if requested == "auto":
        doc.meta["encoding_detected"] = True

    if kind == DocKind.TEXT:
        doc.text = path.read_text(encoding=enc, errors="replace")
        doc.meta["encoding"] = enc
    elif kind == DocKind.MARKDOWN:
        from instantlensdoc.core.markdown_io import markdown_to_html

        doc.text = path.read_text(encoding=enc, errors="replace")
        doc.meta["encoding"] = enc
        doc.meta["html"] = markdown_to_html(doc.text)
        doc.meta["rich_text"] = True
    elif kind == DocKind.CSV:
        from ild_pdf.tables import (
            import_csv,
            sniff_csv_delimiter,
            table_to_html_document,
            table_to_markdown,
        )

        table = import_csv(path, encoding=enc)
        raw = path.read_bytes()[:4096]
        try:
            sample = raw.decode(enc if enc != "utf-8" else "utf-8-sig", errors="replace")
        except Exception:
            sample = path.read_text(encoding=enc, errors="replace")[:4096]
        delim = sniff_csv_delimiter(sample)
        html = table_to_html_document(table, title=path.stem)
        doc.text = table_to_markdown(table, with_markers=False)
        doc.meta["encoding"] = enc
        doc.meta["html"] = html
        doc.meta["rich_text"] = True
        doc.meta["table"] = table.to_dict()
        doc.meta["csv_delimiter"] = delim
    elif kind == DocKind.HTML:
        doc.text = path.read_text(encoding=enc, errors="replace")
        doc.meta["encoding"] = enc
        doc.meta["html"] = doc.text
        doc.meta["rich_text"] = True
    elif kind == DocKind.DOCX:
        try:
            from instantlensdoc.core.richtext_docx import docx_plain_and_html

            plain, html = docx_plain_and_html(path)
            doc.text = plain
            doc.meta["html"] = html
            doc.meta["rich_text"] = True
            # Standardschrift (Normal-Stil/docDefaults) für den Editor — 2.6.53
            try:
                from instantlensdoc.core.richtext_docx import docx_default_font

                fam, size_pt = docx_default_font(path)
                if fam:
                    doc.meta["font_family"] = str(fam)
                if size_pt:
                    doc.meta["font_size_pt"] = float(size_pt)
            except Exception:
                pass
            try:
                from instantlensdoc.core.richtext_docx import docx_header_footer_author

                header, footer, author = docx_header_footer_author(path)
                if header:
                    doc.meta["header"] = header
                if footer:
                    doc.meta["footer"] = footer
                if author:
                    doc.meta.setdefault("author", author)
            except Exception:
                pass
            try:
                from instantlensdoc.core.field_tokens import (
                    extract_field_tokens,
                    load_field_tokens_sidecar,
                    merge_field_specs,
                    read_docx_custom_field_tokens,
                    serialize_field_specs,
                )

                merged = merge_field_specs(
                    extract_field_tokens((doc.text or "") + "\n" + str(doc.meta.get("html") or "")),
                    read_docx_custom_field_tokens(path),
                    load_field_tokens_sidecar(path),
                    (doc.meta or {}).get("field_tokens"),
                )
                if merged:
                    doc.meta["field_tokens"] = serialize_field_specs(merged)
            except Exception:
                pass
        except Exception as rich_err:
            # Nicht still degradieren: Grund loggen + im Meta vermerken — 2.6.52
            import logging

            logging.getLogger("instantlensdoc.documents").warning(
                "DOCX-Rich-Text-Import fehlgeschlagen, Plaintext-Fallback (%s): %s",
                path,
                rich_err,
            )
            doc.meta["rich_text_error"] = str(rich_err)
            try:
                from instantlensdoc.core.export import import_document_text

                loaded = import_document_text(path)
                doc.text = loaded.get("text") or ""
                doc.meta.update(loaded.get("meta") or {})
            except Exception:
                try:
                    from docx import Document as DocxDocument

                    d = DocxDocument(str(path))
                    doc.text = "\n".join(p.text for p in d.paragraphs)
                except ImportError:
                    doc.text = "[python-docx nicht installiert]"
                    doc.meta["error"] = "python-docx fehlt"
    elif kind == DocKind.RTF:
        from instantlensdoc.core.export import import_rtf

        doc.text = import_rtf(path)
        doc.meta["encoding"] = "utf-8"
        # RTF-Import bleibt plaintext; Editor-Formate beim Speichern via meta.html
    elif kind == DocKind.XLSX:
        from ild_pdf.tables import (
            import_spreadsheet,
            list_spreadsheet_sheets,
            table_to_html_document,
            table_to_markdown,
        )

        sheets = list_spreadsheet_sheets(path)
        table = import_spreadsheet(path, sheet=0)
        active = sheets[0] if sheets else "Tabelle1"
        html = table_to_html_document(table, title=active)
        doc.text = table_to_markdown(table, with_markers=False)
        doc.meta["html"] = html
        doc.meta["rich_text"] = True
        doc.meta["table"] = table.to_dict()
        doc.meta["sheets"] = sheets
        doc.meta["active_sheet"] = active
        if path.suffix.lower() == ".xls":
            doc.meta["xls_legacy"] = True
    elif kind == DocKind.PDF:
        doc.text = ""  # PDF wird über Viewer gerendert
        doc.meta["pdf"] = str(path)
    elif kind == DocKind.IMAGE:
        doc.text = ""
        doc.meta["image"] = str(path)
    else:
        doc.text = path.read_text(encoding=enc, errors="replace")
        doc.kind = DocKind.TEXT
        doc.meta["encoding"] = enc
    try:
        from instantlensdoc.core.field_tokens import (
            load_field_tokens_sidecar,
            merge_field_specs,
            serialize_field_specs,
        )

        side = load_field_tokens_sidecar(path)
        if side:
            merged = merge_field_specs((doc.meta or {}).get("field_tokens"), side)
            doc.meta["field_tokens"] = serialize_field_specs(merged)
    except Exception:
        pass
    try:
        from instantlensdoc.core.document_author import ensure_author, load_author

        kind_name = str(getattr(kind, "name", kind) or "")
        author = load_author(path, kind=kind_name, meta=doc.meta)
        doc.meta["author"] = ensure_author({"author": author})
    except Exception:
        pass
    try:
        from instantlensdoc.core.write_protect import detect_protection

        info = detect_protection(path)
        if info.protected:
            doc.meta["write_protect"] = True
            if info.password_hash:
                doc.meta["write_protect_hash"] = info.password_hash
    except Exception:
        pass
    return doc


def backup_existing(path: Path) -> Path | None:
    """Kopiert vorhandene Datei nach path.bak (überschreibt älteres .bak)."""
    path = Path(path)
    if not path.is_file():
        return None
    bak = Path(str(path) + ".bak")
    shutil.copy2(path, bak)
    return bak


def backup_ildbak(path: Path, max_backups: int = 3) -> Path | None:
    """
    Autosave-Backup als path.ext.ildbak mit Rotation (max. 1–10) — 0.9.9.

    Neueste Kopie: ``dateiname.ext.ildbak``;
    ältere: ``dateiname.ext.1.ildbak`` … ``dateiname.ext.(n-1).ildbak``.
    """
    path = Path(path)
    if not path.is_file():
        return None
    try:
        n = int(max_backups)
    except (TypeError, ValueError):
        n = 3
    n = max(1, min(10, n))
    base = str(path)
    newest = Path(base + ".ildbak")
    if n == 1:
        shutil.copy2(path, newest)
        return newest
    oldest = Path(f"{base}.{n - 1}.ildbak")
    if oldest.is_file():
        try:
            oldest.unlink()
        except OSError:
            pass
    for i in range(n - 2, 0, -1):
        src = Path(f"{base}.{i}.ildbak")
        dst = Path(f"{base}.{i + 1}.ildbak")
        if src.is_file():
            try:
                src.replace(dst)
            except OSError:
                try:
                    shutil.copy2(src, dst)
                    src.unlink()
                except OSError:
                    pass
    if newest.is_file():
        slot1 = Path(f"{base}.1.ildbak")
        try:
            newest.replace(slot1)
        except OSError:
            try:
                shutil.copy2(newest, slot1)
                newest.unlink()
            except OSError:
                pass
    shutil.copy2(path, newest)
    return newest


def save_document(
    doc: Document,
    path: Optional[Path] = None,
    *,
    encoding: str | None = None,
) -> Path:
    from instantlensdoc.core.fs_path import as_native_path

    target = as_native_path(path or doc.path or "unbenannt.txt")
    kind = detect_kind(target) if path else doc.kind

    # Optionale Backup-Kopie vor Überschreiben (Einstellungen → Backup .bak)
    try:
        from instantlensdoc.core.app_settings import get_backup_on_save

        if get_backup_on_save() and target.is_file():
            backup_existing(target)
    except Exception:
        pass

    if encoding is None:
        encoding = doc.meta.get("encoding")
    if encoding is None:
        try:
            from instantlensdoc.core.app_settings import get_editor_text_encoding

            encoding = get_editor_text_encoding()
        except Exception:
            encoding = "utf-8"
    # Beim Speichern nie "auto" schreiben — erkanntes/gewähltes Encoding nutzen
    enc = resolve_text_encoding(target if normalize_text_encoding(encoding) == "auto" else None, encoding)
    if normalize_text_encoding(encoding) == "auto" and doc.meta.get("encoding") in ("utf-8", "latin-1"):
        enc = normalize_text_encoding(doc.meta.get("encoding"))

    if kind == DocKind.DOCX:
        html = (doc.meta or {}).get("html")
        if html:
            from instantlensdoc.core.richtext_docx import html_to_docx

            html_to_docx(
                str(html),
                target,
                title=doc.title,
                header=str((doc.meta or {}).get("header") or ""),
                footer=str((doc.meta or {}).get("footer") or ""),
                author=str((doc.meta or {}).get("author") or ""),
            )
        else:
            from instantlensdoc.core.export import export_docx

            export_docx(doc.text, target, title=doc.title, author=str((doc.meta or {}).get("author") or ""))
        try:
            from instantlensdoc.core.field_tokens import persist_document_field_tokens

            persist_document_field_tokens(
                target,
                (doc.meta or {}).get("field_tokens"),
                kind="docx",
            )
        except Exception:
            pass
    elif kind == DocKind.RTF:
        html = (doc.meta or {}).get("html")
        if html:
            from instantlensdoc.core.richtext_docx import html_to_rtf

            html_to_rtf(str(html), target, title=doc.title)
        else:
            from instantlensdoc.core.export import export_rtf

            export_rtf(doc.text, target, title=doc.title)
    elif kind == DocKind.MARKDOWN:
        from instantlensdoc.core.markdown_io import html_to_markdown, strip_pilcrow

        html = (doc.meta or {}).get("html")
        md_src = (doc.meta or {}).get("markdown_source")
        if html:
            text = html_to_markdown(str(html))
        elif md_src:
            text = str(md_src)
        else:
            text = doc.text or ""
        text = strip_pilcrow(text)
        if "¶" in text:
            text = text.replace("¶", "")
        target.write_text(text, encoding=enc, errors="replace")
        doc.meta["encoding"] = enc
        doc.meta["html"] = html or doc.meta.get("html")
    elif kind == DocKind.CSV:
        from ild_pdf.tables import (
            create_table,
            export_table_csv,
        )

        delim = str((doc.meta or {}).get("csv_delimiter") or ";")
        if delim not in (",", ";", "\t", "|"):
            delim = ";"
        table = _table_from_doc(doc)
        if table is None:
            rows = [[line] for line in (doc.text or "").splitlines()] or [[""]]
            table = create_table(data=rows, header=False)
        csv_enc = enc if enc != "utf-8" else "utf-8-sig"
        export_table_csv(table, target, delimiter=delim, encoding=csv_enc)
        doc.meta["encoding"] = csv_enc
        doc.meta["csv_delimiter"] = delim
    elif kind == DocKind.XLSX:
        from ild_pdf.tables import create_table, export_table_xlsx

        table = _table_from_doc(doc)
        if table is None:
            from instantlensdoc.core.export import export_xlsx_from_text

            export_xlsx_from_text(doc.text or "", target)
        else:
            sheet = str((doc.meta or {}).get("active_sheet") or "Tabelle1")
            if target.suffix.lower() == ".xls":
                _save_xls(table, target, sheet_name=sheet, doc=doc)
            else:
                export_table_xlsx(table, target, sheet_name=sheet)
    elif kind == DocKind.HTML:
        html = (doc.meta or {}).get("html") or doc.text or ""
        stripped = (html or "").lstrip().lower()
        if stripped.startswith("<!doctype") or stripped.startswith("<html") or "<p" in stripped or "<b" in stripped:
            target.write_text(html, encoding=enc, errors="replace")
            doc.meta["encoding"] = enc
            doc.meta["html"] = html
        else:
            from instantlensdoc.core.export import export_html

            export_html(doc.text, target, title=doc.title or target.stem)
    elif kind == DocKind.PDF:
        # Ziel .pdf: NIE eine Nicht-PDF-Quelle byte-identisch kopieren. Genau das
        # erzeugte im Feld (≤ 2.6.42: „Speichern unter … .pdf“ aus einem DOCX) eine
        # Datei mit .pdf-Endung und ZIP-Inhalt → PDFium „Data format error“, pikepdf
        # „unable to find trailer dictionary“, Edge „ungültiges Dateiformat“ — 2.6.54
        from ild_pdf.pdf_sniff import assert_valid_pdf, sniff_file

        src_is_pdf = False
        if doc.kind == DocKind.PDF and doc.path and doc.path.is_file():
            src_is_pdf = sniff_file(doc.path).is_pdf_like
        if src_is_pdf:
            # Echte PDF-Datei: Inhalt wird über Annotation-Sidecar / pikepdf verwaltet
            if doc.path.resolve() != target.resolve():
                shutil.copy2(doc.path, target)
                assert_valid_pdf(target)
        else:
            # Editor-Dokument (TXT/MD/HTML/DOCX/RTF …) → echtes PDF über den Exporter;
            # formatiertes HTML (DOCX/HTML) wird — wenn Qt läuft — mit Formaten gesetzt.
            from instantlensdoc.core.export import export_pdf

            html = (doc.meta or {}).get("html") if doc.kind != DocKind.PDF else None
            export_pdf(
                doc.text or "",
                target,
                title=doc.title or target.stem,
                html=str(html) if html else None,
            )
    elif kind == DocKind.IMAGE:
        if doc.path and doc.path.resolve() != target.resolve():
            shutil.copy2(doc.path, target)
    else:
        target.write_text(doc.text, encoding=enc, errors="replace")
        doc.meta["encoding"] = enc
    doc.path = target
    doc.kind = kind
    doc.dirty = False
    doc.title = target.name
    try:
        from instantlensdoc.core.field_tokens import persist_document_field_tokens

        if kind != DocKind.DOCX:
            persist_document_field_tokens(
                target,
                (doc.meta or {}).get("field_tokens"),
                kind=kind.name if hasattr(kind, "name") else str(kind),
            )
    except Exception:
        pass
    kind_name = kind.name if hasattr(kind, "name") else str(kind)
    try:
        from instantlensdoc.core.document_author import ensure_author, persist_author

        author = ensure_author(doc.meta)
        doc.meta["author"] = author
        persist_author(target, author, kind=kind_name)
    except Exception:
        pass
    try:
        from instantlensdoc.core.write_protect import persist_protection

        persist_protection(
            target,
            kind=kind_name,
            protected=bool((doc.meta or {}).get("write_protect")),
            password_hash_value=str((doc.meta or {}).get("write_protect_hash") or ""),
            password=(doc.meta or {}).get("write_protect_password"),
        )
    except Exception:
        pass
    return target


def _table_from_doc(doc: Document):
    """Aktuelle Tabelle aus Meta oder HTML (CSV/Excel-Speichern)."""
    from ild_pdf.tables import (
        DocumentTable,
        TableFormat,
        cells_from_html_tables,
        create_table,
    )

    raw = (doc.meta or {}).get("table")
    if isinstance(raw, dict) and raw.get("cells"):
        fmt_raw = raw.get("format") if isinstance(raw.get("format"), dict) else {}
        allowed = {k: fmt_raw[k] for k in ("header", "border", "align", "header_bold", "style") if k in fmt_raw}
        fmt = TableFormat(**allowed) if allowed else TableFormat()
        return DocumentTable(
            cells=list(raw.get("cells") or []),
            format=fmt,
            table_id=str(raw.get("id") or "t1"),
        ).normalized()
    html = str((doc.meta or {}).get("html") or "")
    tables = cells_from_html_tables(html)
    if tables:
        return create_table(data=tables[0], header=True)
    return None


def _save_xls(table, target: Path, *, sheet_name: str, doc: Document) -> None:
    """Echtes .xls via xlwt, sonst XLSX-Bytes unter .xls + deutsche Warnung."""
    from ild_pdf.tables import export_table_xlsx

    try:
        import xlwt  # type: ignore
    except ImportError:
        export_table_xlsx(table, target, sheet_name=sheet_name)
        doc.meta["xls_write_warning"] = XLS_WRITE_VIA_XLSX_DE
        return
    wb = xlwt.Workbook()
    ws = wb.add_sheet((sheet_name or "Tabelle1")[:31] or "Tabelle1")
    t = table.normalized()
    for r_i, row in enumerate(t.cells):
        for c_i, val in enumerate(row):
            ws.write(r_i, c_i, val)
    wb.save(str(target))
    doc.meta.pop("xls_write_warning", None)
