"""DOCX ↔ HTML Rich-Text (Bold/Italic/Underline/Absätze) — 2.6.49.

Lädt DOCX via python-docx mit echten Zeichenformaten in HTML für QTextDocument.
Exportiert HTML (inkl. Qt-``toHtml()``-Spans) zurück nach DOCX/RTF ohne Markdown-Marker.
"""

from __future__ import annotations

import html as html_lib
import logging
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Optional


_log = logging.getLogger("instantlensdoc.richtext_docx")

# Word-Highlight-Index → HTML-Farbe (für Textmarker im Editor) — 2.6.52
_HIGHLIGHT_INDEX_TO_HEX: dict[str, str] = {
    "YELLOW": "#ffff00",
    "BRIGHT_GREEN": "#00ff00",
    "TURQUOISE": "#00ffff",
    "PINK": "#ff00ff",
    "BLUE": "#0000ff",
    "RED": "#ff0000",
    "DARK_BLUE": "#000080",
    "TEAL": "#008080",
    "GREEN": "#008000",
    "VIOLET": "#800080",
    "DARK_RED": "#800000",
    "DARK_YELLOW": "#808000",
    "GRAY_50": "#808080",
    "GRAY_25": "#c0c0c0",
    "BLACK": "#000000",
    "WHITE": "#ffffff",
}


def _escape(text: str) -> str:
    return html_lib.escape(text or "", quote=False).replace("\n", "<br/>")


def _expand_hex(value: str) -> str | None:
    v = (value or "").strip().lower()
    if not v:
        return None
    if not v.startswith("#"):
        v = "#" + v
    if re.fullmatch(r"#[0-9a-f]{3}", v):
        v = "#" + "".join(ch * 2 for ch in v[1:])
    if re.fullmatch(r"#[0-9a-f]{6}", v):
        return v
    return None


def _closest_highlight_index(hex_color: str):
    """HTML-Hintergrund → nächstliegender Word-Highlight-Index (WD_COLOR_INDEX)."""
    from docx.enum.text import WD_COLOR_INDEX

    hx = _expand_hex(hex_color) or "#ffff00"
    r, g, b = int(hx[1:3], 16), int(hx[3:5], 16), int(hx[5:7], 16)
    best_name = "YELLOW"
    best_dist = None
    for name, ref in _HIGHLIGHT_INDEX_TO_HEX.items():
        if name in ("BLACK", "WHITE"):
            continue
        rr, gg, bb = int(ref[1:3], 16), int(ref[3:5], 16), int(ref[5:7], 16)
        dist = (r - rr) ** 2 + (g - gg) ** 2 + (b - bb) ** 2
        if best_dist is None or dist < best_dist:
            best_dist = dist
            best_name = name
    return getattr(WD_COLOR_INDEX, best_name, WD_COLOR_INDEX.YELLOW)


def _style_chain(style: Any) -> list[Any]:
    """Style + base_style-Kette (zyklussicher)."""
    out: list[Any] = []
    seen: set[int] = set()
    cur = style
    while cur is not None and id(cur) not in seen and len(out) < 16:
        seen.add(id(cur))
        out.append(cur)
        try:
            cur = cur.base_style
        except Exception:
            cur = None
    return out


def _font_attr(font: Any, name: str) -> Any:
    try:
        return getattr(font, name)
    except Exception:
        return None


def _effective_tri(run: Any, para: Any, attr: str) -> bool:
    """Effektiver Bool-Wert (bold/italic/underline) inkl. Zeichen-/Absatzstil.

    python-docx liefert ``None`` wenn nicht direkt gesetzt (= „erben“). Viele
    Word-Dokumente formatieren über Stile (Strong, Heading 1, …) — vor 2.6.52
    ging diese Formatierung komplett verloren („DOCX wie Plaintext“).
    """
    # 1) direkt am Run
    val = _font_attr(getattr(run, "font", None), attr)
    if val is None:
        val = getattr(run, attr, None)
    if val is not None:
        return bool(val)
    # 2) Zeichenstil-Kette des Runs
    try:
        rstyle = run.style
    except Exception:
        rstyle = None
    for st in _style_chain(rstyle):
        v = _font_attr(getattr(st, "font", None), attr)
        if v is not None:
            return bool(v)
    # 3) Absatzstil-Kette
    try:
        pstyle = para.style if para is not None else None
    except Exception:
        pstyle = None
    for st in _style_chain(pstyle):
        v = _font_attr(getattr(st, "font", None), attr)
        if v is not None:
            return bool(v)
    return False


def _effective_size_pt(run: Any, para: Any) -> float | None:
    for owner in (run,):
        sz = _font_attr(getattr(owner, "font", None), "size")
        if sz is not None:
            try:
                return float(sz.pt)
            except Exception:
                pass
    for st in _style_chain(getattr(run, "style", None)):
        sz = _font_attr(getattr(st, "font", None), "size")
        if sz is not None:
            try:
                return float(sz.pt)
            except Exception:
                pass
    return None


def _effective_font_name(run: Any, para: Any) -> str | None:
    """Explizite Schriftart am Run bzw. Zeichenstil (None = Dokument-Standard) — 2.6.53."""
    name = _font_attr(getattr(run, "font", None), "name")
    if name:
        return str(name)
    for st in _style_chain(getattr(run, "style", None)):
        v = _font_attr(getattr(st, "font", None), "name")
        if v:
            return str(v)
    return None


def _docx_default_font_from_styles_xml(d: Any) -> tuple[str | None, float | None]:
    """docDefaults ``w:rPrDefault`` (rFonts/ sz) — python-docx hat dafür kein API."""
    try:
        styles_el = d.styles.element
        ns = styles_el.nsmap.get("w") or "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        rpr = styles_el.find(f"{{{ns}}}docDefaults/{{{ns}}}rPrDefault/{{{ns}}}rPr")
        if rpr is None:
            return None, None
        fam = None
        fonts = rpr.find(f"{{{ns}}}rFonts")
        if fonts is not None:
            fam = fonts.get(f"{{{ns}}}ascii") or fonts.get(f"{{{ns}}}hAnsi")
            if not fam:
                theme = fonts.get(f"{{{ns}}}asciiTheme") or ""
                # Word-Theme-Defaults (Office 2013+): minor = Calibri, major = Calibri Light
                if theme.startswith("minor"):
                    fam = "Calibri"
                elif theme.startswith("major"):
                    fam = "Calibri Light"
        size = None
        sz = rpr.find(f"{{{ns}}}sz")
        if sz is not None:
            try:
                size = float(sz.get(f"{{{ns}}}val")) / 2.0
            except (TypeError, ValueError):
                size = None
        return fam, size
    except Exception:
        return None, None


def docx_default_font(path: str | Path) -> tuple[str | None, float | None]:
    """(Schriftfamilie, Größe pt) des Normal-Stils bzw. der docDefaults — 2.6.53.

    Der Editor nutzt dies als ``QTextDocument.defaultFont`` statt der Editor-
    Monospace-Schrift (Feld: „DOCX-Fließtext sieht aus wie Plaintext/Consolas“).
    """
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return None, None
    try:
        d = DocxDocument(str(path))
    except Exception:
        return None, None
    fam: str | None = None
    size: float | None = None
    try:
        normal = d.styles["Normal"]
        for st in _style_chain(normal):
            if fam is None:
                v = _font_attr(getattr(st, "font", None), "name")
                if v:
                    fam = str(v)
            if size is None:
                sz = _font_attr(getattr(st, "font", None), "size")
                if sz is not None:
                    try:
                        size = float(sz.pt)
                    except Exception:
                        size = None
    except Exception:
        pass
    if fam is None or size is None:
        fam2, size2 = _docx_default_font_from_styles_xml(d)
        fam = fam or fam2
        size = size or size2
    return fam, size


def _effective_color_hex(run: Any) -> str | None:
    try:
        color = run.font.color
    except Exception:
        return None
    try:
        rgb = color.rgb if color is not None else None
    except Exception:
        rgb = None
    if rgb is None:
        return None
    try:
        hexv = str(rgb)
    except Exception:
        return None
    if re.fullmatch(r"[0-9A-Fa-f]{6}", hexv):
        return "#" + hexv.lower()
    return None


def _effective_highlight_hex(run: Any) -> str | None:
    try:
        hl = run.font.highlight_color
    except Exception:
        return None
    if hl is None:
        return None
    name = getattr(hl, "name", None) or str(hl)
    name = str(name).split(".")[-1].upper()
    if name in ("AUTO", "INHERITED", "NONE"):
        return None
    return _HIGHLIGHT_INDEX_TO_HEX.get(name)


def _run_to_html(run: Any, para: Any = None) -> str:
    text = run.text or ""
    if not text:
        return ""
    piece = _escape(text)
    styles: list[str] = []
    try:
        fam = _effective_font_name(run, para)
        if fam and re.fullmatch(r"[\w .+-]{1,60}", fam):
            styles.append(f"font-family:'{fam}'")
    except Exception:
        pass
    try:
        size = _effective_size_pt(run, para)
        if size and 4.0 <= size <= 200.0:
            styles.append(f"font-size:{size:g}pt")
    except Exception:
        pass
    try:
        color = _effective_color_hex(run)
        if color and color != "#000000":
            styles.append(f"color:{color}")
    except Exception:
        pass
    try:
        hl = _effective_highlight_hex(run)
        if hl:
            styles.append(f"background-color:{hl}")
    except Exception:
        pass
    if styles:
        piece = f'<span style="{";".join(styles)};">{piece}</span>'
    # Unterstrich zuerst innen, dann Kursiv/Fett — wie typische Writer
    if _effective_tri(run, para, "underline"):
        piece = f"<u>{piece}</u>"
    if _effective_tri(run, para, "italic"):
        piece = f"<i>{piece}</i>"
    if _effective_tri(run, para, "bold"):
        piece = f"<b>{piece}</b>"
    return piece


def _iter_para_runs(para: Any):
    """Runs inkl. Hyperlink-Runs (``para.runs`` überspringt Hyperlinks!) — 2.6.52.

    Liefert Tupel (run, href|None).
    """
    iter_inner = getattr(para, "iter_inner_content", None)
    if callable(iter_inner):
        try:
            for item in iter_inner():
                if hasattr(item, "runs") and hasattr(item, "address"):
                    href = ""
                    try:
                        href = str(item.address or "")
                    except Exception:
                        href = ""
                    for r in item.runs:
                        yield r, (href or None)
                else:
                    yield item, None
            return
        except Exception as e:  # pragma: no cover - fallback
            _log.debug("iter_inner_content fehlgeschlagen: %s", e)
    for r in para.runs:
        yield r, None


def _para_inner_html(para: Any) -> str:
    bits: list[str] = []
    link_open: str | None = None
    for run, href in _iter_para_runs(para):
        piece = _run_to_html(run, para)
        if not piece:
            continue
        if href != link_open:
            if link_open is not None:
                bits.append("</a>")
            if href:
                bits.append(f'<a href="{html_lib.escape(href, quote=True)}">')
            link_open = href
        bits.append(piece)
    if link_open is not None:
        bits.append("</a>")
    inner = "".join(bits)
    if not inner:
        try:
            txt = para.text or ""
        except Exception:
            txt = ""
        if txt:
            # Fallback wenn Runs leer, Text aber vorhanden
            inner = _escape(txt)
    return inner


def _heading_tag(style_name: str | None) -> str:
    name = (style_name or "").strip()
    m = re.match(r"Heading\s*([1-6])\b", name, re.I)
    if m:
        return f"h{m.group(1)}"
    m = re.match(r"Überschrift\s*([1-6])\b", name, re.I)
    if m:
        return f"h{m.group(1)}"
    if name.lower() in {"title", "titel"}:
        return "h1"
    if name.lower() in {"subtitle", "untertitel"}:
        return "h2"
    return "p"


def _para_align_style(para: Any) -> str:
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        a = para.alignment
        if a is None and para.style is not None:
            try:
                a = para.style.paragraph_format.alignment
            except Exception:
                a = None
        if a == WD_ALIGN_PARAGRAPH.CENTER:
            return ' style="text-align:center;"'
        if a == WD_ALIGN_PARAGRAPH.RIGHT:
            return ' style="text-align:right;"'
        if a == WD_ALIGN_PARAGRAPH.JUSTIFY:
            return ' style="text-align:justify;"'
    except Exception:
        pass
    return ""


def _para_to_html(para: Any) -> str:
    style_name = ""
    try:
        style_name = para.style.name if para.style is not None else ""
    except Exception:
        style_name = ""
    tag = _heading_tag(style_name)
    inner = _para_inner_html(para)
    if not inner:
        inner = "<br/>"
    prefix = _list_prefix_html(para, style_name)
    if prefix and not inner.startswith(prefix):
        inner = prefix + inner
    return f"<{tag}{_para_align_style(para)}>{inner}</{tag}>"


def _list_prefix_html(para: Any, style_name: str) -> str:
    """Word-Listen → sichtbares • / 1. (kein Steuerzeichen)."""
    name = (style_name or "").lower()
    num_pr = None
    try:
        ppr = para._p.pPr
        num_pr = ppr.numPr if ppr is not None else None
    except Exception:
        num_pr = None
    ilvl = 0
    try:
        if num_pr is not None and num_pr.ilvl is not None:
            ilvl = int(num_pr.ilvl.val)
    except Exception:
        ilvl = 0
    is_num = "list number" in name or "nummer" in name
    is_bullet = "list bullet" in name or "aufzähl" in name or "bullet" in name
    if not is_num and not is_bullet and num_pr is None:
        return ""
    if is_num:
        return "1. "
    glyphs = ("• ", "◦ ", "▪ ")
    return glyphs[min(max(0, ilvl), 2)]


def docx_to_html(path: str | Path) -> str:
    """DOCX → HTML mit <b>/<i>/<u>, Größe/Farbe/Textmarker, Links, Absätzen.

    Fehler in einzelnen Absätzen/Tabellen werden geloggt und degradieren nur
    diesen Absatz zu Plaintext — nie das ganze Dokument (2.6.52).
    Dokumentreihenfolge: Absätze und Tabellen wie im Body (``iter_inner_content``).
    """
    try:
        from docx import Document as DocxDocument
    except ImportError as e:
        raise RuntimeError("python-docx fehlt zum DOCX-Import") from e

    d = DocxDocument(str(path))
    parts: list[str] = [
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"/></head><body>"
    ]
    para_errors = 0

    def _table_html(tbl: Any) -> str:
        bits = ["<table border=\"1\" cellpadding=\"4\" cellspacing=\"0\">"]
        for row in tbl.rows:
            bits.append("<tr>")
            for cell in row.cells:
                cell_bits: list[str] = []
                for p in cell.paragraphs:
                    try:
                        cell_bits.append(_para_inner_html(p) or _escape(p.text))
                    except Exception:
                        cell_bits.append(_escape(getattr(p, "text", "") or ""))
                bits.append(f"<td>{'<br/>'.join(cell_bits)}</td>")
            bits.append("</tr>")
        bits.append("</table>")
        return "".join(bits)

    body_items: list[Any] = []
    iter_inner = getattr(d, "iter_inner_content", None)
    if callable(iter_inner):
        try:
            body_items = list(iter_inner())
        except Exception as e:
            _log.debug("Document.iter_inner_content fehlgeschlagen: %s", e)
            body_items = []
    if not body_items:
        body_items = list(d.paragraphs) + list(d.tables)

    for item in body_items:
        try:
            if hasattr(item, "rows") and hasattr(item, "columns"):
                parts.append(_table_html(item))
            else:
                parts.append(_para_to_html(item))
        except Exception as e:
            para_errors += 1
            txt = ""
            try:
                txt = item.text or ""
            except Exception:
                txt = ""
            parts.append(f"<p>{_escape(txt) or '<br/>'}</p>")
            if para_errors <= 3:
                _log.warning("DOCX-Absatz nur als Text übernommen (%s): %s", path, e)

    parts.append("</body></html>")
    if para_errors:
        _log.warning("DOCX %s: %d Absatz/Tabelle ohne Formatierung übernommen", path, para_errors)
    return "".join(parts)


def docx_plain_and_html(path: str | Path) -> tuple[str, str]:
    """Liefert (Plaintext, HTML) aus einer DOCX-Datei."""
    html = docx_to_html(path)
    # Plain: einfache Tag-Entfernung
    plain = re.sub(r"(?i)<br\s*/?>", "\n", html)
    plain = re.sub(r"(?i)</p\s*>", "\n", plain)
    plain = re.sub(r"(?i)</h[1-6]\s*>", "\n", plain)
    plain = re.sub(r"(?i)</tr\s*>", "\n", plain)
    plain = re.sub(r"(?i)</td\s*>", "\t", plain)
    plain = re.sub(r"<[^>]+>", "", plain)
    plain = html_lib.unescape(plain)
    plain = re.sub(r"\n{3,}", "\n\n", plain).strip() + ("\n" if plain.strip() else "")
    return plain, html


class _HtmlToDocxParser(HTMLParser):
    """Minimaler HTML→DOCX-Parser für Qt-toHtml und docx_to_html-Ausgabe."""

    def __init__(self, document: Any) -> None:
        super().__init__(convert_charrefs=True)
        self.doc = document
        self._para: Any | None = None
        self._bold = 0
        self._italic = 0
        self._underline = 0
        self._skip = 0
        self._pending_align: str | None = None
        self._heading_level: int | None = None
        self._span_stack: list[tuple[bool, bool, bool]] = []
        # Farbe / Größe / Textmarker aus Qt-Spans (innerster Wert gewinnt) — 2.6.52
        self._style_stack: list[dict[str, Any]] = []
        self._cur_table: list[list[str]] | None = None
        self._cur_row: list[str] | None = None
        self._cell_buf: list[str] | None = None
        self._in_cell = False
        self._th_bold = False

    def _ensure_para(self) -> Any:
        if self._para is None:
            if self._heading_level:
                self._para = self.doc.add_heading("", level=self._heading_level)
            else:
                self._para = self.doc.add_paragraph("")
            if self._pending_align:
                try:
                    from docx.enum.text import WD_ALIGN_PARAGRAPH

                    mapping = {
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                    }
                    al = mapping.get(self._pending_align)
                    if al is not None:
                        self._para.alignment = al
                except Exception:
                    pass
        return self._para

    def _current_extra(self) -> dict[str, Any]:
        merged: dict[str, Any] = {}
        for st in self._style_stack:
            for k, v in st.items():
                if v is not None:
                    merged[k] = v
        return merged

    def _flush_docx_table(self, rows: list[list[str]]) -> None:
        if not rows:
            return
        cols = max((len(r) for r in rows), default=0)
        if cols < 1:
            return
        try:
            table = self.doc.add_table(rows=len(rows), cols=cols)
            try:
                table.style = "Table Grid"
            except Exception:
                pass
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    val = row[ci] if ci < len(row) else ""
                    table.cell(ri, ci).text = val
        except Exception:
            for row in rows:
                self._para = None
                para = self._ensure_para()
                para.add_run(" | ".join(row))
                self._para = None
        self._para = None

    def _add_text(self, text: str) -> None:
        if not text:
            return
        if self._in_cell and self._cell_buf is not None:
            cleaned = (
                str(text)
                .replace("\x0c", " ")
                .replace("\u00b6", " ")
                .replace("\ufffd", " ")
            )
            self._cell_buf.append(cleaned)
            return
        if self._cur_table is not None:
            return
        para = self._ensure_para()
        run = para.add_run(text)
        run.bold = self._bold > 0
        run.italic = self._italic > 0
        run.underline = self._underline > 0
        extra = self._current_extra()
        try:
            size = extra.get("size")
            if size:
                from docx.shared import Pt

                run.font.size = Pt(float(size))
        except Exception:
            pass
        try:
            color = extra.get("color")
            if color:
                from docx.shared import RGBColor

                run.font.color.rgb = RGBColor.from_string(str(color).lstrip("#").upper())
        except Exception:
            pass
        try:
            bg = extra.get("background")
            if bg:
                run.font.highlight_color = _closest_highlight_index(str(bg))
        except Exception:
            pass
        try:
            fam = extra.get("font")
            if fam:
                run.font.name = str(fam)
        except Exception:
            pass

    def _style_flags(self, style: str) -> tuple[bool, bool, bool]:
        s = (style or "").lower().replace(" ", "")
        bold = (
            "font-weight:600" in s
            or "font-weight:700" in s
            or "font-weight:800" in s
            or "font-weight:900" in s
            or "font-weight:bold" in s
        )
        italic = "font-style:italic" in s
        underline = "text-decoration:underline" in s or "text-decoration-line:underline" in s
        return bold, italic, underline

    @staticmethod
    def _style_extra(style: str) -> dict[str, Any]:
        """color / font-size / background-color aus einem style-Attribut."""
        s = (style or "").lower()
        out: dict[str, Any] = {"color": None, "size": None, "background": None, "font": None}
        # Schriftfamilie (erste Familie, ohne generische Fallbacks) — 2.6.53
        m = re.search(r"font-family\s*:\s*([^;]+)", style or "", re.IGNORECASE)
        if m:
            first = m.group(1).split(",")[0].strip().strip("'\"")
            if first and first.lower() not in {
                "serif", "sans-serif", "monospace", "cursive", "fantasy", "system-ui",
            }:
                out["font"] = first
        m = re.search(r"(?<![\w-])color\s*:\s*(#[0-9a-f]{6}|#[0-9a-f]{3})", s)
        if m:
            out["color"] = _expand_hex(m.group(1))
        m = re.search(r"background(?:-color)?\s*:\s*(#[0-9a-f]{6}|#[0-9a-f]{3})", s)
        if m:
            out["background"] = _expand_hex(m.group(1))
        m = re.search(r"font-size\s*:\s*([0-9]+(?:\.[0-9]+)?)\s*(pt|px)", s)
        if m:
            val = float(m.group(1))
            if m.group(2) == "px":
                val = val * 0.75
            if 4.0 <= val <= 200.0:
                out["size"] = val
        return out

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        t = tag.lower()
        ad = {k.lower(): (v or "") for k, v in attrs}
        if t in {"script", "style", "head", "meta", "title"}:
            self._skip += 1
            return
        if self._skip:
            return
        if t in {"b", "strong"}:
            self._bold += 1
        elif t in {"i", "em"}:
            self._italic += 1
        elif t == "u":
            self._underline += 1
        elif t in {"span", "font", "a"}:
            b, i, u = self._style_flags(ad.get("style", ""))
            if b:
                self._bold += 1
            if i:
                self._italic += 1
            if u:
                self._underline += 1
            self._span_stack.append((b, i, u))
            extra = self._style_extra(ad.get("style", ""))
            if t == "font" and ad.get("color"):
                extra["color"] = _expand_hex(ad.get("color", ""))
            self._style_stack.append(extra)
        elif t in {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6"}:
            if self._cur_table is not None:
                return
            self._para = None
            self._pending_align = None
            style = ad.get("style", "")
            if "text-align:center" in style.replace(" ", "").lower():
                self._pending_align = "center"
            elif "text-align:right" in style.replace(" ", "").lower():
                self._pending_align = "right"
            elif "text-align:justify" in style.replace(" ", "").lower():
                self._pending_align = "justify"
            if t.startswith("h") and t[1:].isdigit():
                self._heading_level = int(t[1])
            else:
                self._heading_level = None
        elif t == "br":
            if self._in_cell:
                self._add_text("\n")
            else:
                self._add_text("\n")
        elif t == "li":
            if self._cur_table is not None:
                return
            self._para = None
            self._heading_level = None
            self._pending_align = None
        elif t == "table":
            self._para = None
            self._cur_table = []
            self._cur_row = None
            self._cell_buf = None
            self._in_cell = False
        elif t == "tr":
            if self._cur_table is not None:
                self._cur_row = []
        elif t in {"td", "th"}:
            if self._cur_row is not None:
                self._cell_buf = []
                self._in_cell = True
                if t == "th":
                    self._bold += 1
                    self._th_bold = True
        elif t in {"thead", "tbody", "tfoot", "colgroup", "col"}:
            return

    def handle_endtag(self, tag: str) -> None:
        t = tag.lower()
        if t in {"script", "style", "head", "meta", "title"}:
            if self._skip:
                self._skip -= 1
            return
        if self._skip:
            return
        if t in {"b", "strong"}:
            self._bold = max(0, self._bold - 1)
        elif t in {"i", "em"}:
            self._italic = max(0, self._italic - 1)
        elif t == "u":
            self._underline = max(0, self._underline - 1)
        elif t in {"span", "font", "a"}:
            stack = self._span_stack
            if stack:
                b, i, u = stack.pop()
                if b:
                    self._bold = max(0, self._bold - 1)
                if i:
                    self._italic = max(0, self._italic - 1)
                if u:
                    self._underline = max(0, self._underline - 1)
            if self._style_stack:
                self._style_stack.pop()
        elif t in {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
            if self._cur_table is not None:
                return
            if self._para is None:
                # leerer Absatz
                self._ensure_para()
            self._para = None
            self._heading_level = None
            self._pending_align = None
        elif t in {"td", "th"}:
            if self._cur_row is not None:
                text = "".join(self._cell_buf or []).strip()
                self._cur_row.append(text)
                self._cell_buf = None
                self._in_cell = False
            if self._th_bold:
                self._bold = max(0, self._bold - 1)
                self._th_bold = False
            self._para = None
        elif t == "tr":
            if self._cur_table is not None and self._cur_row is not None:
                self._cur_table.append(self._cur_row)
            self._cur_row = None
            self._para = None
        elif t == "table":
            rows = self._cur_table or []
            self._cur_table = None
            self._cur_row = None
            self._cell_buf = None
            self._in_cell = False
            self._flush_docx_table(rows)
            self._para = None
        elif t in {"thead", "tbody", "tfoot", "colgroup", "col"}:
            return

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if not data:
            return
        # Qt liefert oft Newlines zwischen Tags — nur echte Textstücke schreiben
        if data.strip() == "" and "\n" in data and data.replace("\n", "").strip() == "":
            return
        self._add_text(data)


def html_to_docx(
    html: str,
    path: str | Path,
    *,
    title: Optional[str] = None,
    header: Optional[str] = None,
    footer: Optional[str] = None,
    author: Optional[str] = None,
) -> Path:
    """HTML (inkl. Qt-Spans) → DOCX mit Bold/Italic/Underline."""
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt
    except ImportError as e:
        raise RuntimeError("python-docx fehlt zum DOCX-Export") from e

    path = Path(path)
    d = DocxDocument()
    if title:
        try:
            d.core_properties.title = title
        except Exception:
            pass
    if author:
        try:
            d.core_properties.author = str(author)
        except Exception:
            pass
    try:
        style = d.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
    except Exception:
        pass

    parser = _HtmlToDocxParser(d)
    parser.feed(html or "")
    parser.close()
    if parser._para is None and not d.paragraphs:
        d.add_paragraph("")
    _apply_docx_header_footer(d, header, footer, html=html)
    d.save(str(path))
    return path


def _clean_hf_text(value: str | None) -> str:
    s = str(value or "")
    return (
        s.replace("\x0c", " ")
        .replace("\u00b6", " ")
        .replace("\u2028", " ")
        .replace("\u2029", " ")
        .strip()
    )


def _apply_docx_header_footer(
    d,
    header: str | None,
    footer: str | None,
    *,
    html: str = "",
) -> None:
    """Kopf-/Fußzeile in DOCX-Section, nicht in den Body (keine ¶/Form-Feed)."""
    h = _clean_hf_text(header)
    f = _clean_hf_text(footer)
    if (not h or not f) and html:
        for m in re.finditer(r"<!--\s*ild-(header|footer)\s+(.*?)\s*-->", html or "", re.I | re.S):
            kind = (m.group(1) or "").lower()
            text = _clean_hf_text(html_lib.unescape(m.group(2) or ""))
            if kind == "header" and not h:
                h = text
            elif kind == "footer" and not f:
                f = text
    if not h and not f:
        return
    try:
        section = d.sections[0]
    except Exception:
        return
    if h:
        try:
            hp = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
            hp.text = h
        except Exception:
            pass
    if f:
        try:
            fp = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
            fp.text = f
        except Exception:
            pass


def docx_header_footer_author(path: str | Path) -> tuple[str, str, str]:
    """Kopf-/Fußzeile und Autor aus DOCX-Section/Core-Properties."""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return "", "", ""
    try:
        d = DocxDocument(str(path))
    except Exception:
        return "", "", ""
    header = ""
    footer = ""
    author = ""
    try:
        section = d.sections[0]
        header = _clean_hf_text("\n".join(p.text for p in section.header.paragraphs))
        footer = _clean_hf_text("\n".join(p.text for p in section.footer.paragraphs))
    except Exception:
        pass
    try:
        author = str(getattr(d.core_properties, "author", "") or "").strip()
    except Exception:
        author = ""
    return header, footer, author


def html_to_rtf(html: str, path: str | Path, *, title: Optional[str] = None) -> Path:
    """HTML → einfaches RTF mit \\b/\\i/\\ul."""

    class _Rtf(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.parts: list[str] = []
            self.bold = 0
            self.italic = 0
            self.underline = 0
            self.skip = 0
            self._span_stack: list[tuple[bool, bool, bool]] = []

        def _esc(self, s: str) -> str:
            out: list[str] = []
            for ch in s:
                o = ord(ch)
                if ch in {"\\", "{", "}"}:
                    out.append("\\" + ch)
                elif ch == "\n":
                    out.append("\\par\n")
                elif o < 128:
                    out.append(ch)
                else:
                    out.append(f"\\u{o}?")
            return "".join(out)

        def _style_flags(self, style: str) -> tuple[bool, bool, bool]:
            s = (style or "").lower().replace(" ", "")
            bold = "font-weight:600" in s or "font-weight:700" in s or "font-weight:bold" in s
            italic = "font-style:italic" in s
            underline = "text-decoration:underline" in s
            return bold, italic, underline

        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            t = tag.lower()
            ad = {k.lower(): (v or "") for k, v in attrs}
            if t in {"script", "style", "head"}:
                self.skip += 1
                return
            if self.skip:
                return
            if t in {"b", "strong"}:
                self.bold += 1
                self.parts.append("\\b ")
            elif t in {"i", "em"}:
                self.italic += 1
                self.parts.append("\\i ")
            elif t == "u":
                self.underline += 1
                self.parts.append("\\ul ")
            elif t == "span":
                b, i, u = self._style_flags(ad.get("style", ""))
                self._span_stack.append((b, i, u))
                if b:
                    self.bold += 1
                    self.parts.append("\\b ")
                if i:
                    self.italic += 1
                    self.parts.append("\\i ")
                if u:
                    self.underline += 1
                    self.parts.append("\\ul ")
            elif t == "br":
                self.parts.append("\\par\n")
            elif t in {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li"}:
                if self.parts and not self.parts[-1].endswith("\\par\n"):
                    pass

        def handle_endtag(self, tag: str) -> None:
            t = tag.lower()
            if t in {"script", "style", "head"}:
                if self.skip:
                    self.skip -= 1
                return
            if self.skip:
                return
            if t in {"b", "strong"}:
                self.bold = max(0, self.bold - 1)
                self.parts.append("\\b0 ")
            elif t in {"i", "em"}:
                self.italic = max(0, self.italic - 1)
                self.parts.append("\\i0 ")
            elif t == "u":
                self.underline = max(0, self.underline - 1)
                self.parts.append("\\ul0 ")
            elif t == "span":
                if self._span_stack:
                    b, i, u = self._span_stack.pop()
                    if u:
                        self.underline = max(0, self.underline - 1)
                        self.parts.append("\\ul0 ")
                    if i:
                        self.italic = max(0, self.italic - 1)
                        self.parts.append("\\i0 ")
                    if b:
                        self.bold = max(0, self.bold - 1)
                        self.parts.append("\\b0 ")
            elif t in {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr"}:
                self.parts.append("\\par\n")

        def handle_data(self, data: str) -> None:
            if self.skip or not data:
                return
            if data.strip() == "" and data.replace("\n", "").strip() == "":
                return
            self.parts.append(self._esc(data))

    path = Path(path)
    header = [
        r"{\rtf1\ansi\deff0",
        r"{\fonttbl{\f0 Times New Roman;}{\f1 Segoe UI;}}",
    ]
    if title:
        from instantlensdoc.core.export import _rtf_escape

        header.append(r"{\info{\title " + _rtf_escape(title) + "}}")
    header.append(r"\f0\fs22 ")
    parser = _Rtf()
    parser.feed(html or "")
    parser.close()
    path.write_text("".join(header) + "".join(parser.parts) + "}", encoding="utf-8", errors="replace")
    return path


def html_has_char_formats(html: str) -> bool:
    """True wenn HTML Bold/Italic/Underline-Marker enthält."""
    h = html or ""
    if re.search(r"(?i)</?(b|strong|i|em|u)\b", h):
        return True
    if re.search(r"(?i)font-weight\s*:\s*(bold|[6-9]00)", h):
        return True
    if re.search(r"(?i)font-style\s*:\s*italic", h):
        return True
    if re.search(r"(?i)text-decoration[^;]*underline", h):
        return True
    if re.search(r"(?i)background(-color)?\s*:\s*#", h):
        return True
    return False
