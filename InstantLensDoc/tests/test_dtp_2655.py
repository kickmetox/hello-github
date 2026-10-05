"""Nutzbarer DTP-Layout-Modus: Caret, Resize, Flow, Import, Master, Ebenen — 2.6.54."""

from __future__ import annotations

from pathlib import Path

from instantlensdoc.dtp.import_text import read_import_text
from instantlensdoc.dtp.model import DtpDocument, DtpFrame
from instantlensdoc.dtp.presets import apply_book_preset


def test_inline_edit_and_resize(qapp):
    from instantlensdoc.dtp.canvas import DtpPane

    doc = DtpDocument()
    fr = doc.add_text_frame("Hallo", x=40, y=40, width=120, height=50)
    pane = DtpPane(doc=doc)
    pane.resize(900, 700)
    pane.show()
    qapp.processEvents()
    item = pane.scene._items[fr.id]
    item.begin_edit()
    item.set_plain_text("Inline-Caret-Text")
    item.end_edit()
    qapp.processEvents()
    assert "Inline-Caret-Text" in (fr.text or "") or doc.story_text(fr).startswith("Inline")
    orig_w = fr.width
    item.resize_from("br", 80, 40, (fr.x, fr.y, fr.width, fr.height))
    assert fr.width > orig_w
    assert abs(item.rect().width() - fr.width) < 0.5
    pane.close()


def test_link_reflows_overflow():
    doc = DtpDocument()
    a = doc.add_text_frame("", x=40, y=40, width=90, height=36)
    b = doc.add_text_frame("", x=40, y=90, width=90, height=80)
    doc.link_frames(a.id, b.id)
    a.text = ("Wort " * 60).strip()
    result = doc.reflow_chain(a, auto_extend=False)
    assert a.id in result and b.id in result
    assert len(b.text) > 0
    assert a.next_id == b.id
    # Overflow-Fortsetzung auf neuer Seite
    tiny = DtpDocument()
    apply_book_preset(tiny, "A6")
    t = tiny.add_text_frame("X " * 400, x=20, y=20, width=70, height=28)
    tiny.reflow_chain(t, auto_extend=True)
    assert tiny.page_count >= 2
    assert tiny.chain_members(tiny.chain_head(t))[-1].page >= 1


def test_guide_drag_persists(qapp):
    from instantlensdoc.dtp.canvas import DtpPane, GuideItem

    doc = DtpDocument()
    g = doc.add_guide("vertical", 80.0)
    pane = DtpPane(doc=doc)
    pane.show()
    qapp.processEvents()
    gi = next(it for it in pane.scene.items() if isinstance(it, GuideItem))
    gi.setPos(120 + 40, 0)  # PAGE_OFFSET = 40
    qapp.processEvents()
    assert abs(g.position_pt - 120) < 1.5
    pane.close()


def test_book_preset_sets_margins_and_master():
    doc = DtpDocument()
    apply_book_preset(doc, "Taschenbuch")
    assert doc.geometry.name == "Taschenbuch"
    assert doc.geometry.margin_left_pt > 10
    assert doc.geometry.columns >= 1
    m = doc.add_master("Kapitel", header="Kap.", footer="{n}")
    doc.add_page()
    doc.apply_master(m.id, pages=[1])
    assert doc.page_master[1] == m.id
    st = doc.apply_style(doc.add_text_frame("H1", style_id="body").id, "h1")
    assert st.style_id == "h1"
    assert st.font_size >= 14


def test_import_txt_docx_ocr(tmp_path: Path):
    txt = tmp_path / "in.txt"
    txt.write_text("Erste Zeile.\nZweite Zeile aus OCR.", encoding="utf-8")
    ocr = tmp_path / "scan.ildocr.txt"
    ocr.write_text("Erkannter Text vom Scanner.", encoding="utf-8")
    docx_path = tmp_path / "t.docx"
    from docx import Document as Dx

    d = Dx()
    d.add_paragraph("DOCX-Absatz im Rahmen")
    d.save(str(docx_path))
    assert "Erste Zeile" in read_import_text(txt)
    assert "Scanner" in read_import_text(ocr)
    assert "DOCX-Absatz" in read_import_text(docx_path)
    doc = DtpDocument()
    fr = doc.add_text_frame("", width=200, height=120)
    doc.import_into_frame(fr.id, str(txt), auto_extend=False)
    assert "Erste Zeile" in (fr.text or "")


def test_image_frame_replace_keeps_kind(tmp_path: Path):
    png = tmp_path / "p.png"
    png.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00"
        b"\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    doc = DtpDocument()
    fr = doc.add_image_frame("", x=10, y=10, width=80, height=60)
    assert fr.kind == "image"
    doc.set_image(fr.id, str(png))
    assert fr.kind == "image" and fr.image_path.endswith("p.png")


def test_layers_toggle_and_align():
    from instantlensdoc.dtp.geometry import align_frames

    doc = DtpDocument()
    a = doc.add_shape("rectangle", x=10, y=10, width=20, height=10)
    b = doc.add_shape("rectangle", x=50, y=30, width=20, height=10)
    a.layer_id = "images"
    doc.set_layer_visible("images", False)
    vis = [f.id for f in doc.sorted_frames(0)]
    assert a.id not in vis
    doc.set_layer_visible("images", True)
    align_frames([a, b], "left")
    assert a.x == b.x


def test_export_pdf_header_valid(qapp, tmp_path: Path):
    from instantlensdoc.dtp.export import export_pdf

    doc = DtpDocument.sample("A5")
    dest = tmp_path / "layout.pdf"
    export_pdf(doc, dest)
    raw = dest.read_bytes()
    assert raw.startswith(b"%PDF")
    import pikepdf

    with pikepdf.open(dest) as pdf:
        assert len(pdf.pages) >= 1


def test_pane_layers_and_import(qapp, tmp_path: Path):
    from instantlensdoc.dtp.canvas import DtpPane

    src = tmp_path / "body.txt"
    src.write_text("Importierter Fliesstext " * 20, encoding="utf-8")
    pane = DtpPane(doc=DtpDocument.sample("A5"))
    pane.resize(1200, 800)
    pane.show()
    qapp.processEvents()
    assert pane.layer_list.count() >= 3
    pane.import_text(str(src))
    qapp.processEvents()
    joined = " ".join(f.text for f in pane.doc.frames if f.kind == "text")
    assert "Importierter" in joined
    pane.close()
