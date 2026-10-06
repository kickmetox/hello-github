"""Dokument-Tabellen 2.6.14: erstellen, formatieren, sortieren; CSV/Excel-Import.

Markdown-Pipe-Tabellen mit optionalen ``ild-table``-Markern im Editor-Text.
Kein CMYK/Bleed/PDF/X, kein volles i18n, keine KI-Wizards, kein PDF-Compare.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, List, Optional, Sequence, Union
from xml.etree import ElementTree as ET

PathLike = Union[str, Path]

_TABLE_OPEN_RE = re.compile(
    r"^<!--\s*ild-table:(?P<meta>[^>]*)\s*-->\s*$",
    re.IGNORECASE,
)
_TABLE_CLOSE_RE = re.compile(r"^<!--\s*/ild-table\s*-->\s*$", re.IGNORECASE)
_ALIGN_MAP = {"l": "left", "c": "center", "r": "right", "left": "l", "center": "c", "right": "r"}


@dataclass
class TableFormat:
    """Darstellungshinweise für eine Tabelle."""

    header: bool = True
    border: bool = True
    align: str = ""  # z. B. "lcr" je Spalte
    header_bold: bool = True
    style: str = "default"  # default|striped|compact

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    def align_chars(self, cols: int) -> str:
        chars = (self.align or "").lower().replace("left", "l").replace("center", "c").replace("right", "r")
        chars = re.sub(r"[^lcr]", "", chars)
        if not chars:
            return "l" * max(1, cols)
        if len(chars) < cols:
            chars = chars + chars[-1] * (cols - len(chars))
        return chars[:cols]


@dataclass
class DocumentTable:
    """Einfache Dokument-Tabelle (Zeilen × Spalten als Strings)."""

    cells: list[list[str]] = field(default_factory=list)
    format: TableFormat = field(default_factory=TableFormat)
    table_id: str = "t1"

    @property
    def rows(self) -> int:
        return len(self.cells)

    @property
    def cols(self) -> int:
        if not self.cells:
            return 0
        return max(len(r) for r in self.cells)

    def normalized(self) -> "DocumentTable":
        cols = self.cols
        out: list[list[str]] = []
        for row in self.cells:
            r = [str(c) if c is not None else "" for c in row]
            if len(r) < cols:
                r = r + [""] * (cols - len(r))
            out.append(r[:cols])
        return DocumentTable(cells=out, format=self.format, table_id=self.table_id)

    def to_dict(self) -> dict[str, Any]:
        t = self.normalized()
        return {
            "id": t.table_id,
            "rows": t.rows,
            "cols": t.cols,
            "cells": t.cells,
            "format": t.format.to_dict(),
        }

    def header_row(self) -> list[str]:
        if self.format.header and self.cells:
            return list(self.cells[0])
        return [f"Spalte {i + 1}" for i in range(self.cols)]

    def data_rows(self) -> list[list[str]]:
        if self.format.header and self.cells:
            return [list(r) for r in self.cells[1:]]
        return [list(r) for r in self.cells]


def create_table(
    rows: int = 3,
    cols: int = 3,
    *,
    header: bool = True,
    data: Sequence[Sequence[Any]] | None = None,
    table_id: str = "t1",
    align: str = "",
    border: bool = True,
    style: str = "default",
) -> DocumentTable:
    """Leere oder befüllte Tabelle anlegen."""
    if data is not None:
        cells = [[("" if c is None else str(c)) for c in row] for row in data]
        if not cells:
            cells = [["" for _ in range(max(1, cols))]]
    else:
        r = max(1, int(rows))
        c = max(1, int(cols))
        cells = [["" for _ in range(c)] for _ in range(r)]
        if header:
            cells[0] = [f"Spalte {i + 1}" for i in range(c)]
    fmt = TableFormat(
        header=bool(header),
        border=bool(border),
        align=str(align or ""),
        style=str(style or "default"),
    )
    return DocumentTable(cells=cells, format=fmt, table_id=str(table_id or "t1")).normalized()


def format_table(
    table: DocumentTable | str,
    *,
    align: str | None = None,
    border: bool | None = None,
    header: bool | None = None,
    header_bold: bool | None = None,
    style: str | None = None,
) -> DocumentTable:
    """Format-Attribute setzen (Align je Spalte, Rahmen, Header, Style)."""
    t = parse_table(table) if isinstance(table, str) else table.normalized()
    if align is not None:
        t.format.align = str(align)
    if border is not None:
        t.format.border = bool(border)
    if header is not None:
        t.format.header = bool(header)
    if header_bold is not None:
        t.format.header_bold = bool(header_bold)
    if style is not None:
        t.format.style = str(style)
    return t


def _sort_key(value: str, *, numeric: bool):
    s = (value or "").strip()
    if numeric:
        try:
            return (0, float(s.replace(",", ".").replace(" ", "")))
        except ValueError:
            pass
    return (1, s.casefold())


def sort_table(
    table: DocumentTable | str,
    column: int = 0,
    *,
    reverse: bool = False,
    numeric: bool = True,
) -> DocumentTable:
    """Tabelle nach Spaltenindex sortieren (0-basiert); Header bleibt oben."""
    t = parse_table(table) if isinstance(table, str) else table.normalized()
    col = int(column)
    if col < 0 or col >= t.cols:
        raise IndexError(f"Spalte außerhalb: {col} (cols={t.cols})")
    header = t.cells[0] if t.format.header and t.cells else None
    body = t.data_rows()
    body.sort(key=lambda row: _sort_key(row[col] if col < len(row) else "", numeric=numeric), reverse=reverse)
    cells = ([header] + body) if header is not None else body
    return DocumentTable(cells=cells, format=t.format, table_id=t.table_id)


def table_to_markdown(table: DocumentTable, *, with_markers: bool = True) -> str:
    """Markdown-Pipe-Tabelle (+ optionale ild-table Marker)."""
    t = table.normalized()
    cols = t.cols
    if cols == 0:
        return ""
    align = t.format.align_chars(cols)
    sep_parts = []
    for ch in align:
        if ch == "c":
            sep_parts.append(":---:")
        elif ch == "r":
            sep_parts.append("---:")
        else:
            sep_parts.append("---")

    def row_md(cells: list[str]) -> str:
        return "| " + " | ".join((c or "").replace("|", "\\|") for c in cells) + " |"

    lines: list[str] = []
    if with_markers:
        meta = (
            f"id={t.table_id};header={1 if t.format.header else 0};"
            f"align={align};border={1 if t.format.border else 0};"
            f"style={t.format.style}"
        )
        lines.append(f"<!-- ild-table:{meta} -->")
    if t.cells:
        lines.append(row_md(t.cells[0]))
        lines.append("| " + " | ".join(sep_parts) + " |")
        for row in t.cells[1:]:
            lines.append(row_md(row))
    if with_markers:
        lines.append("<!-- /ild-table -->")
    return "\n".join(lines)


def table_to_html(table: DocumentTable) -> str:
    t = table.normalized()
    border = "1" if t.format.border else "0"
    align = t.format.align_chars(t.cols)
    parts = [f'<table border="{border}" data-ild-table="{t.table_id}" data-style="{t.format.style}">']
    body = t.cells
    if t.format.header and body:
        parts.append("<thead><tr>")
        for i, c in enumerate(body[0]):
            a = {"l": "left", "c": "center", "r": "right"}.get(align[i], "left")
            weight = "font-weight:bold;" if t.format.header_bold else ""
            parts.append(f'<th style="text-align:{a};{weight}">{_escape_html(c)}</th>')
        parts.append("</tr></thead>")
        body = body[1:]
    parts.append("<tbody>")
    for ri, row in enumerate(body):
        cls = ' class="stripe"' if t.format.style == "striped" and ri % 2 else ""
        parts.append(f"<tr{cls}>")
        for i, c in enumerate(row):
            a = {"l": "left", "c": "center", "r": "right"}.get(align[i], "left")
            parts.append(f'<td style="text-align:{a}">{_escape_html(c)}</td>')
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def _escape_html(s: str) -> str:
    return (
        (s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _parse_meta(meta: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for part in (meta or "").split(";"):
        part = part.strip()
        if not part or "=" not in part:
            continue
        k, v = part.split("=", 1)
        out[k.strip().lower()] = v.strip()
    return out


def _parse_md_row(line: str) -> list[str] | None:
    s = line.strip()
    if not s.startswith("|"):
        return None
    # Separator |---|
    inner = s.strip("|")
    cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", inner)]
    if cells and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells if c != ""):
        return None  # separator row → skip via caller
    if all(re.fullmatch(r":?-+:?", (c or "").replace(" ", "")) for c in cells) and cells:
        return None
    return cells


def _is_sep_row(line: str) -> bool:
    s = line.strip().strip("|")
    if not s:
        return False
    parts = [p.strip() for p in s.split("|")]
    return bool(parts) and all(re.fullmatch(r":?-+:?", p.replace(" ", "") or "-") for p in parts)


def parse_markdown_table(text: str) -> DocumentTable:
    """Eine Markdown-Tabelle (optional mit Markern) parsen."""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    meta: dict[str, str] = {}
    body_lines: list[str] = []
    in_marked = False
    for line in lines:
        m = _TABLE_OPEN_RE.match(line.strip())
        if m:
            meta = _parse_meta(m.group("meta"))
            in_marked = True
            body_lines = []
            continue
        if _TABLE_CLOSE_RE.match(line.strip()):
            break
        if in_marked or line.strip().startswith("|"):
            body_lines.append(line)
            if not in_marked and line.strip().startswith("|"):
                in_marked = True
    cells: list[list[str]] = []
    align_from_sep = ""
    for line in body_lines:
        if _is_sep_row(line):
            parts = [p.strip() for p in line.strip().strip("|").split("|")]
            chars = []
            for p in parts:
                left = p.startswith(":")
                right = p.endswith(":")
                if left and right:
                    chars.append("c")
                elif right:
                    chars.append("r")
                else:
                    chars.append("l")
            align_from_sep = "".join(chars)
            continue
        row = _parse_md_row(line)
        if row is not None:
            cells.append(row)
    header = meta.get("header", "1") not in ("0", "false", "no")
    border = meta.get("border", "1") not in ("0", "false", "no")
    align = meta.get("align") or align_from_sep
    fmt = TableFormat(
        header=header,
        border=border,
        align=align,
        style=meta.get("style") or "default",
    )
    return DocumentTable(
        cells=cells,
        format=fmt,
        table_id=meta.get("id") or "t1",
    ).normalized()


def parse_table(text: str) -> DocumentTable:
    """Alias: Markdown-Tabelle aus Text extrahieren/parsen."""
    return parse_markdown_table(text)


def find_tables_in_text(text: str) -> list[tuple[int, int, DocumentTable]]:
    """Alle ild-table-Blöcke bzw. Pipe-Tabellen finden: (start, end, table)."""
    raw = text or ""
    lines = raw.replace("\r\n", "\n").split("\n")
    # Offsets je Zeilenanfang
    starts = []
    pos = 0
    for line in lines:
        starts.append(pos)
        pos += len(line) + 1
    results: list[tuple[int, int, DocumentTable]] = []
    i = 0
    while i < len(lines):
        if _TABLE_OPEN_RE.match(lines[i].strip()):
            j = i + 1
            while j < len(lines) and not _TABLE_CLOSE_RE.match(lines[j].strip()):
                j += 1
            end_line = j if j < len(lines) else j - 1
            block = "\n".join(lines[i : end_line + 1])
            t = parse_markdown_table(block)
            start_off = starts[i]
            end_off = starts[end_line] + len(lines[end_line]) if end_line < len(lines) else len(raw)
            results.append((start_off, end_off, t))
            i = end_line + 1
            continue
        if lines[i].strip().startswith("|"):
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            block = "\n".join(lines[i:j])
            t = parse_markdown_table(block)
            if t.rows >= 1:
                start_off = starts[i]
                end_off = starts[j - 1] + len(lines[j - 1])
                results.append((start_off, end_off, t))
            i = j
            continue
        i += 1
    return results


def insert_table_into_text(
    text: str,
    table: DocumentTable,
    *,
    at: int | None = None,
    replace_index: int | None = None,
) -> str:
    """Tabelle am Cursor (Zeichenoffset) einfügen oder bestehende ersetzen."""
    md = table_to_markdown(table, with_markers=True)
    if replace_index is not None:
        found = find_tables_in_text(text)
        if replace_index < 0 or replace_index >= len(found):
            raise IndexError(f"Tabelle #{replace_index} nicht gefunden")
        a, b, _ = found[replace_index]
        return (text or "")[:a] + md + (text or "")[b:]
    src = text or ""
    if at is None or at < 0 or at > len(src):
        if src and not src.endswith("\n"):
            src += "\n"
        return src + md + "\n"
    before = src[:at]
    after = src[at:]
    pad_before = "" if before.endswith("\n") or before == "" else "\n"
    pad_after = "" if after.startswith("\n") or after == "" else "\n"
    return before + pad_before + md + pad_after + after


def sniff_csv_delimiter(sample: str) -> str:
    """Trennzeichen erkennen: Komma, Semikolon, Tab (wie Seriendruck)."""
    text = sample or ""
    try:
        from instantlensdoc.core.mail_merge import sniff_csv_delimiter as _mm_sniff

        return _mm_sniff(text)
    except Exception:
        pass
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        if dialect.delimiter in ",;\t|":
            return dialect.delimiter
    except csv.Error:
        pass
    counts = {",": text.count(","), ";": text.count(";"), "\t": text.count("\t")}
    return max(counts, key=counts.get) if any(counts.values()) else ","


def import_csv(
    path: PathLike,
    *,
    delimiter: str | None = None,
    encoding: str | None = "auto",
    header: bool = True,
    table_id: str = "csv1",
) -> DocumentTable:
    """CSV in DocumentTable laden (Delimiter- + Encoding-Autodetect)."""
    p = Path(path)
    enc = encoding
    if enc in (None, "", "auto", "detect"):
        try:
            from instantlensdoc.core.documents import detect_file_encoding

            enc = detect_file_encoding(p)
        except Exception:
            enc = "utf-8"
    if str(enc).lower() in ("utf-8", "utf8"):
        enc = "utf-8-sig"
    raw = p.read_text(encoding=enc, errors="replace")
    sample = raw[:4096]
    if delimiter is None:
        delimiter = sniff_csv_delimiter(sample)
    reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
    cells = [[str(c) for c in row] for row in reader if any(str(x).strip() for x in row) or row]
    if not cells:
        cells = [[""]]
    return create_table(data=cells, header=header, table_id=table_id)


def _xlsx_shared_strings(z: zipfile.ZipFile) -> list[str]:
    try:
        data = z.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(data)
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    out: list[str] = []
    for si in root.findall("m:si", ns):
        texts = [t.text or "" for t in si.findall(".//m:t", ns)]
        out.append("".join(texts))
    return out


def _col_row(cell_ref: str) -> tuple[int, int]:
    m = re.match(r"([A-Z]+)(\d+)", cell_ref.upper())
    if not m:
        return 0, 0
    col_s, row_s = m.group(1), m.group(2)
    col = 0
    for ch in col_s:
        col = col * 26 + (ord(ch) - 64)
    return col - 1, int(row_s) - 1


def list_spreadsheet_sheets(path: PathLike) -> list[str]:
    """Blattnamen aus .xlsx/.xls (openpyxl / OOXML / xlrd)."""
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".xls" and not zipfile.is_zipfile(p):
        try:
            import xlrd  # type: ignore

            wb = xlrd.open_workbook(str(p), on_demand=True)
            try:
                return [str(n) for n in wb.sheet_names()]
            finally:
                wb.release_resources()
        except ImportError:
            return []
        except Exception:
            return []
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        try:
            return [str(n) for n in wb.sheetnames]
        finally:
            wb.close()
    except ImportError:
        pass
    except Exception:
        pass
    try:
        with zipfile.ZipFile(p, "r") as z:
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            names = []
            for s in wb_xml.findall("m:sheets/m:sheet", ns):
                n = s.get("name")
                if n:
                    names.append(str(n))
            return names
    except Exception:
        return []


def import_xls(
    path: PathLike,
    *,
    sheet: str | int | None = 0,
    header: bool = True,
    table_id: str = "xls1",
) -> DocumentTable:
    """Excel 97–2003 .xls (xlrd) oder .xls-mit-xlsx-Inhalt."""
    p = Path(path)
    if zipfile.is_zipfile(p):
        return import_xlsx(p, sheet=sheet, header=header, table_id=table_id)
    try:
        import xlrd  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "xlrd fehlt — .xls (Excel 97–2003) kann nicht gelesen werden. "
            "Bitte als .xlsx speichern oder xlrd installieren."
        ) from e
    wb = xlrd.open_workbook(str(p))
    if isinstance(sheet, str):
        ws = wb.sheet_by_name(sheet)
    else:
        idx = int(sheet or 0)
        ws = wb.sheet_by_index(idx)
    cells: list[list[str]] = []
    for r in range(ws.nrows):
        cells.append(
            ["" if ws.cell_value(r, c) is None else str(ws.cell_value(r, c)) for c in range(ws.ncols)]
        )
    while cells and all(not str(c).strip() for c in cells[-1]):
        cells.pop()
    if not cells:
        cells = [[""]]
    return create_table(data=cells, header=header, table_id=table_id)


def import_spreadsheet(
    path: PathLike,
    *,
    sheet: str | int | None = 0,
    header: bool = True,
    table_id: str = "sheet1",
) -> DocumentTable:
    """Erste/gewählte Tabelle aus .xlsx oder .xls."""
    p = Path(path)
    if p.suffix.lower() == ".xls" and not zipfile.is_zipfile(p):
        return import_xls(p, sheet=sheet, header=header, table_id=table_id)
    return import_xlsx(p, sheet=sheet, header=header, table_id=table_id)


def table_to_html_document(table: DocumentTable, *, title: str = "") -> str:
    """Vollständiges HTML mit einer QTextTable-tauglichen Tabelle."""
    inner = table_to_html(table)
    heading = f"<h1>{_escape_html(title)}</h1>" if title else ""
    return (
        "<!DOCTYPE html><html><head><meta charset='utf-8'/></head>"
        f"<body>{heading}{inner}</body></html>"
    )


def cells_from_html_tables(html: str) -> list[list[list[str]]]:
    """Alle HTML-Tabellen → Zellenlisten (für CSV/XLSX-Speichern)."""
    src = html or ""
    if "<table" not in src.lower():
        return []
    tables: list[list[list[str]]] = []
    for m in re.finditer(r"<table\b[^>]*>(.*?)</table>", src, flags=re.I | re.S):
        block = m.group(1)
        rows: list[list[str]] = []
        for rm in re.finditer(r"<tr\b[^>]*>(.*?)</tr>", block, flags=re.I | re.S):
            cells = [
                re.sub(r"<[^>]+>", "", c).replace("&nbsp;", " ").replace("&amp;", "&").strip()
                for c in re.findall(r"<t[dh]\b[^>]*>(.*?)</t[dh]>", rm.group(1), flags=re.I | re.S)
            ]
            if cells:
                rows.append(cells)
        if rows:
            tables.append(rows)
    return tables


def document_table_from_cells(
    cells: Sequence[Sequence[Any]] | None,
    *,
    header: bool = True,
    table_id: str = "t1",
) -> DocumentTable:
    return create_table(data=cells or [[""]], header=header, table_id=table_id)


def import_xlsx(
    path: PathLike,
    *,
    sheet: str | int | None = 0,
    header: bool = True,
    table_id: str = "xlsx1",
) -> DocumentTable:
    """Excel/.xlsx erste/gewählte Tabelle lesen (openpyxl oder Stdlib-OOXML)."""
    p = Path(path)
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        try:
            if isinstance(sheet, str):
                ws = wb[sheet]
            else:
                idx = int(sheet or 0)
                ws = wb.worksheets[idx]
            cells = []
            for row in ws.iter_rows(values_only=True):
                cells.append(["" if c is None else str(c) for c in row])
            # trailing empty rows trim
            while cells and all(not c.strip() for c in cells[-1]):
                cells.pop()
            if not cells:
                cells = [[""]]
            return create_table(data=cells, header=header, table_id=table_id)
        finally:
            wb.close()
    except ImportError:
        pass

    with zipfile.ZipFile(p, "r") as z:
        # sheet path
        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        ns = {
            "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        sheets = wb_xml.findall("m:sheets/m:sheet", ns)
        if not sheets:
            raise ValueError("Keine Arbeitsblätter in XLSX")
        if isinstance(sheet, str):
            sheet_el = next((s for s in sheets if s.get("name") == sheet), None)
            if sheet_el is None:
                raise KeyError(f"Sheet nicht gefunden: {sheet}")
        else:
            idx = int(sheet or 0)
            if idx < 0 or idx >= len(sheets):
                raise IndexError(f"Sheet-Index {idx}")
            sheet_el = sheets[idx]
        rid = sheet_el.get(f"{{{ns['r']}}}id") or sheet_el.get("r:id")
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        target = None
        for rel in rels.findall("r:Relationship", rel_ns):
            if rel.get("Id") == rid:
                target = rel.get("Target")
                break
        if not target:
            target = "worksheets/sheet1.xml"
        sheet_path = "xl/" + target.lstrip("/")
        if sheet_path.startswith("xl/xl/"):
            sheet_path = sheet_path[3:]
        shared = _xlsx_shared_strings(z)
        root = ET.fromstring(z.read(sheet_path))
        grid: dict[tuple[int, int], str] = {}
        max_r = max_c = 0
        for c in root.findall(".//m:c", ns):
            ref = c.get("r") or ""
            col, row = _col_row(ref)
            max_r = max(max_r, row)
            max_c = max(max_c, col)
            t = c.get("t")
            v_el = c.find("m:v", ns)
            if v_el is None or v_el.text is None:
                val = ""
            elif t == "s":
                try:
                    val = shared[int(v_el.text)]
                except (ValueError, IndexError):
                    val = v_el.text
            else:
                val = v_el.text
            grid[(row, col)] = val
        cells = []
        for r in range(max_r + 1):
            cells.append([grid.get((r, c), "") for c in range(max_c + 1)])
        while cells and all(not x.strip() for x in cells[-1]):
            cells.pop()
        if not cells:
            cells = [[""]]
        return create_table(data=cells, header=header, table_id=table_id)


def export_table_csv(
    table: DocumentTable,
    path: PathLike,
    *,
    delimiter: str = ";",
    encoding: str = "utf-8-sig",
) -> Path:
    p = Path(path)
    t = table.normalized()
    with p.open("w", encoding=encoding, newline="") as f:
        w = csv.writer(f, delimiter=delimiter)
        for row in t.cells:
            w.writerow(row)
    return p


def export_table_xlsx(table: DocumentTable, path: PathLike, *, sheet_name: str = "Tabelle1") -> Path:
    """Tabelle als .xlsx schreiben (openpyxl oder minimaler OOXML-Writer)."""
    p = Path(path)
    t = table.normalized()
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31] or "Tabelle1"
        for r_i, row in enumerate(t.cells, start=1):
            for c_i, val in enumerate(row, start=1):
                ws.cell(row=r_i, column=c_i, value=val)
        wb.save(str(p))
        return p
    except ImportError:
        pass
    return _write_minimal_xlsx(t.cells, p, sheet_name=sheet_name)


def _col_name(index: int) -> str:
    """0-basiert → Excel-Spaltenbuchstaben."""
    n = index + 1
    s = ""
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def _write_minimal_xlsx(cells: list[list[str]], path: Path, *, sheet_name: str = "Tabelle1") -> Path:
    """Minimaler XLSX-Writer ohne openpyxl (shared strings)."""
    shared: list[str] = []
    index: dict[str, int] = {}

    def sid(val: str) -> int:
        if val not in index:
            index[val] = len(shared)
            shared.append(val)
        return index[val]

    rows_xml = []
    for r_i, row in enumerate(cells, start=1):
        cells_xml = []
        for c_i, val in enumerate(row):
            ref = f"{_col_name(c_i)}{r_i}"
            i = sid(str(val))
            cells_xml.append(f'<c r="{ref}" t="s"><v>{i}</v></c>')
        rows_xml.append(f'<row r="{r_i}">{"".join(cells_xml)}</row>')

    def esc(s: str) -> str:
        return (
            s.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )

    ss_items = "".join(f"<si><t>{esc(x)}</t></si>" for x in shared)
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(rows_xml)}</sheetData></worksheet>'
    )
    shared_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        f'count="{len(shared)}" uniqueCount="{len(shared)}">{ss_items}</sst>'
    )
    wb_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<sheets><sheet name="{esc(sheet_name[:31] or "Tabelle1")}" sheetId="1" r:id="rId1"/></sheets>'
        "</workbook>"
    )
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""
    wb_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("xl/workbook.xml", wb_xml)
        z.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
        z.writestr("xl/sharedStrings.xml", shared_xml)
    return path


def list_table_styles() -> list[dict[str, Any]]:
    return [
        {"id": "default", "name": "Standard"},
        {"id": "striped", "name": "Gestreift"},
        {"id": "compact", "name": "Kompakt"},
    ]
