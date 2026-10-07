#!/usr/bin/env python3
"""Offscreen End-to-End-Audit 2.6.52 — echte MainWindow-Pfade, keine Needles.

Deckt die Feldfehler 2.6.36–2.6.51 ab und beweist die Fixes:

1. Fenster-Mindestgröße: PDF-Toolbar/Sidebar/Ribbon/Statusleiste dürfen das
   Hauptfenster nicht breiter/höher als einen Laptop-Monitor erzwingen
   (vorher 7286×2175 px → Seite außerhalb des Bildschirms = „weiße Ansicht“).
2. PDF über ``MainWindow.open_path``: Canvas sichtbar **innerhalb** des
   Viewports, ``grab()`` nicht weiß, echte Seitenanzahl (Hintergrund-Refresh),
   Navigation ▶, Thumbs für alle Seiten mit Tinte.
3. Editor: Fett/Kursiv/Unterstrichen/Textmarker unabhängig (kein Überschreiben),
   ``setPlainText`` erbt kein Format vom vorherigen Dokument.
4. DOCX: Stil-basierte Formate, Hyperlink-Text, Größe/Farbe; Roundtrip mit
   Textmarker nach DOCX.
5. Geräte-Menü + Ribbon-Tab + Scan-Dialog öffnen ohne Absturz.

Aufruf: ``QT_QPA_PLATFORM=offscreen python3 scripts/test_ui_audit_2652.py``
"""

from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["ILD_SMOKE_QT"] = "1"
os.environ.setdefault("ILD_SKIP_DEPS_CHECK", "1")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from test_pdf_canvas_not_blank import _make_multipage_pdf, _nonwhite_count  # noqa: E402

# Laptop-Klasse: 1366×768 muss ohne Clipping funktionieren
MAX_MIN_WIDTH = 1200
MAX_MIN_HEIGHT = 700


def _pump(app, seconds: float = 0.3, until=None) -> None:
    t0 = time.time()
    while time.time() - t0 < seconds:
        app.processEvents()
        if until is not None and until():
            return
        time.sleep(0.01)


def _make_docx(path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = Document()
    d.add_heading("Überschrift Eins", level=1)
    p = d.add_paragraph()
    p.add_run("Normal ")
    r = p.add_run("Fett")
    r.bold = True
    p.add_run(" und ")
    r = p.add_run("Kursiv")
    r.italic = True
    p.add_run(" und ")
    r = p.add_run("Unterstrichen")
    r.underline = True
    p2 = d.add_paragraph()
    p2.add_run("Via Style: ")
    r = p2.add_run("StrongStyle")
    r.style = d.styles["Strong"]
    r = p2.add_run(" EmphasisStyle")
    r.style = d.styles["Emphasis"]
    p3 = d.add_paragraph()
    r = p3.add_run("GrossRot")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p4 = d.add_paragraph()
    p4.add_run("Link: ")
    r_id = d.part.relate_to(
        "https://example.org",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    run_el.append(rpr)
    t = OxmlElement("w:t")
    t.text = "HyperlinkText"
    run_el.append(t)
    hl.append(run_el)
    p4._p.append(hl)
    d.save(str(path))


def _runs(document):
    from PySide6.QtGui import QFont

    out = []
    blk = document.begin()
    while blk.isValid():
        it = blk.begin()
        while not it.atEnd():
            frag = it.fragment()
            f = frag.charFormat()
            out.append(
                {
                    "text": frag.text(),
                    "bold": f.fontWeight() >= QFont.Bold,
                    "italic": bool(f.fontItalic()),
                    "underline": bool(f.fontUnderline()),
                    "size": float(f.fontPointSize() or 0.0),
                    "color": f.foreground().color().name(),
                }
            )
            it += 1
        blk = blk.next()
    return out


def main() -> int:  # noqa: C901 - ein Skript, bewusst linear
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QFont, QTextCursor
    from PySide6.QtWidgets import QApplication, QDialog, QMenu

    from instantlensdoc.license import LicenseManager
    from instantlensdoc.ui.image_qt import qpixmap_has_ink
    from instantlensdoc.ui.main_window import MainWindow
    from instantlensdoc.ui.ribbon_bar import RibbonBar

    app = QApplication.instance() or QApplication([])
    lm = LicenseManager()
    lm.ensure_trial_started()
    win = MainWindow(lm)
    win.resize(1600, 900)
    win.show()
    _pump(app, 0.2)

    # ---- 1) Mindestgröße ------------------------------------------------
    msh = win.minimumSizeHint()
    assert msh.width() <= MAX_MIN_WIDTH, (
        f"MainWindow-Mindestbreite {msh.width()} px > {MAX_MIN_WIDTH} — Toolbar/Ribbon/"
        "Statusleiste erzwingen Off-Screen-Layout (weiße Ansicht)"
    )
    assert msh.height() <= MAX_MIN_HEIGHT, (
        f"MainWindow-Mindesthöhe {msh.height()} px > {MAX_MIN_HEIGHT} — Sidebar ohne Scroll"
    )
    pv = win.pdf_view
    assert pv.minimumSizeHint().width() <= 400, (
        f"PdfViewer-Mindestbreite {pv.minimumSizeHint().width()} px — Toolbar bricht nicht um"
    )
    assert win.size().width() <= 1600 + 8 and win.size().height() <= 900 + 8, (
        f"Fenster wurde über resize() hinaus aufgeblasen: {win.size()}"
    )
    print("OK  min size", msh.width(), "x", msh.height())

    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        pdf = tdp / "audit_multi.pdf"
        pages = ["PageOneAAA", "PageTwoBBB", "PageThreeCCC", "PageFourDDD", "PageFiveEEE"]
        _make_multipage_pdf(pdf, pages)

        # ---- 2) PDF via MainWindow.open_path -----------------------------
        win.open_path(str(pdf))
        _pump(app, 2.0, until=lambda: int(pv.page_count or 0) >= len(pages))
        assert win.stack.currentWidget() is pv, "Stack zeigt nach PDF-Open nicht den PdfViewer"
        assert pv.isVisible(), "PdfViewer unsichtbar"
        assert pv._canvas_has_page_image(), (
            f"Canvas ohne Seitenbild: err={pv._last_refresh_error!r}"
        )
        # Seitenanzahl aus Hintergrund-Refresh (vorher dauerhaft 1)
        assert int(pv.page_count) == len(pages), (
            f"page_count={pv.page_count} statt {len(pages)} — Hintergrund-Refresh kam nie an"
        )
        # Canvas liegt innerhalb des sichtbaren Viewports
        vp = pv.scroll.viewport()
        canvas_pos = pv.canvas.mapTo(vp, pv.canvas.rect().topLeft())
        assert 0 <= canvas_pos.x() < vp.width(), (
            f"Canvas x={canvas_pos.x()} außerhalb Viewport-Breite {vp.width()}"
        )
        assert vp.width() <= win.width(), "Viewport breiter als Fenster"
        grab = pv.grab().toImage()
        nw = _nonwhite_count(grab, 10)
        assert nw >= 20, f"PdfViewer.grab() fast weiß (nonwhite={nw})"
        central = win.stack.grab().toImage()
        assert _nonwhite_count(central, 10) >= 20, "Zentraler Stack grab() fast weiß"
        print("OK  pdf open: page_count", pv.page_count, "canvas@", canvas_pos.x(), canvas_pos.y())

        # Navigation ▶ (vorher tot wegen page_count=1)
        pv.next_page()
        _pump(app, 0.3)
        assert pv.page_index == 1, f"next_page blieb auf {pv.page_index}"
        assert pv._canvas_has_page_image(), "Seite 2 ohne Bild"
        pv.goto_page(4)
        _pump(app, 0.3)
        assert pv.page_index == 4, f"goto_page(4) -> {pv.page_index}"
        assert pv._canvas_has_page_image(), "Seite 5 ohne Bild"
        # Thumbs: alle Seiten, mit Tinte (Lazy-Loader)
        thumbs = win.sidebar.thumbs
        _pump(app, 2.0, until=lambda: thumbs.count() >= len(pages))
        assert thumbs.count() == len(pages), f"Thumbs: {thumbs.count()} statt {len(pages)}"

        def _all_thumbs_ink() -> bool:
            for i in range(thumbs.count()):
                pm = thumbs.item(i).icon().pixmap(96, 128)
                if pm.isNull() or not qpixmap_has_ink(pm):
                    return False
            return True

        _pump(app, 3.0, until=_all_thumbs_ink)
        assert _all_thumbs_ink(), "Mindestens ein Thumb grau/leer"
        print("OK  navigation + thumbs", thumbs.count())

        # ---- 3) Editor-Formate ---------------------------------------------
        txt = tdp / "plain.txt"
        txt.write_text("hello world\nzweite Zeile\n", encoding="utf-8")
        ed = win.editor
        # Format-Erbe provozieren: zuerst fettes Rich-Dokument laden
        ed.set_rich_html("<p><b><u>MMMM</u></b></p>")
        win.open_path(str(txt))
        _pump(app, 0.2)
        assert win.stack.currentWidget() is win.editor_pane
        # Frisch geöffnet darf nie „geändert“ sein (Stack-Wechsel vor Inhalt → sticky dirty)
        assert not win._current_is_dirty(), "TXT direkt nach Öffnen als geändert markiert"
        probe = QTextCursor(ed.document())
        probe.setPosition(1)
        f0 = probe.charFormat()
        assert f0.fontWeight() < QFont.Bold and not f0.fontUnderline(), (
            "TXT hat Fett/Unterstrichen vom vorherigen Dokument geerbt"
        )
        cur = ed.textCursor()
        cur.setPosition(0)
        cur.setPosition(5, QTextCursor.KeepAnchor)
        ed.setTextCursor(cur)

        def fmt_at(pos: int):
            c = QTextCursor(ed.document())
            c.setPosition(pos + 1)
            f = c.charFormat()
            return (
                f.fontWeight() >= QFont.Bold,
                bool(f.fontItalic()),
                bool(f.fontUnderline()),
                f.background().style() != Qt.NoBrush,
            )

        win._toggle_bold()
        assert fmt_at(0) == (True, False, False, False), fmt_at(0)
        win._toggle_underline()
        assert fmt_at(0) == (True, False, True, False), f"Unterstreichen löschte Fett: {fmt_at(0)}"
        win._on_editor_toolbar_action("mark")
        assert fmt_at(0) == (True, False, True, True), f"Markieren löschte Formate: {fmt_at(0)}"
        win._toggle_italic()
        assert fmt_at(0) == (True, True, True, True), f"Kursiv löschte Formate: {fmt_at(0)}"
        win._toggle_bold()
        assert fmt_at(0) == (False, True, True, True), f"Fett-aus löschte Formate: {fmt_at(0)}"
        assert all(fmt_at(i) == (False, True, True, True) for i in range(5)), "Auswahl uneinheitlich"
        assert fmt_at(6) == (False, False, False, False), "Format lief über die Auswahl hinaus"
        html = ed.to_rich_html()
        assert "background-color" in html and "underline" in html, "HTML ohne Marker/Unterstrich"
        win._clear_editor_marks()
        assert fmt_at(0) == (False, True, True, False), f"Markierungen löschen zerstörte Formate: {fmt_at(0)}"
        print("OK  editor B/I/U/Textmarker unabhängig")

        # ---- 4) DOCX -------------------------------------------------------
        docx_path = tdp / "audit.docx"
        _make_docx(docx_path)
        win.open_path(str(docx_path))
        _pump(app, 0.3)
        assert win.stack.currentWidget() is win.editor_pane
        assert win.doc is not None and not win.doc.meta.get("rich_text_error"), (
            f"DOCX Rich-Text-Fehler: {win.doc.meta.get('rich_text_error')}"
        )
        assert not win._current_is_dirty(), "DOCX direkt nach Öffnen als geändert markiert"
        runs = {r["text"].strip(): r for r in _runs(ed.document()) if r["text"].strip()}
        assert runs["Fett"]["bold"] and not runs["Fett"]["italic"], runs["Fett"]
        assert runs["Kursiv"]["italic"], runs["Kursiv"]
        assert runs["Unterstrichen"]["underline"], runs["Unterstrichen"]
        assert runs["StrongStyle"]["bold"], f"Zeichenstil Strong verloren: {runs['StrongStyle']}"
        assert runs["EmphasisStyle"]["italic"], f"Zeichenstil Emphasis verloren: {runs['EmphasisStyle']}"
        assert runs["Überschrift Eins"]["bold"], "Heading 1 nicht fett"
        assert "HyperlinkText" in runs, "Hyperlink-Text im DOCX verloren"
        assert runs["HyperlinkText"]["bold"], "Hyperlink-Run-Format verloren"
        assert abs(runs["GrossRot"]["size"] - 20.0) < 0.6, f"Schriftgröße verloren: {runs['GrossRot']}"
        assert runs["GrossRot"]["color"].lower() == "#c00000", f"Farbe verloren: {runs['GrossRot']}"
        # visuell: Editor-Grab nicht leer
        assert _nonwhite_count(ed.grab().toImage(), 6) > 30, "Editor grab() leer"
        # Roundtrip: „Normal“ unterstreichen + markieren → speichern → neu laden
        pos = ed.document().toPlainText().index("Normal")
        cur = ed.textCursor()
        cur.setPosition(pos)
        cur.setPosition(pos + 6, QTextCursor.KeepAnchor)
        ed.setTextCursor(cur)
        win._toggle_underline()
        win._on_editor_toolbar_action("mark")
        out = tdp / "audit_roundtrip.docx"
        win.doc.path = out
        assert win.save_doc(), "DOCX speichern fehlgeschlagen"
        from docx import Document as DocxDocument

        dd = DocxDocument(str(out))
        flat = {r.text: r for p in dd.paragraphs for r in p.runs}
        assert flat["Normal"].underline, "Unterstreichen nicht im DOCX"
        assert str(flat["Normal"].font.highlight_color).startswith("YELLOW"), (
            f"Textmarker nicht im DOCX: {flat['Normal'].font.highlight_color}"
        )
        assert flat["Fett"].bold and flat["Kursiv"].italic, "B/I nach Roundtrip verloren"
        assert flat["StrongStyle"].bold, "Stil-Fett nach Roundtrip verloren"
        assert any(r.text == "HyperlinkText" for r in flat.values()), "Hyperlink-Text nach Roundtrip verloren"
        win.open_path(str(out))
        _pump(app, 0.3)
        runs2 = {r["text"].strip(): r for r in _runs(ed.document()) if r["text"].strip()}
        assert runs2["Normal"]["underline"], "Roundtrip: Unterstrich verloren"
        assert not win._current_is_dirty(), "Roundtrip-DOCX direkt nach Öffnen als geändert markiert"
        print("OK  docx formats + roundtrip")

        # ---- 5) Geräte / Scan ------------------------------------------------
        titles = [a.text().replace("&", "") for a in win.menuBar().actions()]
        assert "Geräte" in titles, f"Menü Geräte fehlt: {titles}"
        m = win.findChild(QMenu, "menuDevices")
        names = {a.objectName() for a in m.actions()}
        for need in ("actDevicesScanner", "actDevicesPrinters", "actDevicesDiscover", "actDevicesRefresh"):
            assert need in names, f"Geräte-Aktion fehlt: {need}"
        rb = win.findChild(RibbonBar)
        assert rb is not None and "Geräte" in [b.text() for b in rb._cat_buttons], "Ribbon-Tab Geräte fehlt"
        from instantlensdoc.ui.scan_dialog import ScanDialog

        dlg = ScanDialog(pv, win, auto_launch_scantuxio=False)
        dlg.show()
        _pump(app, 0.1)
        assert dlg.isVisible(), "ScanDialog nicht sichtbar"
        dlg.refresh_devices()
        _pump(app, 0.5)
        dlg.close()
        # Menü-Aktionen öffnen modale Dialoge → automatisch schließen
        from PySide6.QtCore import QTimer

        opened: list[str] = []

        def _close_modal() -> None:
            w = QApplication.activeModalWidget()
            if isinstance(w, QDialog):
                opened.append(w.objectName() or w.windowTitle())
                w.reject()

        for name in ("actDevicesScanner", "actDevicesDiscover"):
            act = next(a for a in m.actions() if a.objectName() == name)
            QTimer.singleShot(400, _close_modal)
            act.trigger()
            _pump(app, 0.1)
        assert len(opened) == 2, f"Geräte-Dialoge nicht geöffnet: {opened}"
        print("OK  devices menu/ribbon/scan dialog", opened)

    assert not win._current_is_dirty(), "Dokument vor close() unerwartet dirty (Beenden-Dialog)"
    win.close()
    print("OK test_ui_audit_2652")
    return 0


if __name__ == "__main__":
    import faulthandler

    # Nie still hängen (modale Dialoge offscreen): Traceback + Exit nach 180 s
    faulthandler.dump_traceback_later(180, exit=True)
    raise SystemExit(main())
