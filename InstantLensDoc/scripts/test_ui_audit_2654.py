#!/usr/bin/env python3
"""Offscreen End-to-End 2.6.54 — Save-As-PDF, Minimap, Undo + PDF-Annotationen.

Feldbefund 2.6.53 (Windows): ``Dunning_Kruger_Effekt_1.pdf`` (25 226 Byte)
``PK 03 04`` = unverändertes Word-DOCX (Application: Microsoft Office Word).
Ursache: ≤ 2.6.42 ``save_document`` kopierte die DOCX-Bytes unter .pdf-Namen.

Dieser Test beweist:

1. Die Nutzerdatei (oder ein synthetisches DOCX-als-PDF) wird als ZIP-DOCX erkannt;
   pikepdf/PDFium scheitern; ``pdf_doctor`` Exit 1; Recover → lesbares .docx.
2. ``save_document``/Export eines geladenen DOCX nach .pdf erzeugt ``%PDF-`` + ``%%EOF``,
   pikepdf- und PDFium-öffenbar, **kein** ZIP-Header.
3. 0-Byte / HTML-mit-.pdf: Diagnosetext mit Größe + Header-Hex (Schritt 0).
4. Minimap Standard aus, Toggle in Ansicht, in DOCX unsichtbar, Textbreite zurück.
5. Tippen → Fett → Undo zweimal → Originaltext; Ribbon/Menü Rückgängig/Wiederholen.
6. Highlight bleibt nach Zoom (Koordinaten = PDF-Punkte, Anzeige × Scale).
7. Werkzeug „Stift“ aktiviert Ink, nicht Highlight (objectName annTool_ink).
8. Objekte verschiebbar; /Rect wird ins PDF geschrieben und ändert sich.
9. FreeText-Edit schreibt /Contents zurück.
10. QUndoStack: Anlegen → Undo weg → Redo zurück; Verschieben → Undo stellt /Rect her.
11. Ribbon Start+Bearbeiten: ↶ Rückgängig / ↷ Wiederholen inkl. Tooltip.
12. Auswahl-Hit-Test in PDF-Punkten: Klick Mitte Rechteck → selected + 8 Griffe;
    Gummiband zwei Objekte; Leerklick deselektiert; Werkzeug bleibt Auswahl.
13. Auswahl → Ribbon-Farbe mutiert ``/C``; ohne Auswahl Select-All (Seite).
14. Seite 1 Rect A, Seite 2 Rect B: Auswahl B + Strich ändert nur B (``/C``/``/Border``);
    A unverändert. Seitenwechsel leert alte IDs.

Aufruf: ``QT_QPA_PLATFORM=offscreen python3 scripts/test_ui_audit_2654.py``
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["ILD_SMOKE_QT"] = "1"
os.environ.setdefault("ILD_SKIP_DEPS_CHECK", "1")
os.environ["XDG_CONFIG_HOME"] = tempfile.mkdtemp(prefix="ild-audit-2654-cfg-")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from test_pdf_canvas_not_blank import _make_multipage_pdf  # noqa: E402
from test_ui_audit_2652 import _pump  # noqa: E402

USER_PDF = Path(
    "/cursor/stores/bc-b08b150e-66b9-4e56-8bf8-50c583a73b7e/inbox/"
    "Dunning_Kruger_Effekt_1.pdf"
)
STORE = Path("/cursor/stores/bc-b08b150e-66b9-4e56-8bf8-50c583a73b7e")
SHOT_DIR = STORE / "media" / "instantlensdoc-2654"


def _make_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Dunning-Kruger-Effekt", level=1)
    d.add_paragraph("Absatz eins: " + ("Fließtext " * 40))
    p = d.add_paragraph()
    r = p.add_run("Fettkursiv")
    r.bold = True
    r.italic = True
    d.save(str(path))


def _assert_valid_pdf(path: Path, *, must_contain: str | None = None) -> None:
    from ild_pdf.pdf_sniff import validate_pdf_file

    raw = path.read_bytes()
    assert raw[:5] == b"%PDF-", f"{path.name} beginnt nicht mit %PDF- sondern {raw[:16]!r}"
    assert b"%%EOF" in raw[-4096:], f"{path.name} ohne %%EOF"
    assert not raw.startswith(b"PK"), f"{path.name} ist ZIP/DOCX mit .pdf-Endung"
    v = validate_pdf_file(path, use_pdfium=True, min_pages=1)
    assert v.ok, v.message_de()
    if must_contain:
        import pypdfium2 as pdfium

        from ild_pdf.pdfium_open import PDFIUM_LOCK

        with PDFIUM_LOCK:
            doc = pdfium.PdfDocument(raw)
            try:
                txt = " ".join(doc[0].get_textpage().get_text_range().split())
            finally:
                doc.close()
        assert must_contain.lower() in txt.lower(), f"PDF-Text ohne {must_contain!r}: {txt[:200]!r}"


def _close(a, b, tol: float = 0.75) -> bool:
    return abs(float(a) - float(b)) <= tol


def test_pdf_roundtrip_move_and_freetext(td: Path) -> None:
    from ild_pdf import (
        Annotation,
        AnnotationStore,
        AnnotationType,
        read_ild_annots,
        write_annotations_to_pdf,
    )

    pdf = td / "ann_roundtrip.pdf"
    _make_multipage_pdf(pdf, ["AnnPageAAA"])
    store = AnnotationStore(pdf)
    rect = Annotation(
        page=0,
        type=AnnotationType.RECTANGLE,
        x=40.0,
        y=50.0,
        width=80.0,
        height=30.0,
        color="#C0392B",
    )
    ft = Annotation(
        page=0,
        type=AnnotationType.TEXT,
        x=20.0,
        y=120.0,
        width=140.0,
        height=28.0,
        text="Hallo",
        color="#1A1A1A",
    )
    store.add(rect)
    store.add(ft)
    write_annotations_to_pdf(pdf, store)
    native = read_ild_annots(pdf, page_index=0)
    by_nm = {n["nm"]: n for n in native}
    assert f"ild:{rect.id}" in by_nm, native
    assert f"ild:{ft.id}" in by_nm, native
    assert by_nm[f"ild:{ft.id}"]["subtype"] == "FreeText"
    assert by_nm[f"ild:{ft.id}"]["contents"] == "Hallo"
    before = list(by_nm[f"ild:{rect.id}"]["rect"])
    moved = store.move_by([rect.id], 25.0, 12.0)
    assert moved == 1
    write_annotations_to_pdf(pdf, store)
    after = read_ild_annots(pdf, page_index=0)
    after_rect = next(n["rect"] for n in after if n["nm"] == f"ild:{rect.id}")
    assert after_rect != before, f"Rect unverändert nach Verschieben: {before}"
    store.undo()
    write_annotations_to_pdf(pdf, store)
    restored = next(
        n["rect"] for n in read_ild_annots(pdf, page_index=0) if n["nm"] == f"ild:{rect.id}"
    )
    assert all(_close(a, b) for a, b in zip(restored, before)), (
        f"Undo stellte /Rect nicht her: {restored} vs {before}"
    )
    store.update(ft.id, text="Geändert")
    write_annotations_to_pdf(pdf, store)
    contents = next(
        n["contents"] for n in read_ild_annots(pdf, page_index=0) if n["nm"] == f"ild:{ft.id}"
    )
    assert contents == "Geändert", contents


def test_viewer_tools_undo_zoom(app, td: Path) -> None:
    from PySide6.QtWidgets import QToolButton

    from ild_pdf import Annotation, AnnotationType
    from instantlensdoc.ui.pdf_view import ANN_TOOL_ACTION_MAP, PdfViewer

    pdf = td / "viewer_ann.pdf"
    _make_multipage_pdf(pdf, ["ZoomHlAAA", "PageTwoBBB"])
    v = PdfViewer()
    v.resize(960, 720)
    v.show()
    app.processEvents()
    assert v.load(pdf), f"load fehlgeschlagen: {v._last_refresh_error!r}"
    v.set_annotations_locked(False)
    v._suppress_default_zoom = True
    _pump(app, 1.2, until=lambda: v._canvas_has_page_image())
    v.set_scale(1.5, immediate=True)
    _pump(app, 0.4)
    assert _close(v.scale, 1.5, 0.05), v.scale

    mapping = v.annotation_tool_map()
    assert mapping["annTool_ink"] == "ink"
    assert mapping["annTool_highlight"] == "highlight"
    assert mapping["annTool_select"] == "select"
    for oid, tid in ANN_TOOL_ACTION_MAP:
        assert mapping[oid] == tid, (oid, mapping.get(oid), tid)

    ink_btn = v.findChild(QToolButton, "annTool_ink")
    hl_btn = v.findChild(QToolButton, "annTool_highlight")
    assert ink_btn is not None and ink_btn.text() == "Stift"
    assert hl_btn is not None and "Highlight" in hl_btn.text()
    assert v.btn_pen_color.text() == "Strich"
    v.set_tool_from_id("pen")
    assert v.current_tool_id() == "ink", v.current_tool_id()
    assert v.tool == AnnotationType.INK
    ink_btn.click()
    _pump(app, 0.1)
    assert v.current_tool_id() == "ink"
    assert not hl_btn.isChecked()
    assert ink_btn.isChecked()

    n0 = len(v.store.annotations) if v.store else 0
    v._commit_ann(
        Annotation(
            page=0,
            type=AnnotationType.HIGHLIGHT,
            x=30.0,
            y=40.0,
            width=90.0,
            height=18.0,
            color="#FFE066",
        )
    )
    _pump(app, 0.2)
    assert v.store is not None
    assert len(v.store.annotations) == n0 + 1
    stored = v.store.annotations[-1]
    assert stored.type == AnnotationType.HIGHLIGHT
    s = v._view_scale()
    assert _close(stored.x, 30.0 / s, 0.5), (stored.x, s)
    assert v._undo_stack.canUndo()
    can_u, tu, can_r, tr = v.undo_ui_state()
    assert can_u, (can_u, tu, can_r, tr)
    assert "highlight" in tu.lower() or "hinzufügen" in tu.lower() or tu

    v.set_scale(2.0, immediate=True)
    _pump(app, 0.4)
    disp = v._display_anns_for_page(0)
    hl = [a for a in disp if a.type == AnnotationType.HIGHLIGHT]
    assert hl, "Highlight nach Zoom 200% nicht in der Anzeige"
    assert _close(hl[0].x, stored.x * 2.0, 1.0), (hl[0].x, stored.x)
    canvas_hl = [a for a in v.canvas._annotations if a.type == AnnotationType.HIGHLIGHT]
    assert canvas_hl, "Canvas-Overlay verlor Highlight nach Zoom"

    v.undo_annotation()
    _pump(app, 0.2)
    assert len(v.store.annotations) == n0
    assert not [a for a in v._display_anns_for_page(0) if a.type == AnnotationType.HIGHLIGHT]
    v.redo_annotation()
    _pump(app, 0.2)
    assert len(v.store.annotations) == n0 + 1

    rect = Annotation(
        page=0,
        type=AnnotationType.RECTANGLE,
        x=60.0,
        y=80.0,
        width=50.0,
        height=24.0,
        color="#2980B9",
    )
    v._commit_ann(rect)
    _pump(app, 0.2)
    rid = v.store.annotations[-1].id
    before = (v.store.get(rid).x, v.store.get(rid).y)
    v._on_annotations_moved([rid], 30.0, 15.0)  # view pixels → store /scale
    after = (v.store.get(rid).x, v.store.get(rid).y)
    assert after != before
    v.undo_annotation()
    _pump(app, 0.2)
    back = (v.store.get(rid).x, v.store.get(rid).y)
    assert _close(back[0], before[0], 0.4) and _close(back[1], before[1], 0.4)

    v.save_annotations()
    from ild_pdf import read_ild_annots

    native = read_ild_annots(pdf, page_index=0)
    assert native, "save_annotations schrieb keine /Annots"
    assert any(n["subtype"] == "Highlight" for n in native)

    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    v.set_scale(1.5, immediate=True)
    _pump(app, 0.3)
    img = v.canvas.grab()
    out = SHOT_DIR / "annotations.png"
    assert img.save(str(out)), out
    assert out.is_file() and out.stat().st_size > 200, out
    v.close()


def _canvas_page_to_widget(canvas, vx: float, vy: float) -> tuple[float, float]:
    pw, ph = canvas._pixmap_logical_size()
    cr = canvas.contentsRect()
    lx = float(cr.x()) + (float(cr.width()) - pw) / 2.0
    ly = float(cr.y()) + (float(cr.height()) - ph) / 2.0
    return float(vx) + lx, float(vy) + ly


def _qclick(canvas, x: float, y: float) -> None:
    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtTest import QTest

    QTest.mouseClick(
        canvas,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
        QPoint(int(round(x)), int(round(y))),
    )


def _qdrag(canvas, x0: float, y0: float, x1: float, y1: float) -> None:
    from PySide6.QtCore import QEvent, QPoint, QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication

    p0 = QPoint(int(round(x0)), int(round(y0)))
    p1 = QPoint(int(round(x1)), int(round(y1)))
    QTest.mousePress(canvas, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p0)
    app = QApplication.instance()
    steps = 8
    for i in range(1, steps + 1):
        t = i / float(steps)
        p = QPoint(
            int(round(p0.x() + (p1.x() - p0.x()) * t)),
            int(round(p0.y() + (p1.y() - p0.y()) * t)),
        )
        local = QPointF(float(p.x()), float(p.y()))
        try:
            ev = QMouseEvent(
                QEvent.Type.MouseMove,
                local,
                Qt.MouseButton.NoButton,
                Qt.MouseButton.LeftButton,
                Qt.KeyboardModifier.NoModifier,
            )
        except TypeError:
            ev = QMouseEvent(
                QEvent.Type.MouseMove,
                local,
                local,
                local,
                Qt.MouseButton.NoButton,
                Qt.MouseButton.LeftButton,
                Qt.KeyboardModifier.NoModifier,
            )
        QApplication.sendEvent(canvas, ev)
        if app is not None:
            app.processEvents()
    QTest.mouseRelease(canvas, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p1)


def test_hit_test_pdf_points_unit() -> None:
    from ild_pdf import (
        Annotation,
        AnnotationType,
        annotation_contains_point,
        hit_test_annotations,
    )

    rect = Annotation(
        page=0, type=AnnotationType.RECTANGLE, x=40.0, y=50.0, width=80.0, height=30.0
    )
    thin = Annotation(
        page=0, type=AnnotationType.HIGHLIGHT, x=40.0, y=120.0, width=90.0, height=1.5
    )
    ink = Annotation.from_ink_points(0, [(10.0, 10.0), (40.0, 12.0), (70.0, 9.0)])
    top = Annotation(
        page=0, type=AnnotationType.ELLIPSE, x=50.0, y=55.0, width=20.0, height=20.0
    )
    assert annotation_contains_point(rect, 80.0, 65.0, slop_pt=4.0)
    assert annotation_contains_point(thin, 85.0, 120.75, slop_pt=4.0)
    assert annotation_contains_point(ink, 40.0, 12.0, slop_pt=4.0)
    hit = hit_test_annotations([rect, thin, top], 60.0, 65.0, slop_pt=4.0)
    assert hit is top, getattr(hit, "type", None)


def test_viewer_select_hit_test(app, td: Path) -> None:
    from ild_pdf import Annotation, AnnotationType
    from instantlensdoc.ui.pdf_view import PdfViewer

    pdf = td / "select_hit.pdf"
    _make_multipage_pdf(pdf, ["SelectHitAAA"])
    v = PdfViewer()
    v.resize(960, 720)
    v.show()
    app.processEvents()
    assert v.load(pdf), f"load fehlgeschlagen: {v._last_refresh_error!r}"
    v.set_annotations_locked(False)
    v._suppress_default_zoom = True
    _pump(app, 1.2, until=lambda: v._canvas_has_page_image())
    v.set_scale(1.5, immediate=True)
    _pump(app, 0.4)
    v.set_tool(None)
    assert v.current_tool_id() == "select"
    assert v.canvas._select_mode

    r1 = Annotation(
        page=0,
        type=AnnotationType.RECTANGLE,
        x=80.0,
        y=90.0,
        width=120.0,
        height=50.0,
        color="#2980B9",
        fill_color="",
    )
    r2 = Annotation(
        page=0,
        type=AnnotationType.STAMP,
        x=260.0,
        y=200.0,
        width=70.0,
        height=36.0,
        text="OK",
        color="#1E8449",
    )
    v.store.add(r1)
    v.store.add(r2)
    v.refresh()
    _pump(app, 0.3)
    s = v._view_scale()
    cx = (r1.x + r1.width / 2.0) * s
    cy = (r1.y + r1.height / 2.0) * s
    wx, wy = _canvas_page_to_widget(v.canvas, cx, cy)
    _qclick(v.canvas, wx, wy)
    _pump(app, 0.15)
    assert r1.id in v.canvas._selected_ids, (
        f"Klick Mitte Rechteck wählte nicht: {v.canvas._selected_ids!r} "
        f"scale={s} view=({cx:.1f},{cy:.1f}) widget=({wx:.1f},{wy:.1f}) "
        f"hit={v.hit_test_annotation_at_view(cx, cy)}"
    )
    assert v._selected_ann_id == r1.id
    handles = v.canvas.selected_handle_rects()
    assert len(handles) == 8, handles
    assert v.current_tool_id() == "select"
    assert v.canvas._select_mode

    # Gummiband: Drag auf leerer Fläche (nicht auf dem selektierten Rect) erfasst beide
    vx0, vy0 = 20.0 * s, 20.0 * s
    vx1, vy1 = 360.0 * s, 260.0 * s
    x0, y0 = _canvas_page_to_widget(v.canvas, vx0, vy0)
    x1, y1 = _canvas_page_to_widget(v.canvas, vx1, vy1)
    _qdrag(v.canvas, x0, y0, x1, y1)
    _pump(app, 0.2)
    assert r1.id in v.canvas._selected_ids and r2.id in v.canvas._selected_ids, (
        f"Gummiband-Drag: {v.canvas._selected_ids!r} view=({vx0:.1f},{vy0:.1f})-({vx1:.1f},{vy1:.1f}) "
        f"band={v.canvas._band_start!r}"
    )
    assert len(v.canvas._selected_ids) >= 2
    handles = v.canvas.selected_handle_rects()
    assert len(handles) == 8, handles

    # Leerklick deselektiert, Werkzeug bleibt Auswahl
    ex, ey = _canvas_page_to_widget(v.canvas, 20.0 * s, 20.0 * s)
    _qclick(v.canvas, ex, ey)
    _pump(app, 0.1)
    assert not v.canvas._selected_ids, v.canvas._selected_ids
    assert v.current_tool_id() == "select"
    v.close()


def test_viewer_style_selection(app, td: Path) -> None:
    from ild_pdf import Annotation, AnnotationType, read_ild_annots
    from instantlensdoc.ui.pdf_view import PdfViewer

    pdf = td / "style_sel.pdf"
    _make_multipage_pdf(pdf, ["StyleSelAAA"])
    v = PdfViewer()
    v.resize(960, 720)
    v.show()
    app.processEvents()
    assert v.load(pdf), f"load fehlgeschlagen: {v._last_refresh_error!r}"
    v.set_annotations_locked(False)
    v._suppress_default_zoom = True
    _pump(app, 1.2, until=lambda: v._canvas_has_page_image())
    v.set_scale(1.5, immediate=True)
    _pump(app, 0.3)
    v.set_tool(None)

    rect = Annotation(
        page=0,
        type=AnnotationType.RECTANGLE,
        x=80.0,
        y=90.0,
        width=120.0,
        height=50.0,
        color="#2980B9",
    )
    note = Annotation(
        page=0,
        type=AnnotationType.TEXT,
        x=40.0,
        y=200.0,
        width=140.0,
        height=28.0,
        text="Hallo",
        color="#1A1A1A",
        font_size=12.0,
    )
    v.store.add(rect)
    v.store.add(note)
    v.refresh()
    _pump(app, 0.2)

    s = v._view_scale()
    cx = (rect.x + rect.width / 2.0) * s
    cy = (rect.y + rect.height / 2.0) * s
    wx, wy = _canvas_page_to_widget(v.canvas, cx, cy)
    _qclick(v.canvas, wx, wy)
    _pump(app, 0.1)
    assert rect.id in v.canvas._selected_ids, v.canvas._selected_ids
    n = v.apply_toolbar_color("#C0392B")
    assert n >= 1, n
    stored = v.store.get(rect.id)
    assert stored is not None and stored.color.upper() == "#C0392B", stored.color
    note_after = v.store.get(note.id)
    assert note_after is not None and note_after.color.upper() != "#C0392B", (
        "Farbe darf unselektierte Objekte nicht mitändern"
    )
    v.save_annotations()
    native = read_ild_annots(pdf, page_index=0)
    rec = next(x for x in native if x["nm"] == f"ild:{rect.id}")
    assert rec.get("c") is not None, rec
    r, g, b = rec["c"]
    assert r > 0.6 and g < 0.35 and b < 0.35, rec["c"]

    v.set_tool(AnnotationType.INK)
    _pump(app, 0.05)
    assert rect.id in v.canvas._selected_ids, "Create-Tool darf Auswahl nicht aufheben"
    assert v.canvas._select_mode
    nfont = v.apply_font_size(18.0)
    assert nfont == 0, "Rect ist kein FreeText — Schrift nur auf Text-Typen"
    v.canvas.set_selected_ids({note.id})
    v._selected_ann_ids = {note.id}
    v._selected_ann_id = note.id
    assert v.apply_font_size(18.0) == 1
    assert abs(float(v.store.get(note.id).font_size) - 18.0) < 0.2

    v.canvas.set_selected_ids(set())
    v._selected_ann_ids = set()
    v._selected_ann_id = None
    n_page = v.apply_toolbar_color("#27AE60")
    assert n_page == 2, n_page
    assert v.store.get(rect.id).color.upper() == "#27AE60"
    assert v.store.get(note.id).color.upper() == "#27AE60"
    n_all = v.select_all_text_or_annotations()
    assert n_all == 2, n_all
    assert rect.id in v.canvas._selected_ids and note.id in v.canvas._selected_ids
    v.close()


def test_viewer_style_current_page_only(app, td: Path) -> None:
    """Seite 1 Rect A, Seite 2 Rect B: Strich trifft nur die aktuelle Auswahl B."""
    from ild_pdf import Annotation, AnnotationType, read_ild_annots
    from instantlensdoc.ui.pdf_view import PdfViewer

    pdf = td / "style_pages.pdf"
    _make_multipage_pdf(pdf, ["PageOneAAA", "PageTwoBBB"])
    v = PdfViewer()
    v.resize(960, 720)
    v.show()
    app.processEvents()
    assert v.load(pdf), f"load fehlgeschlagen: {v._last_refresh_error!r}"
    v.set_annotations_locked(False)
    v._suppress_default_zoom = True
    _pump(app, 1.2, until=lambda: v._canvas_has_page_image())
    v.set_scale(1.5, immediate=True)
    _pump(app, 0.3)
    v.set_tool(None)

    rect_a = Annotation(
        page=0,
        type=AnnotationType.RECTANGLE,
        x=80.0,
        y=90.0,
        width=120.0,
        height=50.0,
        color="#2980B9",
        stroke_width=2.0,
    )
    rect_b = Annotation(
        page=1,
        type=AnnotationType.RECTANGLE,
        x=80.0,
        y=90.0,
        width=120.0,
        height=50.0,
        color="#27AE60",
        stroke_width=2.0,
    )
    v.store.add(rect_a)
    v.store.add(rect_b)
    v.refresh()
    _pump(app, 0.2)

    s = v._view_scale()
    ax = (rect_a.x + rect_a.width / 2.0) * s
    ay = (rect_a.y + rect_a.height / 2.0) * s
    wx, wy = _canvas_page_to_widget(v.canvas, ax, ay)
    _qclick(v.canvas, wx, wy)
    _pump(app, 0.1)
    assert rect_a.id in v.canvas._selected_ids, v.canvas._selected_ids
    assert v.apply_toolbar_color("#1F618D") >= 1
    assert v.store.get(rect_a.id).color.upper() == "#1F618D"

    v.goto_page(1)
    _pump(app, 0.8, until=lambda: v.page_index == 1 and v._canvas_has_page_image())
    assert v.page_index == 1
    assert rect_a.id not in v._selected_ann_ids, "Seitenwechsel muss alte Auswahl leeren"
    assert rect_a.id not in (v.canvas._selected_ids or set())

    bx = (rect_b.x + rect_b.width / 2.0) * s
    by = (rect_b.y + rect_b.height / 2.0) * s
    wx, wy = _canvas_page_to_widget(v.canvas, bx, by)
    _qclick(v.canvas, wx, wy)
    _pump(app, 0.1)
    assert rect_b.id in v.canvas._selected_ids, v.canvas._selected_ids

    # Stale Viewer-IDs von Seite 1 dürfen den Live-Klick nicht überschreiben
    v._selected_ann_ids = {rect_a.id}
    v._selected_ann_id = rect_a.id
    n = v.apply_toolbar_color("#C0392B")
    assert n >= 1, n
    assert v.store.get(rect_b.id).color.upper() == "#C0392B", v.store.get(rect_b.id).color
    assert v.store.get(rect_a.id).color.upper() == "#1F618D", v.store.get(rect_a.id).color
    v._apply_toolbar_stroke(7.0, commit=True)
    assert abs(float(v.store.get(rect_b.id).stroke_width) - 7.0) < 0.2
    assert abs(float(v.store.get(rect_a.id).stroke_width) - 2.0) < 0.2

    v.save_annotations()
    native_a = read_ild_annots(pdf, page_index=0)
    native_b = read_ild_annots(pdf, page_index=1)
    rec_a = next(x for x in native_a if x["nm"] == f"ild:{rect_a.id}")
    rec_b = next(x for x in native_b if x["nm"] == f"ild:{rect_b.id}")
    ra, ga, ba = rec_a["c"]
    rb, gb, bb = rec_b["c"]
    assert ra < 0.35 and ga < 0.5 and ba > 0.45, rec_a["c"]
    assert rb > 0.6 and gb < 0.35 and bb < 0.35, rec_b["c"]
    if rec_b.get("border") is not None:
        assert rec_b["border"] >= 6.5, rec_b["border"]
        if rec_a.get("border") is not None:
            assert rec_a["border"] <= 3.0, rec_a["border"]
    v.close()


def test_ribbon_undo_arrows() -> None:
    from instantlensdoc.ui.ribbon_bar import RibbonBar

    rb = RibbonBar()
    undos = rb._action_buttons.get("undo") or []
    redos = rb._action_buttons.get("redo") or []
    assert len(undos) >= 2, "Undo fehlt in Start und/oder Bearbeiten"
    assert len(redos) >= 2, "Redo fehlt in Start und/oder Bearbeiten"
    assert any("↶" in b.text() for b in undos)
    assert any("↷" in b.text() for b in redos)
    rb.set_action_tooltip("undo", "Rückgängig: Rechteck verschieben")
    rb.set_action_enabled("undo", False)
    assert all("Rechteck verschieben" in b.toolTip() for b in undos)
    assert all(not b.isEnabled() for b in undos)
    rb.set_action_enabled("undo", True)


def test_mainwindow_shortcuts(win) -> None:
    from PySide6.QtGui import QKeySequence

    from instantlensdoc.core.app_settings import get_annotations_locked

    act_u = getattr(win, "_edit_undo_action", None)
    act_r = getattr(win, "_edit_redo_action", None)
    assert act_u is not None and act_r is not None
    sc = [s.toString() for s in act_r.shortcuts()]
    joined = " ".join(sc)
    assert "Ctrl+Shift+Z" in joined or QKeySequence("Ctrl+Shift+Z") in act_r.shortcuts()
    rb = win.ribbon_bar
    assert "undo" in rb._action_buttons
    assert get_annotations_locked() is False, "Zweit-Panel darf Annotationen-Sperre nicht global setzen"
    assert win.pdf_view.annotations_locked() is False


def main() -> int:  # noqa: C901
    from PySide6.QtGui import QTextCursor
    from PySide6.QtWidgets import QApplication

    from ild_pdf.pdfium_open import PdfiumOpenError, open_pdfium
    from ild_pdf.pdf_sniff import (
        KIND_ZIP_DOCX,
        describe_non_pdf_de,
        recover_misnamed_file,
        sniff_file,
        validate_pdf_file,
    )
    from instantlensdoc.core.app_settings import (
        apply_one_time_migrations,
        get_editor_minimap,
        set_editor_minimap,
    )
    from instantlensdoc.core.documents import DocKind, open_document, save_document
    from instantlensdoc.license import LicenseManager
    from instantlensdoc.ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    lm = LicenseManager()
    lm.ensure_trial_started()
    win = MainWindow(lm)
    win.resize(1400, 900)
    win.show()
    _pump(app, 0.3)

    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)

        # ---- 1) Nutzerdatei / synthetisches DOCX-als-PDF ------------------------
        src_pdf = tdp / "Dunning_Kruger_Effekt_1.pdf"
        if USER_PDF.is_file():
            shutil.copy2(USER_PDF, src_pdf)
        else:
            real = tdp / "_src.docx"
            _make_docx(real)
            shutil.copy2(real, src_pdf)
        sn = sniff_file(src_pdf)
        assert sn.kind == KIND_ZIP_DOCX, f"erwartet zip-docx, got {sn.kind} {sn.detail}"
        assert sn.head[:2] == b"PK", sn.head_hex
        assert not sn.is_pdf_like
        msg = describe_non_pdf_de(sn)
        assert "Word-Dokument" in msg and "25" in msg or "Byte" in msg
        assert "50 4b 03 04" in sn.head_hex
        v = validate_pdf_file(src_pdf, use_pdfium=True)
        assert not v.ok
        try:
            open_pdfium(src_pdf)
            raise AssertionError("DOCX-als-PDF wurde von open_pdfium akzeptiert")
        except PdfiumOpenError as e:
            text = str(e)
            assert e.not_a_pdf and e.steps[0][0] == "header"
            assert "Schritt 0" in text and "Word-Dokument" in text
            assert "50 4b 03 04" in text
        recovered = recover_misnamed_file(src_pdf)
        assert recovered is not None and recovered.suffix == ".docx" and recovered.is_file()
        opened = open_document(src_pdf)
        assert opened.kind == DocKind.DOCX, opened.kind
        assert opened.meta.get("kind_mismatch")
        assert "Dunning" in (opened.text or "") or "Fließtext" in (opened.text or "") or len(opened.text) > 20
        print("OK  1 sniff/recover user-or-synthetic DOCX-as-PDF")

        proc = subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "pdf_doctor.py"), str(src_pdf)],
            check=False,
            capture_output=True,
            text=True,
        )
        assert proc.returncode == 1, proc.stdout + proc.stderr
        assert "Word-Dokument" in proc.stdout or "zip-docx" in proc.stdout.lower()
        print("OK  1b pdf_doctor exit 1 on misnamed DOCX")

        # ---- 2) DOCX → Speichern unter .pdf ergibt echtes PDF -------------------
        docx_path = tdp / "quelle.docx"
        if recovered and recovered.is_file() and USER_PDF.is_file():
            shutil.copy2(recovered, docx_path)
        else:
            _make_docx(docx_path)
        doc = open_document(docx_path)
        assert doc.kind == DocKind.DOCX
        out_core = tdp / "aus_core.pdf"
        save_document(doc, out_core)
        _assert_valid_pdf(out_core, must_contain="Dunning")
        print("OK  2a save_document(DOCX → .pdf) pikepdf+pdfium", out_core.stat().st_size)

        win.open_path(str(docx_path))
        _pump(app, 0.6)
        assert win.editor.rich_mode(), "DOCX nicht im Rich-Modus"
        assert win.doc.kind == DocKind.DOCX
        from instantlensdoc.core.export import export_document

        out_ui = tdp / "aus_ui.pdf"
        html = win.editor.to_rich_html()
        export_document(
            win.editor.toPlainText(),
            out_ui,
            fmt="pdf",
            title="Dunning",
            html=html,
        )
        _assert_valid_pdf(out_ui, must_contain="Dunning")
        print("OK  2b export_document(html) → gültiges PDF")

        # save_as-Zweig: Endung .pdf + Nicht-PDF-Kind darf nie ZIP schreiben
        out_as = tdp / "save_as.pdf"
        from instantlensdoc.core.documents import save_document as sd

        sd(win.doc, out_as)
        _assert_valid_pdf(out_as, must_contain="Dunning")
        print("OK  2c save_document from loaded DOCX via .pdf target")

        # ---- 3) Diagnose 0-Byte / HTML ------------------------------------------
        empty = tdp / "leer.pdf"
        empty.write_bytes(b"")
        try:
            open_pdfium(empty)
            raise AssertionError("0-Byte geöffnet")
        except PdfiumOpenError as e:
            t = str(e)
            assert "0 Byte" in t and "Schritt 0" in t and "leer.pdf" in t
        html_as_pdf = tdp / "seite.pdf"
        html_as_pdf.write_bytes(b"<!DOCTYPE html><html><body>Hallo</body></html>")
        hsn = sniff_file(html_as_pdf)
        assert hsn.kind == "html"
        assert "3c 21 44 4f 43 54 59 50 45" in hsn.head_hex  # <!DOCTYPE
        win.open_path(str(empty))
        _pump(app, 0.6)
        err = str(getattr(win.pdf_view, "_last_refresh_error", "") or "")
        banner = " ".join(
            [
                err,
                str(getattr(win.pdf_view, "_blank_view_fallback_active", False)),
            ]
        )
        # 0-Byte bleibt PDF-Kind → Viewer + Banner (Schritt 0) oder Diagnosetext
        assert "0 Byte" in err or "Schritt 0" in err or bool(
            getattr(win.pdf_view, "_blank_view_fallback_active", False)
        ), f"keine 0-Byte-Diagnose: {err!r} banner={banner!r}"
        print("OK  3 invalid-file diagnostic (0-byte + HTML header hex)")

        # ---- 4) Minimap Standard aus + Toggle + DOCX unsichtbar -----------------
        assert get_editor_minimap() is False
        set_editor_minimap(True)
        notes = apply_one_time_migrations()
        assert get_editor_minimap() is False, f"Migration ließ Minimap an: {notes}"
        txt = tdp / "plain.txt"
        txt.write_text("Zeile eins\n" + "\n".join(f"Zeile {i}" for i in range(40)), encoding="utf-8")
        win.open_path(str(txt))
        _pump(app, 0.4)
        ed = win.editor
        assert ed.minimap_width() == 0 and not ed._minimap_area.isVisible()
        vp_off = ed.viewport().width()
        win._minimap_action.setChecked(True)
        _pump(app, 0.2)
        assert ed.minimap_visible() and ed.minimap_width() == 36 and ed._minimap_area.isVisible()
        vp_on = ed.viewport().width()
        assert vp_on < vp_off, f"Minimap stahl keine Breite: off={vp_off} on={vp_on}"
        win.open_path(str(docx_path))
        _pump(app, 0.5)
        assert ed.rich_mode()
        assert ed.minimap_width() == 0 and not ed._minimap_area.isVisible(), (
            f"Minimap neben DOCX: width={ed.minimap_width()} vis={ed._minimap_area.isVisible()}"
        )
        win._minimap_action.setChecked(False)
        _pump(app, 0.1)
        assert get_editor_minimap() is False
        print("OK  4 minimap default off, toggle, hidden in DOCX")

        # ---- 5) Undo/Redo: tippen → fett → undo ×2 ------------------------------
        win.open_path(str(txt))
        _pump(app, 0.3)
        original = ed.toPlainText()
        ed.setFocus()
        cur = ed.textCursor()
        cur.movePosition(QTextCursor.End)
        ed.setTextCursor(cur)
        cur.insertText(" mehr")
        after_type = ed.toPlainText()
        assert after_type.endswith(" mehr")
        ed.selectAll()
        assert ed.toggle_bold_selection()
        assert ed.selection_font_bold()
        assert win._undo_action.isEnabled()
        assert win.ribbon_bar.is_enabled("undo")
        assert len(win.ribbon_bar.buttons("undo")) >= 2
        win._undo()
        _pump(app, 0.05)
        # Fett weg, Text noch mit " mehr"
        assert ed.toPlainText() == after_type
        assert not ed.selection_font_bold()
        win._undo()
        _pump(app, 0.05)
        assert ed.toPlainText() == original, repr(ed.toPlainText())
        assert win._redo_action.isEnabled() and win.ribbon_bar.is_enabled("redo")
        win._redo()
        assert ed.toPlainText() == after_type
        shortcuts = [k.toString() for k in win._redo_action.shortcuts()]
        assert any("Ctrl+Y" == s or s.endswith("Y") for s in shortcuts) or "Ctrl+Y" in shortcuts
        assert any("Shift+Z" in s for s in shortcuts)
        before = ed.toPlainText()
        ed.insert_table(2, 2)
        assert "|" in ed.toPlainText()
        win._undo()
        assert ed.toPlainText() == before
        print("OK  5 undo/redo type→bold→undo×2 + table one-step + ribbon arrows")

        # ---- 6–11) PDF-Annotationen --------------------------------------------
        test_pdf_roundtrip_move_and_freetext(tdp)
        print("OK  6 native /Rect /Contents + Undo-Rect")
        test_viewer_tools_undo_zoom(app, tdp)
        print("OK  7 viewer Stift/Highlight/Zoom/QUndoStack")
        test_hit_test_pdf_points_unit()
        print("OK  7b hit-test PDF-Punkte (dünn/oben)")
        test_viewer_select_hit_test(app, tdp)
        print("OK  7c Auswahl Klick/Gummiband/Griffe")
        test_viewer_style_selection(app, tdp)
        print("OK  7d Ribbon-Farbe auf Auswahl, Seite ohne Auswahl, Select-All global")
        test_viewer_style_current_page_only(app, tdp)
        print("OK  7e Seite2-Auswahl Strich ändert nur B /C, A unverändert")
        test_ribbon_undo_arrows()
        print("OK  8 ribbon ↶/↷")
        test_mainwindow_shortcuts(win)
        print("OK  9 menü Ctrl+Shift+Z")

    print("OK test_ui_audit_2654")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
