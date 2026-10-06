#!/usr/bin/env python3
"""OCR-Fixture → Word-Suite-Editor → Fett / Suchen / Speichern DOCX.

Beweist (User 2.6.54): OCR- und PDF-Text landen als rich ``QTextDocument``
(gleiche Dokumentart wie ein geöffnetes DOCX), nicht als Plain-Blob/Bild.
Alle Textfunktionen (Fett, Finden, Speichern DOCX) greifen.

Aufruf: ``QT_QPA_PLATFORM=offscreen python3 scripts/test_ocr_editable.py``
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["ILD_SMOKE_QT"] = "1"
os.environ.setdefault("ILD_SKIP_DEPS_CHECK", "1")
os.environ.setdefault("ILD_NO_SESSION", "1")
os.environ.setdefault("ILD_NO_SPLASH", "1")
os.environ["XDG_CONFIG_HOME"] = tempfile.mkdtemp(prefix="ild-ocr-editable-cfg-")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

FIXTURE = ROOT / "tests" / "fixtures" / "ocr_sample.ildocr.txt"


def _fail(msg: str) -> None:
    print(f"FAIL: {msg}", file=sys.stderr)
    raise SystemExit(1)


def _ok(msg: str) -> None:
    print(f"OK: {msg}")


def test_core_ocr_fixture_to_docx() -> None:
    from instantlensdoc.core.documents import DocKind, save_document
    from instantlensdoc.core.ocr_word_suite import (
        blocks_to_word_suite_html,
        import_ildocr_sidecar,
        open_ocr_result,
        word_suite_to_document,
    )

    if not FIXTURE.is_file():
        _fail(f"Fixture fehlt: {FIXTURE}")
    ws = import_ildocr_sidecar(FIXTURE, auto_format=False)
    if "EINLEITUNG" not in ws.text:
        _fail("Fixture-Text fehlt im Word-Suite-Dokument")
    if not (ws.html or "").strip():
        _fail("Word-Suite-HTML fehlt (kein QTextDocument-Pfad)")
    if "<h1" not in ws.html.lower() and "<p" not in ws.html.lower():
        _fail(f"HTML ohne Absätze: {ws.html[:240]}")
    if "page-break" not in ws.html and ws.block_count < 2:
        _fail("erwartete Blöcke/Seitenumbrüche")
    if not ws.source_comment or "OCR-Quelle" not in ws.source_comment:
        _fail("Quell-Kommentar (Sidecar/PDF) fehlt")
    if not ws.sidecar:
        _fail("Sidecar-Link fehlt")

    html = blocks_to_word_suite_html(ws.blocks, text=ws.text, title=ws.title)
    if "EINLEITUNG" not in html:
        _fail("HTML ohne Fixture-Überschrift")

    doc = word_suite_to_document(ws)
    if doc.kind != DocKind.DOCX:
        _fail(f"Arbeitsdokument kind={doc.kind} — erwartet DOCX wie geöffnetes Word")
    if not doc.meta.get("rich_text") or not doc.meta.get("html"):
        _fail("Document.meta.html/rich_text fehlt")

    opened = open_ocr_result(FIXTURE, auto_format=False, title="Word-Suite — Fixture")
    if opened.kind != DocKind.DOCX:
        _fail(f"open_ocr_result kind={opened.kind}")
    if "EINLEITUNG" not in opened.text:
        _fail("open_ocr_result ohne Fixture-Text")

    typed = open_ocr_result(text="Getippter Absatz eins.\n\nZweiter Absatz.", auto_format=False)
    if typed.kind != DocKind.DOCX or "Getippter" not in typed.text:
        _fail("getippter Text nicht als Word-Suite-DOCX")

    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "ocr-fixture.docx"
        save_document(opened, out)
        if not out.is_file() or out.stat().st_size < 200:
            _fail("save_document DOCX leer/fehlt")
        from docx import Document as DocxDocument

        d = DocxDocument(str(out))
        body = "\n".join(p.text for p in d.paragraphs)
        if "EINLEITUNG" not in body and "Fließtext" not in body:
            _fail(f"DOCX ohne OCR-Text: {body[:300]!r}")
    _ok("core: fixture ildocr → HTML/DOCX + Sidecar-Kommentar + save")


def test_editor_ocr_bold_find_save() -> None:
    from PySide6.QtGui import QFont, QTextCursor
    from PySide6.QtWidgets import QApplication

    from instantlensdoc.core.documents import DocKind, save_document
    from instantlensdoc.license import LicenseManager
    from instantlensdoc.ui.main_window import MainWindow

    app = QApplication.instance() or QApplication(sys.argv)
    lm = LicenseManager()
    lm.ensure_trial_started()
    win = MainWindow(lm)

    assert callable(getattr(win, "open_ocr_result", None))
    assert callable(getattr(win, "_handoff_ocr_to_word_suite", None))
    assert callable(getattr(win, "_present_word_suite_document", None))

    ok = win.open_ocr_result(
        path=FIXTURE,
        title="Word-Suite — OCR-Fixture",
        auto_format=False,
        source_path="scan-demo.pdf",
        source_page=2,
    )
    if not ok:
        _fail("open_ocr_result/handoff schlug fehl")
    if win.stack.currentWidget() is not win.editor_pane:
        _fail("OCR landete nicht im Word-Suite-Editor")
    if not win.editor.rich_mode():
        _fail("Editor nicht im Rich-Modus (kein QTextDocument wie DOCX)")
    if win.doc is None or win.doc.kind != DocKind.DOCX:
        _fail(f"doc.kind={getattr(win.doc, 'kind', None)} — erwartet DOCX")
    if not win._editor_document_active():
        _fail("Editor-Aktionen (Fett/Suchen) nach OCR nicht aktiv")

    plain = win.editor.toPlainText()
    if "EINLEITUNG" not in plain or "Fließtext" not in plain:
        _fail(f"Editor-Plain ohne Fixture: {plain[:240]!r}")
    if "**" in plain or "__EINLEITUNG" in plain:
        _fail("Markdown-Marker im OCR-Dokument")

    needle = "Fließtext"
    if needle not in plain:
        needle = "EINLEITUNG"
    n_find = win.editor.find_and_highlight(needle)
    if n_find < 1:
        _fail(f"Finden findet {needle!r} nicht in OCR-Text")

    idx = plain.index(needle)
    cur = win.editor.textCursor()
    cur.setPosition(idx)
    cur.setPosition(idx + len(needle), QTextCursor.KeepAnchor)
    win.editor.setTextCursor(cur)
    if not win.editor.toggle_bold_selection():
        _fail("toggle_bold_selection auf OCR-Text fehlgeschlagen")
    win._sync_editor_rich_meta()
    probe = win.editor.textCursor()
    probe.setPosition(idx)
    probe.setPosition(idx + len(needle), QTextCursor.KeepAnchor)
    if probe.charFormat().fontWeight() < QFont.Bold:
        _fail("Fett nicht als QTextCharFormat auf OCR-Text")
    if "*" in win.editor.toPlainText().replace(needle, ""):
        # Sternchen nur akzeptieren wenn sie im Fixture-Fließtext vorkamen
        if "**" in win.editor.toPlainText():
            _fail("Fett setzte Markdown-Asterisks")

    html = win.editor.to_rich_html()
    if "font-weight" not in html.lower() and "<b" not in html.lower() and "700" not in html:
        if "600" not in html:
            _fail(f"to_rich_html ohne Bold nach OCR-Fett: {html[:300]}")

    src_comment = (win.doc.meta or {}).get("source_comment") or ""
    sidecar = (win.doc.meta or {}).get("sidecar") or ""
    if not sidecar and "Quelle" not in src_comment and "scan-demo" not in str(
        (win.doc.meta or {}).get("source_path") or ""
    ):
        _fail("Quelle Scan/PDF nicht verknüpft (sidecar/comment)")

    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "ocr-editor.docx"
        save_document(win.doc, out)
        if not out.is_file():
            _fail("DOCX nach Editor-Fett nicht geschrieben")
        from docx import Document as DocxDocument

        d = DocxDocument(str(out))
        flags = []
        joined = []
        for para in d.paragraphs:
            joined.append(para.text)
            for run in para.runs:
                if needle in (run.text or ""):
                    flags.append(bool(run.bold))
        body = "\n".join(joined)
        if needle not in body and "EINLEITUNG" not in body:
            _fail(f"gespeichertes DOCX ohne OCR-Text: {body[:300]!r}")
        if flags and not any(flags):
            _fail(f"Bold-Run nach OCR→Editor→DOCX verloren: {flags}")

    _ok("qt: OCR-Fixture → Editor (DOCX/rich) → Fett/Finden/save docx")

    test_ocr_selection_or_document_on_editor(win)

    test_editor_no_control_glyphs_highlight_font_docx(win)

    assert app is not None


def test_ocr_selection_or_document_on_editor(win) -> None:
    """Auswahl → Tool nur auf Selektion; keine Auswahl → gesamter OCR-Text.

    Bearbeiten/Ribbon laufen über ``self.editor`` nach ``_present_word_suite_document``.
    """
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QTextCursor

    from instantlensdoc.core.documents import DocKind

    if win.stack.currentWidget() is not win.editor_pane:
        _fail("OCR-Dokument nicht im gleichen Editor-Stack (editor_pane)")
    if win.doc is None or win.doc.kind != DocKind.DOCX:
        _fail(f"OCR-Dokument kind={getattr(win.doc, 'kind', None)}")
    if not win._editor_document_active() or win._pdf_tab_active():
        _fail("Bearbeiten/Ribbon treffen OCR nicht (_editor_document_active)")
    if not callable(getattr(win.editor, "selection_or_document_cursor", None)):
        _fail("Editor ohne selection_or_document_cursor")

    try:
        win.editor.setFocus(Qt.OtherFocusReason)
    except Exception:
        pass

    original = win.editor.toPlainText()
    needle = "durchsuchbar"
    if needle not in original:
        _fail("Fixture-Needle fehlt für Auswahl-Test")

    idx = original.index(needle)
    cur = win.editor.textCursor()
    cur.setPosition(idx)
    cur.setPosition(idx + len(needle), QTextCursor.KeepAnchor)
    win.editor.setTextCursor(cur)
    if not win.editor.transform_document_case("upper"):
        _fail("transform_document_case auf OCR-Auswahl fehlgeschlagen")
    after_sel = win.editor.toPlainText()
    if needle.upper() not in after_sel:
        _fail("Auswahl nicht umgewandelt")
    if "EINLEITUNG" not in after_sel:
        _fail("ohne-Auswahl-Text außerhalb der Selektion verändert (EINLEITUNG weg)")
    if after_sel.replace(needle.upper(), needle) != original:
        _fail("Case-Tool auf Auswahl hat Text außerhalb der Selektion verändert")

    cur = win.editor.textCursor()
    cur.clearSelection()
    cur.setPosition(0)
    win.editor.setTextCursor(cur)
    if cur.hasSelection():
        _fail("Selektion nicht geleert")
    if not win.editor.transform_document_case("lower"):
        _fail("transform_document_case ohne Auswahl (ganzes OCR) fehlgeschlagen")
    after_all = win.editor.toPlainText()
    if after_all != original.lower():
        _fail(f"ohne Auswahl nicht gesamter OCR-Text: {after_all[:180]!r}")

    # Highlight: Selektion vs. gesamtes OCR-Dokument
    win.editor.clear_highlight_formats()
    idx = after_all.index("einleitung")
    cur = win.editor.textCursor()
    cur.setPosition(idx)
    cur.setPosition(idx + len("einleitung"), QTextCursor.KeepAnchor)
    win.editor.setTextCursor(cur)
    if not win.editor.highlight_selection():
        _fail("highlight_selection auf OCR-Auswahl fehlgeschlagen")
    if not win.editor.selection_highlighted():
        _fail("OCR-Auswahl nicht markiert")
    win.editor.clear_highlight_formats()
    cur = win.editor.textCursor()
    cur.clearSelection()
    cur.setPosition(0)
    win.editor.setTextCursor(cur)
    if not win.editor.highlight_selection():
        _fail("highlight_selection ohne Auswahl (ganzes OCR) fehlgeschlagen")
    probe = QTextCursor(win.editor.document())
    probe.setPosition(0)
    probe.setPosition(min(8, len(after_all)), QTextCursor.KeepAnchor)
    p2 = win.editor._selection_probe_format(probe)
    if p2.background().style() == Qt.NoBrush:
        _fail("ohne Auswahl kein Highlight auf gesamtem OCR-Text")

    # Bearbeiten-Menü-Pfad (Fett) trifft denselben Editor
    idx = win.editor.toPlainText().index("fazit") if "fazit" in win.editor.toPlainText() else 0
    cur = win.editor.textCursor()
    cur.setPosition(idx)
    cur.setPosition(idx + 5, QTextCursor.KeepAnchor)
    win.editor.setTextCursor(cur)
    win._toggle_bold()
    if not win.editor.selection_font_bold():
        _fail("Bearbeiten-Fett (_toggle_bold) trifft OCR-Editor nicht")

    _ok("qt: OCR Auswahl→Selektion / keine Auswahl→gesamter Text + Bearbeiten-Hit")


def test_pdf_extract_kind_not_pdf() -> None:
    """PDF-extrahierter Text darf nicht als DocKind.PDF (Editor-Tools aus) bleiben."""
    from instantlensdoc.core.documents import DocKind
    from instantlensdoc.core.ocr_word_suite import open_ocr_result

    doc = open_ocr_result(
        text="Seite Eins Absatz.\n\nSeite Zwei Absatz.",
        title="Word-Suite — PDF-Text",
        auto_format=False,
        source_path="/tmp/demo.pdf",
        source_page=1,
    )
    if doc.kind == DocKind.PDF or doc.kind == DocKind.IMAGE:
        _fail("PDF-Text als Bild/PDF — Word-Suite-Tools greifen nicht")
    if doc.kind != DocKind.DOCX:
        _fail(f"PDF-Text kind={doc.kind}")
    if not doc.meta.get("html"):
        _fail("PDF-Text ohne HTML")
    if "demo.pdf" not in str(doc.meta.get("source_path") or "") and "Seite 1" not in str(
        doc.meta.get("source_comment") or ""
    ):
        _fail("PDF-Quelle nicht verknüpft")
    _ok("pdf-extract: Document ist DOCX mit Quelle")


def test_scan_session_import() -> None:
    from instantlensdoc.core.documents import DocKind
    from instantlensdoc.core.ocr import OcrOutputMode, OcrResult
    from instantlensdoc.core.ocr_word_suite import scan_session_to_word_suite
    from instantlensdoc.core.scan import ScanPageResult, ScanSessionResult

    ocr = OcrResult(
        text="Scanzeile Alpha\n\nScanzeile Beta",
        lang="deu+eng",
        mode=OcrOutputMode.EDITABLE_TEXT,
        source_label="page1.png",
    )
    session = ScanSessionResult(
        pages=[
            ScanPageResult(source_label="page1.png", page_index=0, ocr=ocr, sidecar=None),
            ScanPageResult(
                source_label="page2.png",
                page_index=1,
                ocr=OcrResult(
                    text="Zweite Seite Text",
                    lang="deu+eng",
                    mode=OcrOutputMode.EDITABLE_TEXT,
                ),
            ),
        ],
        pdf_path=Path("/tmp/scan-session.pdf"),
        ocr_enabled=True,
    )
    ws = scan_session_to_word_suite(session, auto_format=False)
    if "Scanzeile Alpha" not in ws.text and "Alpha" not in "".join(b.text for b in ws.blocks):
        _fail("Scan-OCR-Text fehlt")
    pages = {int(b.page) for b in ws.blocks if b.page}
    if pages and pages != {1, 2} and 2 not in pages:
        _fail(f"Scan-Seitenumbrüche fehlen: {pages}")
    if "page-break" not in ws.html and len({b.page for b in ws.blocks if b.page}) < 2:
        _fail("Scan-HTML ohne Seitenumbruch")
    doc = session.to_word_suite_document(auto_format=False)
    if doc.kind != DocKind.DOCX:
        _fail(f"Scan-Import kind={doc.kind}")
    _ok("scan: Session → Word-Suite DOCX mit Seiten")


_CONTROL_GLYPHS = ("\x0c", "\u2028", "\u2029", "\ufffd", "\u00b6", "\u200e", "\u200f", "\u202a")

_HOCR_SAMPLE = """
<div class='ocr_page' title='image "scan.png"; bbox 0 0 800 1100; ppageno 0'>
<p class='ocr_par' title='bbox 220 40 580 88'>
<span class='ocr_line' title='bbox 220 40 580 88'>
<span class='ocrx_word' title='bbox 230 40 560 88; x_wconf 95; x_font Times_New_Roman_Bold; x_fsize 18'><strong>EINLEITUNG</strong></span>
</span>
</p>
<p class='ocr_par' title='bbox 40 120 760 210'>
<span class='ocr_line' title='bbox 40 120 760 160'>
<span class='ocrx_word' title='bbox 40 120 200 150; x_font Calibri; x_fsize 11'>Fliesstext</span>
<span class='ocrx_word' title='bbox 210 120 320 150; x_font Calibri_Italic; x_fsize 11'><em>kursiv</em></span>
</span>
</p>
</div>
"""


def test_core_sanitize_and_hocr_layout() -> None:
    from instantlensdoc.core.ocr_word_suite import (
        blocks_to_word_suite_html,
        open_ocr_result,
        parse_hocr_to_blocks,
        sanitize_ocr_visible_text,
    )

    dirty = "Hallo\x0cWelt\u2028Zeile\ufffd\u200e"
    clean = sanitize_ocr_visible_text(dirty)
    for g in _CONTROL_GLYPHS:
        if g in clean:
            _fail(f"sanitize liess Steuerzeichen {g!r}")
    if "Hallo" not in clean or "Welt" not in clean:
        _fail("sanitize hat Fließtext entfernt")

    blocks = parse_hocr_to_blocks(_HOCR_SAMPLE)
    if len(blocks) < 2:
        _fail(f"hOCR-Absätze fehlen: {len(blocks)}")
    joined = " ".join(b.text for b in blocks)
    if "EINLEITUNG" not in joined or "Fliesstext" not in joined:
        _fail(f"hOCR-Text fehlt: {joined!r}")
    if not any(b.bold or b.is_heading for b in blocks):
        _fail("hOCR ohne Fett/Überschrift")
    if not any(b.italic for b in blocks):
        _fail("hOCR ohne Kursiv")
    if not any("Times" in (b.font_name or "") or "Calibri" in (b.font_name or "") for b in blocks):
        _fail(f"hOCR ohne Fontnamen: {[b.font_name for b in blocks]}")
    html = blocks_to_word_suite_html(blocks, title="hOCR")
    low = html.lower()
    if "font-family" not in low or "font-size" not in low or "text-align" not in low:
        _fail(f"hOCR-HTML ohne Layout: {html[:400]}")
    if "<b>" not in low and "<h1" not in low:
        _fail("hOCR-HTML ohne Fett/Überschrift-Tag")
    if "<i>" not in low:
        _fail("hOCR-HTML ohne Kursiv-Tag")
    for g in _CONTROL_GLYPHS:
        if g in html:
            _fail(f"hOCR-HTML enthält Steuerzeichen {g!r}")

    dumped = open_ocr_result(
        text="Dump\x0cmit\u2028Formfeed\ufffd",
        auto_format=False,
        title="Word-Suite — Dirty",
    )
    blob = (dumped.text or "") + str((dumped.meta or {}).get("html") or "")
    for g in ("\x0c", "\u2028", "\ufffd"):
        if g in blob:
            _fail(f"open_ocr_result dumpte Steuerzeichen {g!r}")
    if "Dump" not in dumped.text:
        _fail("Dirty-OCR-Text verloren")
    _ok("core: Steuerzeichen weg, hOCR Font/Fett/Kursiv/Align")


def test_editor_no_control_glyphs_highlight_font_docx(win) -> None:
    """OCR-Editor: keine sichtbaren Steuerzeichen; Highlighter + Schrift überleben DOCX."""
    from PySide6.QtGui import QFont, QTextCursor

    from instantlensdoc.core.documents import save_document

    ok = win.open_ocr_result(
        path=FIXTURE,
        title="Word-Suite — OCR-Fixture",
        auto_format=False,
    )
    if not ok:
        _fail("open_ocr_result für Highlight/Font-Test fehlgeschlagen")

    ed = win.editor
    try:
        if ed.special_chars_visible():
            _fail("Rich-OCR zeigt Sonderzeichen (ShowTabsAndSpaces)")
    except Exception:
        pass
    plain = ed.toPlainText()
    for g in _CONTROL_GLYPHS:
        if g in plain:
            _fail(f"Editor-Plain enthält Steuerzeichen {g!r}")
    html = ed.to_rich_html()
    for g in ("\x0c", "\u2028", "\ufffd", "&#12;"):
        if g in html:
            _fail(f"Editor-HTML enthält Steuerzeichen {g!r}")

    needle = "EINLEITUNG"
    if needle not in plain:
        _fail("EINLEITUNG fehlt vor Highlight/Font")
    idx = plain.index(needle)
    cur = ed.textCursor()
    cur.setPosition(idx)
    cur.setPosition(idx + len(needle), QTextCursor.KeepAnchor)
    ed.setTextCursor(cur)
    if not ed.apply_font_family("Arial"):
        _fail("apply_font_family auf OCR-Auswahl fehlgeschlagen")
    if not ed.highlight_selection("#FFE066"):
        _fail("highlight_selection auf OCR-Auswahl fehlgeschlagen")
    win._sync_editor_rich_meta()
    html = ed.to_rich_html().lower()
    if "arial" not in html:
        _fail(f"to_rich_html ohne Arial nach Font-Picker-Pfad: {html[:360]}")
    if "background" not in html and "ffe066" not in html:
        _fail(f"to_rich_html ohne Highlight-Hintergrund: {html[:360]}")

    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "ocr-highlight-font.docx"
        save_document(win.doc, out)
        if not out.is_file():
            _fail("DOCX nach Highlight/Font nicht geschrieben")
        from docx import Document as DocxDocument

        d = DocxDocument(str(out))
        names = []
        highs = []
        body_parts = []
        for para in d.paragraphs:
            body_parts.append(para.text)
            for run in para.runs:
                if needle not in (run.text or "") and needle.title() not in (run.text or ""):
                    if "EINLEITUNG" not in (run.text or "").upper():
                        continue
                if run.font.name:
                    names.append(str(run.font.name))
                hl = getattr(run.font, "highlight_color", None)
                if hl is not None:
                    highs.append(str(hl))
        body = "\n".join(body_parts)
        if "EINLEITUNG" not in body.upper():
            _fail(f"DOCX ohne OCR-Überschrift: {body[:240]!r}")
        if names and not any("arial" in n.lower() for n in names):
            _fail(f"DOCX Font nicht Arial: {names}")
        if not highs:
            _fail(f"DOCX ohne persistentes Highlight: names={names}")
    _ok("qt: keine Steuerzeichen; Highlight+Arial überleben DOCX")


def main() -> int:
    test_core_ocr_fixture_to_docx()
    test_core_sanitize_and_hocr_layout()
    test_pdf_extract_kind_not_pdf()
    test_scan_session_import()
    test_editor_ocr_bold_find_save()
    print("ALL OK — ocr-editable Word-Suite QTextDocument")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
