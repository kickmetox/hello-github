#!/usr/bin/env python3
"""Smoke/Unit: DOCX Rich-Text + QTextCharFormat Bold/Italic/Underline — 2.6.49.

Prüft:
1. Sample-DOCX öffnen → Zeichenformate in HTML vorhanden
2. Bold auf Auswahl → keine Asterisks im Plaintext
3. Underline spannt Buchstaben (fontUnderline auf Letters)
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _fail(msg: str) -> None:
    print(f"FAIL: {msg}", file=sys.stderr)
    raise SystemExit(1)


def _ok(msg: str) -> None:
    print(f"OK: {msg}")


def test_docx_roundtrip_formats() -> None:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    from instantlensdoc.core.documents import DocKind, open_document, save_document
    from instantlensdoc.core.richtext_docx import (
        docx_to_html,
        html_has_char_formats,
        html_to_docx,
    )

    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "sample.docx"
        d = DocxDocument()
        p = d.add_paragraph()
        r = p.add_run("Bold")
        r.bold = True
        r2 = p.add_run(" ")
        r3 = p.add_run("Italic")
        r3.italic = True
        r4 = p.add_run(" ")
        r5 = p.add_run("Under")
        r5.underline = True
        r6 = p.add_run(" plain")
        r6.font.size = Pt(11)
        d.add_paragraph("Zweiter Absatz.")
        d.save(str(src))

        html = docx_to_html(src)
        if not html_has_char_formats(html):
            _fail("docx_to_html ohne Bold/Italic/Underline-Tags")
        if "<b>Bold</b>" not in html and "<b>" not in html:
            _fail(f"Bold-Run fehlt in HTML: {html[:300]}")
        if "<i>Italic</i>" not in html and "<i>" not in html:
            _fail(f"Italic-Run fehlt in HTML: {html[:300]}")
        if "<u>Under</u>" not in html and "<u>" not in html:
            _fail(f"Underline-Run fehlt in HTML: {html[:300]}")

        opened = open_document(src)
        assert opened.kind == DocKind.DOCX
        assert opened.meta.get("html"), "open_document meta.html fehlt"
        assert html_has_char_formats(str(opened.meta.get("html")))
        assert "Bold" in opened.text and "Italic" in opened.text
        assert "**" not in opened.text and "__" not in opened.text

        out = td_path / "roundtrip.docx"
        html_to_docx(str(opened.meta["html"]), out, title="RT")
        d2 = DocxDocument(str(out))
        flags = []
        for para in d2.paragraphs:
            for run in para.runs:
                if not (run.text or "").strip():
                    continue
                flags.append(
                    (
                        run.text,
                        bool(run.bold),
                        bool(run.italic),
                        bool(run.underline),
                    )
                )
        by_text = {t: (b, i, u) for t, b, i, u in flags}
        if "Bold" not in by_text or not by_text["Bold"][0]:
            _fail(f"Roundtrip Bold verloren: {flags}")
        if "Italic" not in by_text or not by_text["Italic"][1]:
            _fail(f"Roundtrip Italic verloren: {flags}")
        if "Under" not in by_text or not by_text["Under"][2]:
            _fail(f"Roundtrip Underline verloren: {flags}")

        opened.meta["html"] = opened.meta["html"]
        save_document(opened, td_path / "saved.docx")
        assert (td_path / "saved.docx").is_file()
    _ok("docx open → char formats; roundtrip/save")


def test_qt_char_format_toggles() -> None:
    from PySide6.QtGui import QFont, QTextCursor
    from PySide6.QtWidgets import QApplication

    from instantlensdoc.ui.editor import TextEditor

    app = QApplication.instance() or QApplication(sys.argv)
    ed = TextEditor()
    ed.setPlainText("wort")
    cur = ed.textCursor()
    cur.select(QTextCursor.Document)
    ed.setTextCursor(cur)

    assert ed.toggle_bold_selection()
    plain = ed.toPlainText()
    if "*" in plain or "**" in plain:
        _fail(f"Bold fügte Markdown-Marker ein: {plain!r}")
    if "wort" not in plain:
        _fail(f"Text verloren nach Bold: {plain!r}")
    probe = ed.textCursor()
    probe.setPosition(0)
    probe.setPosition(4, QTextCursor.KeepAnchor)
    if probe.charFormat().fontWeight() < QFont.Bold:
        _fail("Bold nicht als QTextCharFormat gesetzt")

    assert ed.toggle_italic_selection()
    if "*" in ed.toPlainText():
        _fail(f"Italic fügte Sternchen ein: {ed.toPlainText()!r}")
    if not probe.charFormat().fontItalic() and not ed.selection_font_italic():
        # re-probe
        probe2 = ed.textCursor()
        probe2.setPosition(0)
        probe2.setPosition(4, QTextCursor.KeepAnchor)
        if not probe2.charFormat().fontItalic():
            _fail("Italic nicht als QTextCharFormat gesetzt")

    assert ed.toggle_underline_selection()
    if "__" in ed.toPlainText() or "_" in ed.toPlainText().replace("wort", ""):
        # underscore markers must not appear
        if "__" in ed.toPlainText():
            _fail(f"Underline fügte __ Marker ein: {ed.toPlainText()!r}")
    probe3 = ed.textCursor()
    probe3.setPosition(0)
    probe3.setPosition(4, QTextCursor.KeepAnchor)
    if not probe3.charFormat().fontUnderline():
        _fail("Underline nicht auf Buchstaben (fontUnderline fehlt)")
    # Explizit: jedes Zeichen der Selection inkl. Letters
    for i, ch in enumerate("wort"):
        c = ed.textCursor()
        c.setPosition(i)
        c.setPosition(i + 1, QTextCursor.KeepAnchor)
        if ch.isalpha() and not c.charFormat().fontUnderline():
            _fail(f"Underline fehlt auf Buchstabe {ch!r} an Pos {i}")

    html = ed.to_rich_html()
    if "font-weight" not in html.lower() and "<b" not in html.lower():
        # Qt may encode bold as span style
        if "600" not in html and "700" not in html and "bold" not in html.lower():
            _fail(f"to_rich_html ohne Bold-Hinweis: {html[:400]}")
    _ok("toggle bold/italic/underline → QTextCharFormat, keine Asterisks")
    # keep app ref
    assert app is not None


def test_load_docx_into_editor_formats() -> None:
    from docx import Document as DocxDocument
    from PySide6.QtGui import QFont, QTextCursor
    from PySide6.QtWidgets import QApplication

    from instantlensdoc.core.documents import open_document
    from instantlensdoc.ui.editor import TextEditor

    app = QApplication.instance() or QApplication(sys.argv)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "fmt.docx"
        d = DocxDocument()
        p = d.add_paragraph()
        rb = p.add_run("Alpha")
        rb.bold = True
        p.add_run(" ")
        ru = p.add_run("Beta")
        ru.underline = True
        d.save(str(src))
        doc = open_document(src)
        ed = TextEditor()
        ed.set_rich_html(str(doc.meta.get("html") or ""))
        plain = ed.toPlainText()
        if "**" in plain or "__" in plain:
            _fail(f"Editor-Plain nach DOCX enthält Marker: {plain!r}")
        i = plain.index("Alpha")
        c = ed.textCursor()
        c.setPosition(i)
        c.setPosition(i + 5, QTextCursor.KeepAnchor)
        if c.charFormat().fontWeight() < QFont.Bold:
            _fail("Alpha nicht fett nach DOCX→Editor")
        j = plain.index("Beta")
        c.setPosition(j)
        c.setPosition(j + 4, QTextCursor.KeepAnchor)
        if not c.charFormat().fontUnderline():
            _fail("Beta nicht unterstrichen nach DOCX→Editor")
    _ok("open sample docx → editor char formats present")
    assert app is not None


def main() -> int:
    test_docx_roundtrip_formats()
    test_qt_char_format_toggles()
    test_load_docx_into_editor_formats()
    print("ALL OK — 2.6.49 docx-richtext")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
